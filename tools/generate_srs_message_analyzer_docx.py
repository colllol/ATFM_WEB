from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "TaiLieu" / "SRS_Cong_Cu_Phan_Tich_Dien_Van.docx"


def set_font(style, size, bold=False):
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.bold = bold


def shade(cell, fill="D9EAF7"):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def repeat_header(row):
    props = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    props.append(element)


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    repeat_header(table.rows[0])
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        shade(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    document.add_paragraph()


def add_bullets(document, items):
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def build():
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)

    set_font(document.styles["Normal"], 12)
    set_font(document.styles["Title"], 20, True)
    set_font(document.styles["Heading 1"], 15, True)
    set_font(document.styles["Heading 2"], 13, True)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(80)
    run = title.add_run("TÀI LIỆU ĐẶC TẢ YÊU CẦU PHẦN MỀM")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(20)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("CÔNG CỤ PHÂN TÍCH ĐIỆN VĂN")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)
    info = document.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_before = Pt(24)
    info.add_run(
        "Hệ thống tiếp nhận kết quả: ATFM_WEB\n"
        "Giao diện khai thác: Receive_LogFile/Inbox.aspx?Menu_ID=47\n"
        "Phiên bản tài liệu: 1.0\nNgày lập: 12/08/2026"
    )
    document.add_page_break()

    document.add_heading("1. Thông tin chung", level=1)
    add_table(document, ["Nội dung", "Giá trị"], [
        ("Tên tính năng", "Công cụ phân tích điện văn"),
        ("Loại ứng dụng", "Công cụ đọc, phân tích và đồng bộ điện văn"),
        ("Hệ thống nhận kết quả", "ATFM_WEB"),
        ("Màn hình xem kết quả", "/Receive_LogFile/Inbox.aspx?Menu_ID=47"),
        ("Nguồn dữ liệu Inbox", "T_RECEIVE_LOG; T_RECEIVE_LOGFILE"),
        ("Gói truy vấn Inbox", "MESSAGE_PKG"),
        ("Đối tượng sử dụng", "Người vận hành công cụ; người khai thác điện văn"),
    ])

    document.add_heading("2. Mục đích và phạm vi", level=1)
    document.add_paragraph(
        "Công cụ đọc dữ liệu điện văn từ nguồn được cấu hình, nhận diện từng điện văn, "
        "trích xuất thông tin cần thiết, kiểm tra hợp lệ và ghi kết quả vào kho dữ liệu "
        "Inbox của ATFM_WEB. Người khai thác xem kết quả tại Menu_ID=47."
    )
    add_bullets(document, [
        "Cấu hình nguồn đọc điện văn và kết nối cơ sở dữ liệu.",
        "Đọc mới hoặc quét lại dữ liệu nguồn theo phạm vi được chọn.",
        "Tách từng điện văn từ dữ liệu thô và giữ nguyên nội dung gốc.",
        "Trích xuất số điện văn, thời điểm nhận, nguồn gửi, loại và nội dung.",
        "Phát hiện bản ghi trùng, bản ghi lỗi và ghi nhật ký xử lý.",
        "Đồng bộ điện văn hợp lệ để hiển thị tại Inbox của ATFM_WEB.",
    ])
    document.add_paragraph(
        "Tài liệu chỉ tập trung vào form Công cụ phân tích điện văn và giao diện Inbox nhận kết quả. "
        "Không đặc tả nghiệp vụ xử lý kế hoạch bay phát sinh từ từng loại điện văn."
    )

    document.add_heading("3. Tác nhân và hệ thống liên quan", level=1)
    add_table(document, ["Tác nhân/thành phần", "Vai trò"], [
        ("Người vận hành", "Cấu hình, bắt đầu/dừng phân tích, theo dõi tiến độ và xử lý lỗi."),
        ("Nguồn điện văn", "Cung cấp tệp hoặc dữ liệu điện văn thô cho công cụ."),
        ("Công cụ phân tích", "Đọc, tách, chuẩn hóa, kiểm tra và đồng bộ điện văn."),
        ("Oracle ATFM", "Lưu điện văn hợp lệ và nhật ký cần thiết."),
        ("Người khai thác Inbox", "Tìm kiếm, xem chi tiết, chuyển tiếp hoặc in điện văn."),
        ("ATFM_WEB Inbox", "Hiển thị kết quả phân tích qua MESSAGE_PKG."),
    ])

    document.add_heading("4. Luồng xử lý tổng quát", level=1)
    document.add_paragraph(
        "Chọn nguồn/cấu hình → Kiểm tra kết nối → Đọc dữ liệu → Tách điện văn → "
        "Phân tích trường thông tin → Kiểm tra hợp lệ/trùng → Ghi dữ liệu Inbox → "
        "Ghi nhật ký → Người dùng xem tại ATFM_WEB Inbox."
    )
    add_table(document, ["Trạng thái", "Ý nghĩa"], [
        ("Chờ xử lý", "Nguồn đã được phát hiện nhưng chưa phân tích."),
        ("Đang xử lý", "Công cụ đang đọc hoặc phân tích dữ liệu."),
        ("Thành công", "Điện văn hợp lệ đã được ghi vào kho Inbox."),
        ("Trùng", "Điện văn đã tồn tại; không tạo thêm bản ghi."),
        ("Lỗi", "Không đọc, không phân tích hoặc không ghi được dữ liệu."),
        ("Đã bỏ qua", "Dữ liệu không thuộc phạm vi hoặc được người vận hành bỏ qua."),
    ])

    document.add_heading("5. Yêu cầu giao diện Công cụ phân tích điện văn", level=1)
    add_table(document, ["Khu vực", "Yêu cầu tối thiểu"], [
        ("Cấu hình nguồn", "Đường dẫn/nguồn dữ liệu, mẫu tệp nếu áp dụng và phạm vi thời gian."),
        ("Cấu hình đích", "Thông tin kết nối Oracle hoặc cấu hình kết nối đã mã hóa/được quản lý ngoài mã nguồn."),
        ("Điều khiển", "Kiểm tra cấu hình, Bắt đầu, Dừng an toàn, Quét lại và Xóa log hiển thị."),
        ("Trạng thái", "Đang dừng/đang chạy, nguồn hiện tại, thời điểm bắt đầu và thời điểm cập nhật cuối."),
        ("Thống kê", "Tổng đã đọc, thành công, trùng, lỗi, bỏ qua và số còn chờ."),
        ("Nhật ký", "Thời gian, mức độ, nguồn/tệp, định danh điện văn và mô tả kết quả."),
        ("Tiến độ", "Hiển thị tiến độ tổng thể hoặc số lượng đã xử lý trên tổng số phát hiện."),
    ])

    document.add_heading("6. Yêu cầu chức năng phân tích điện văn", level=1)
    requirements = [
        ("FR-MA-01", "Cho phép người vận hành khai báo hoặc chọn cấu hình nguồn điện văn."),
        ("FR-MA-02", "Cho phép kiểm tra khả năng truy cập nguồn và kết nối cơ sở dữ liệu trước khi chạy."),
        ("FR-MA-03", "Không cho phép bắt đầu khi cấu hình bắt buộc thiếu hoặc kết nối đích không sử dụng được."),
        ("FR-MA-04", "Cho phép chạy một lần theo phạm vi được chọn hoặc chạy định kỳ theo cấu hình triển khai."),
        ("FR-MA-05", "Đọc dữ liệu nguồn theo thứ tự xác định và không bỏ sót dữ liệu mới trong phạm vi."),
        ("FR-MA-06", "Tách được từng điện văn hoàn chỉnh khi một nguồn chứa nhiều điện văn."),
        ("FR-MA-07", "Giữ nguyên nội dung điện văn gốc, bao gồm thứ tự dòng và các dấu hiệu đầu/cuối điện văn."),
        ("FR-MA-08", "Trích xuất tối thiểu: NBR/CID/CSN, thời điểm nhận, loại điện văn, nguồn gửi (FROM_PL/ORIGIN) và CONTENT."),
        ("FR-MA-09", "Chuẩn hóa khoảng trắng và ký tự điều khiển phục vụ phân tích nhưng không làm thay đổi bản nội dung gốc được lưu."),
        ("FR-MA-10", "Kiểm tra các trường bắt buộc và đánh dấu lỗi nếu không xác định được nội dung điện văn hợp lệ."),
        ("FR-MA-11", "Phát hiện điện văn trùng bằng khóa nghiệp vụ ổn định, ưu tiên kết hợp nguồn, số điện văn, thời điểm và nội dung."),
        ("FR-MA-12", "Không tạo bản ghi Inbox thứ hai khi điện văn đã được ghi thành công trước đó."),
        ("FR-MA-13", "Ghi điện văn hiện hành vào T_RECEIVE_LOG; dữ liệu lưu trữ/log file có thể được ghi vào T_RECEIVE_LOGFILE theo cấu hình."),
        ("FR-MA-14", "Mỗi bản ghi thành công phải có ID duy nhất, LETTERNBR_PK, NBR, CONTENT, TYPE và thông tin nguồn gửi nếu có."),
        ("FR-MA-15", "Giao dịch ghi một điện văn phải toàn vẹn; không để lại bản ghi dở dang khi lỗi."),
        ("FR-MA-16", "Sau khi ghi thành công, điện văn phải truy vấn được bằng MESSAGE_PKG.GetInboxBySearch và GetInboxDetailBy."),
        ("FR-MA-17", "Cho phép dừng an toàn; hoàn tất hoặc rollback điện văn đang ghi trước khi dừng."),
        ("FR-MA-18", "Cho phép chạy lại dữ liệu lỗi mà không ghi trùng các điện văn đã thành công."),
        ("FR-MA-19", "Ghi nhật ký riêng cho từng điện văn và tổng hợp kết quả khi kết thúc lượt chạy."),
        ("FR-MA-20", "Khi mất kết nối nguồn hoặc Oracle, công cụ thông báo lỗi, không báo thành công sai và cho phép thử lại."),
    ]
    add_table(document, ["Mã", "Yêu cầu"], requirements)

    document.add_heading("7. Dữ liệu đầu vào và đầu ra", level=1)
    add_table(document, ["Nhóm", "Trường", "Yêu cầu"], [
        ("Nguồn", "Source path/source ID", "Xác định được nguồn của dữ liệu thô."),
        ("Điện văn", "LETTERNBR_PK", "Thời điểm nhận; dùng cho tìm kiếm theo ngày và sắp xếp."),
        ("Điện văn", "NBR", "Số/định danh hiển thị tại danh sách CID/CSN."),
        ("Điện văn", "CONTENT", "Toàn bộ nội dung gốc; không được rỗng."),
        ("Điện văn", "TYPE", "Loại điện văn hoặc giá trị phân loại hệ thống."),
        ("Điện văn", "FROM_PL", "Nguồn gửi hiển thị tại trường Origin."),
        ("Điện văn", "ORIGIN", "Thông tin nguồn bổ sung nếu có."),
        ("Hệ thống", "ID", "Khóa duy nhất dùng để mở chi tiết/Forward."),
        ("Nhật ký", "Status/error", "Kết quả xử lý và mô tả lỗi không chứa bí mật kết nối."),
    ])

    document.add_heading("8. Yêu cầu hiển thị kết quả tại Inbox", level=1)
    add_table(document, ["Mã", "Yêu cầu"], [
        ("FR-IN-01", "Inbox mặc định cho phép tìm kiếm theo Từ ngày, Đến ngày, NBR, Origin và Content."),
        ("FR-IN-02", "Ngày tìm kiếm phải theo dd-MM-yyyy; Từ ngày không được lớn hơn Đến ngày."),
        ("FR-IN-03", "Danh sách hiển thị NBR/CID/CSN, tổng số bản ghi và phân trang 100 bản ghi/trang."),
        ("FR-IN-04", "Kết quả được sắp xếp theo thời điểm nhận giảm dần."),
        ("FR-IN-05", "Khi chọn một bản ghi, Inbox tải chi tiết theo TYPE và ID."),
        ("FR-IN-06", "Chi tiết hiển thị Address trích từ nội dung, Time receive, Origin và toàn bộ Content."),
        ("FR-IN-07", "Cho phép Forward điện văn sang SendMessageFlight với ID bản ghi được chọn."),
        ("FR-IN-08", "Cho phép Forward AMHS điện văn được chọn."),
        ("FR-IN-09", "Cho phép in toàn bộ nội dung điện văn."),
        ("FR-IN-10", "Ngăn gửi đồng thời nhiều yêu cầu tìm kiếm trùng khi yêu cầu trước chưa hoàn tất."),
        ("FR-IN-11", "Điện văn lưu tại T_RECEIVE_LOGFILE phải tra cứu và xem chi tiết được qua thủ tục LogFile tương ứng khi chức năng lưu trữ được sử dụng."),
    ])

    document.add_heading("9. Quy tắc nghiệp vụ", level=1)
    add_table(document, ["Mã", "Quy tắc"], [
        ("BR-MA-01", "Một điện văn nguồn chỉ tạo tối đa một bản ghi Inbox hợp lệ."),
        ("BR-MA-02", "CONTENT lưu trữ phải là bản có thể đối chiếu với nguồn ban đầu."),
        ("BR-MA-03", "Không ghi bản ghi Thành công khi thiếu CONTENT hoặc không xác định được thời điểm nhận theo quy tắc cấu hình."),
        ("BR-MA-04", "Lỗi của một điện văn không được làm mất các điện văn hợp lệ khác trong cùng nguồn."),
        ("BR-MA-05", "Chạy lại phải có tính lặp an toàn: bản ghi thành công không bị nhân đôi."),
        ("BR-MA-06", "Dữ liệu chỉ xuất hiện tại Inbox sau khi giao dịch lưu hoàn tất."),
        ("BR-MA-07", "ID, TYPE dùng trên Inbox phải ánh xạ đúng tới bản ghi chi tiết."),
        ("BR-MA-08", "Từ ngày và Đến ngày của Inbox bao gồm trọn ngày kết thúc."),
    ])

    document.add_heading("10. Yêu cầu phi chức năng", level=1)
    add_table(document, ["Mã", "Yêu cầu"], [
        ("NFR-MA-01", "Cấu hình mật khẩu/chuỗi kết nối không được ghi cứng trong mã nguồn hoặc nhật ký."),
        ("NFR-MA-02", "Chỉ tài khoản được phân quyền mới được vận hành công cụ và truy cập Inbox."),
        ("NFR-MA-03", "Công cụ phải xử lý Unicode và tiếng Việt mà không biến đổi thành ký tự lỗi."),
        ("NFR-MA-04", "Thời gian, múi giờ và định dạng ngày phải nhất quán giữa công cụ, Oracle và Inbox."),
        ("NFR-MA-05", "Nhật ký phải đủ để truy vết nguồn, thời điểm, kết quả và lỗi của từng lượt chạy."),
        ("NFR-MA-06", "Nhật ký không được chứa mật khẩu hoặc dữ liệu bí mật của kết nối."),
        ("NFR-MA-07", "Công cụ không khóa giao diện trong suốt tác vụ dài; trạng thái và tiến độ tiếp tục được cập nhật."),
        ("NFR-MA-08", "Dừng ứng dụng hoặc mất kết nối không được làm hỏng dữ liệu đã ghi thành công."),
        ("NFR-MA-09", "Inbox dùng phân trang để tránh tải toàn bộ dữ liệu trong một yêu cầu."),
        ("NFR-MA-10", "Lỗi API/Oracle phải được hiển thị có kiểm soát và không làm lộ thông tin nhạy cảm cho người dùng cuối."),
    ])

    document.add_heading("11. Tiêu chí nghiệm thu", level=1)
    tests = [
        ("TC-MA-01", "Kiểm tra cấu hình nguồn và Oracle hợp lệ", "Thông báo thành công; cho phép bắt đầu"),
        ("TC-MA-02", "Thiếu cấu hình bắt buộc", "Không chạy; chỉ rõ trường thiếu"),
        ("TC-MA-03", "Nguồn có một điện văn hợp lệ", "Tạo đúng một bản ghi Inbox"),
        ("TC-MA-04", "Nguồn có nhiều điện văn", "Tách và ghi đủ từng điện văn"),
        ("TC-MA-05", "Điện văn Unicode/tiếng Việt", "Nội dung lưu và hiển thị không lỗi ký tự"),
        ("TC-MA-06", "Điện văn thiếu CONTENT hoặc sai cấu trúc", "Đánh dấu lỗi; không tạo bản ghi thành công"),
        ("TC-MA-07", "Chạy lại nguồn đã thành công", "Không tạo bản ghi trùng"),
        ("TC-MA-08", "Một điện văn lỗi trong nguồn nhiều điện văn", "Các điện văn hợp lệ khác vẫn được xử lý"),
        ("TC-MA-09", "Mất kết nối Oracle khi ghi", "Rollback điện văn đang ghi; ghi log lỗi"),
        ("TC-MA-10", "Dừng trong khi đang chạy", "Dừng an toàn; thống kê đúng phần đã xử lý"),
        ("TC-MA-11", "Chạy lại bản ghi lỗi", "Ghi thành công khi lỗi đã được khắc phục, không nhân đôi bản khác"),
        ("TC-IN-01", "Tìm Inbox theo khoảng ngày", "Trả đúng dữ liệu và bao gồm trọn ngày kết thúc"),
        ("TC-IN-02", "Tìm theo NBR", "Trả đúng điện văn có NBR tương ứng"),
        ("TC-IN-03", "Tìm theo Origin hoặc Content", "Trả các bản ghi chứa giá trị tìm kiếm"),
        ("TC-IN-04", "Từ ngày lớn hơn Đến ngày", "Không gửi truy vấn; hiển thị cảnh báo"),
        ("TC-IN-05", "Kết quả trên 100 bản ghi", "Tổng số và nút chuyển trang chính xác"),
        ("TC-IN-06", "Chọn điện văn", "Hiển thị đúng Address, Time receive, Origin và Content"),
        ("TC-IN-07", "Forward/Forward AMHS", "Mở đúng chức năng với ID điện văn đã chọn"),
        ("TC-IN-08", "Print", "Bản in chứa đầy đủ nội dung điện văn"),
        ("TC-IN-09", "Hai lần Search liên tiếp khi request đang chạy", "Yêu cầu trùng thứ hai bị chặn"),
        ("TC-IN-10", "Đối chiếu nguồn, Oracle và Inbox", "Nội dung và định danh nhất quán, truy vết được"),
    ]
    add_table(document, ["Mã kiểm thử", "Nội dung", "Kết quả mong đợi"], tests)

    document.add_heading("12. Ma trận truy vết", level=1)
    add_table(document, ["Mã yêu cầu", "Nội dung", "Kiểm thử"], [
        ("FR-MA-01…FR-MA-05", "Cấu hình và đọc nguồn", "TC-MA-01…TC-MA-04"),
        ("FR-MA-06…FR-MA-10", "Tách, phân tích và kiểm tra", "TC-MA-04…TC-MA-06"),
        ("FR-MA-11…FR-MA-16", "Chống trùng và ghi Inbox", "TC-MA-03, TC-MA-07…TC-MA-09, TC-IN-10"),
        ("FR-MA-17…FR-MA-20", "Dừng, chạy lại và xử lý lỗi", "TC-MA-09…TC-MA-11"),
        ("FR-IN-01…FR-IN-06", "Tra cứu và xem chi tiết Inbox", "TC-IN-01…TC-IN-06"),
        ("FR-IN-07…FR-IN-11", "Forward, Print và dữ liệu lưu trữ", "TC-IN-07…TC-IN-10"),
        ("NFR-MA-01…NFR-MA-10", "Bảo mật, Unicode, độ tin cậy và hiệu năng", "TC-MA-05, TC-MA-09…TC-MA-11, TC-IN-05, TC-IN-10"),
    ])

    document.add_heading("13. Giả định và nội dung cần xác nhận khi triển khai", level=1)
    add_bullets(document, [
        "Loại nguồn thực tế (thư mục, tệp log, API hoặc nguồn khác) được khai báo bằng cấu hình triển khai.",
        "Quy tắc nhận diện đầu/cuối điện văn và khóa chống trùng phải được xác nhận bằng bộ điện văn mẫu thực tế.",
        "Tần suất chạy tự động, thời gian lưu log và cơ chế lưu trữ T_RECEIVE_LOGFILE do đơn vị vận hành phê duyệt.",
        "Tài khoản Oracle của công cụ chỉ được cấp quyền tối thiểu cần thiết để ghi Inbox và nhật ký.",
    ])

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Trang ")
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = "PAGE"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.extend([begin, instruction, end])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)


if __name__ == "__main__":
    build()
    print(OUTPUT)
