from __future__ import annotations

import re
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

from models import SlotRecord


class SlotFormatError(ValueError):
    pass


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("\xa0", " ").strip()


def _slot_time(value: object) -> str:
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return f"{value.hour:02d}:{value.minute:02d}"
    text = _text(value).replace(" ", "")
    match = re.fullmatch(r"(\d{1,2}):?(\d{2})", text)
    if not match:
        raise SlotFormatError(f"Mốc giờ không hợp lệ: {value!r}")
    hour, minute = int(match.group(1)), int(match.group(2))
    if hour > 23 or minute > 59:
        raise SlotFormatError(f"Mốc giờ không hợp lệ: {value!r}")
    return f"{hour:02d}:{minute:02d}"


def flight_date_from_root(folder: Path) -> date:
    for part in (folder.parent.name, folder.name):
        match = re.search(r"NGAY\s*(\d{2})(\d{2})", part, re.IGNORECASE)
        if match:
            return date(date.today().year, int(match.group(2)), int(match.group(1)))
    raise SlotFormatError(f"Không lấy được ngày từ tên thư mục: {folder.parent.name}")


def airport_from_filename(path: Path) -> str:
    match = re.match(r"SLOT_([A-Z0-9]{4})_", path.name, re.IGNORECASE)
    if not match:
        raise SlotFormatError(f"Không lấy được sân bay từ tên file: {path.name}")
    return match.group(1).upper()


def parse_slot_file(path: Path, flight_date: date) -> list[SlotRecord]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    records: list[SlotRecord] = []
    from_airport = airport_from_filename(path)
    try:
        sheet = workbook.active
        header = tuple(_text(value).upper() for value in next(sheet.iter_rows(values_only=True))[:6])
        if header[:5] != ("ETD", "CARRIER", "AC", "SEATS", "TO"):
            raise SlotFormatError(f"{path.name}: header A:E không đúng định dạng SLOT")
        current_time = ""
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            values = tuple(row) + (None,) * max(0, 6 - len(row))
            if _text(values[0]):
                current_time = _slot_time(values[0])
            # Cột B có dữ liệu nghĩa là dòng slot cần import; A trống kế thừa mốc giờ gần nhất phía trên.
            if not _text(values[1]):
                continue
            if not current_time:
                raise SlotFormatError(f"{path.name}, dòng {row_number}: có Carrier nhưng chưa có ETD")
            callsign = _text(values[5]).replace(" ", "").upper()
            record = SlotRecord(
                etd_eta=current_time,
                carrier=_text(values[1]).upper(),
                aircraft=_text(values[2]).upper(),
                seats=_text(values[3]),
                from_airport=from_airport,
                to_airport=_text(values[4]).upper(),
                callsign=callsign,
                flight_date=flight_date,
                aero=callsign[:2],
                source_file=path.name,
                source_row=row_number,
            )
            limits = {
                "ETD_ETA": (record.etd_eta, 6), "CARRIE": (record.carrier, 10),
                "AC": (record.aircraft, 10), "SEATS": (record.seats, 10),
                "FROM_AIRP": (record.from_airport, 4), "TO_AIRP": (record.to_airport, 4),
                "CALLSIGN": (record.callsign, 10), "AERO": (record.aero, 2),
            }
            for field, (value, limit) in limits.items():
                if len(value) > limit:
                    raise SlotFormatError(f"{path.name}, dòng {row_number}: {field} vượt {limit} ký tự ({value})")
            records.append(record)
    finally:
        workbook.close()
    return records


def parse_slot_folder(folder: Path, files: list[Path]) -> tuple[list[SlotRecord], list[str]]:
    flight_date = flight_date_from_root(folder)
    all_records: list[SlotRecord] = []
    details: list[str] = []
    for path in files:
        records = parse_slot_file(path, flight_date)
        all_records.extend(records)
        details.append(f"{path.name}: {len(records)} dòng")
    # Giữ nguyên cả các dòng giống nhau vì chúng có thể là hai slot thực tế khác nhau trong Excel.
    return all_records, details


def supports_slot_folder(folder: Path) -> bool:
    return folder.name.strip().upper() in {"SLOT", "SLOT CHK"}

