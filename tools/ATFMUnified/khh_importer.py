from __future__ import annotations

import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from models import FlightRecord


class ImportFormatError(ValueError):
    pass


def _text(value: object) -> str:
    return "" if value is None else str(value).replace("\xa0", " ").strip()


def _flight_number(value: object, prefix: str = "") -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    result = _text(value).replace(" ", "").upper()
    if prefix and result and not result.startswith(prefix):
        result = prefix + result
    return result


def _date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for fmt in ("%d/%m/%y", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _time(value: object, extra_plus: object = None) -> str:
    plus = _text(extra_plus) == "+"
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        result = f"{value.hour:02d}:{value.minute:02d}"
    else:
        text = _text(value).upper().replace("UTC", "").replace(" ", "")
        plus = plus or text.endswith("+")
        text = text.rstrip("+")
        match = re.fullmatch(r"(\d{1,2}):?(\d{2})(?::\d{2})?", text)
        if not match:
            raise ImportFormatError(f"Thời gian không hợp lệ: {value!r}")
        hour, minute = int(match.group(1)), int(match.group(2))
        if hour > 23 or minute > 59:
            raise ImportFormatError(f"Thời gian không hợp lệ: {value!r}")
        result = f"{hour:02d}:{minute:02d}"
    return result + ("+" if plus else "")


def _record(
    flight_date: date | None,
    callsign: object,
    origin: object,
    destination: object,
    etd: object,
    eta: object,
    operator: str,
    path: Path,
    row_number: int,
    prefix: str = "",
    eta_plus: object = None,
) -> FlightRecord:
    if flight_date is None:
        raise ImportFormatError("Thiếu ngày bay")
    record = FlightRecord(
        flight_date=flight_date,
        callsign=_flight_number(callsign, prefix),
        origin=_text(origin).upper(),
        destination=_text(destination).upper(),
        etd=_time(etd),
        eta=_time(eta, eta_plus),
        operator=operator,
        source_file=path.name,
        source_row=row_number,
    )
    limits = {
        "CallSign": (record.callsign, 10),
        "From": (record.origin, 4),
        "To": (record.destination, 4),
        "ETD": (record.etd, 6),
        "ETA": (record.eta, 6),
        "Oper": (record.operator, 10),
    }
    for field, (value, limit) in limits.items():
        if not value:
            raise ImportFormatError(f"Thiếu {field}")
        if len(value) > limit:
            raise ImportFormatError(f"{field} vượt quá {limit} ký tự: {value}")
    return record


def _target_date_from_path(folder: Path) -> date | None:
    for part in (folder.name, folder.parent.name):
        match = re.search(r"NGAY\s*(\d{2})(\d{2})", part, re.IGNORECASE)
        if match:
            return date(date.today().year, int(match.group(2)), int(match.group(1)))
    return None


def _header_map(row: Iterable[object]) -> dict[str, int]:
    return {_text(value).upper(): index for index, value in enumerate(row) if _text(value)}


def _parse_standard(
    path: Path,
    target_date: date | None,
    operator: str,
    callsign_column: str,
    prefix: str = "",
    remove_hyphen: bool = False,
) -> list[FlightRecord]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    records: list[FlightRecord] = []
    try:
        sheet = workbook.active
        header: dict[str, int] | None = None
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            values = tuple(row)
            candidate = _header_map(values)
            if "DATE" in candidate and "FLT" in candidate and "DEP" in candidate and "ARR" in candidate:
                header = candidate
                continue
            if not header:
                continue
            first = _text(values[0] if values else None).upper()
            if not first or first == "TOTALS":
                continue
            flight_date = _date(values[header["DATE"]])
            if flight_date is None or (target_date and flight_date != target_date):
                continue
            std_key = "ETD" if "ETD" in header and _text(values[header["ETD"]]) else "STD"
            sta_key = "ETA" if "ETA" in header and _text(values[header["ETA"]]) else "STA"
            callsign = values[header[callsign_column]]
            if remove_hyphen:
                callsign = _text(callsign).replace("-", "")
            records.append(
                _record(
                    flight_date,
                    callsign,
                    values[header["DEP"]],
                    values[header["ARR"]],
                    values[header[std_key]],
                    values[header[sta_key]],
                    operator,
                    path,
                    row_number,
                    prefix,
                )
            )
    finally:
        workbook.close()
    return records


def _parse_vna(path: Path, target_date: date | None) -> list[FlightRecord]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    records: list[FlightRecord] = []
    try:
        sheet = workbook.active
        header = _header_map(next(sheet.iter_rows(values_only=True)))
        required = {"FLT. NO", "DEP", "ARR", "EDD", "ETD", "ETA"}
        if not required.issubset(header):
            raise ImportFormatError("Không nhận diện được định dạng VNA")
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            flight_date = _date(row[header["EDD"]])
            if not _text(row[header["FLT. NO"]]) or (target_date and flight_date != target_date):
                continue
            records.append(_record(flight_date, row[header["FLT. NO"]], row[header["DEP"]], row[header["ARR"]], row[header["ETD"]], row[header["ETA"]], "VNA", path, row_number))
    finally:
        workbook.close()
    return records


def _parse_pic(path: Path, target_date: date | None) -> list[FlightRecord]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    records: list[FlightRecord] = []
    try:
        candidates: list[tuple[int, tuple[object, ...], date | None]] = []
        for row_number, row in enumerate(workbook.active.iter_rows(min_row=2, values_only=True), 2):
            if not _text(row[0]).upper().startswith("PIC"):
                continue
            flight_date = _date(row[1])
            candidates.append((row_number, tuple(row), flight_date))
        matching = [item for item in candidates if target_date is None or item[2] == target_date]
        if not matching and target_date and candidates:
            source_dates = {item[2] for item in candidates}
            # Một số file PIC bị nhập nhầm tháng cho toàn bộ cột B. Chỉ dùng ngày thư mục
            # khi file đồng nhất và ngày/năm vẫn khớp; sai ngày hoặc lẫn ngày vẫn bị từ chối.
            if len(source_dates) == 1:
                source_date = next(iter(source_dates))
                if source_date and source_date.day == target_date.day and source_date.year == target_date.year:
                    matching = [(row_number, row, target_date) for row_number, row, _source_date in candidates]
        for row_number, row, flight_date in matching:
            records.append(_record(flight_date, row[0], row[4], row[6], row[5], row[7], "PIC", path, row_number))
    finally:
        workbook.close()
    return records


def _parse_vfc(path: Path, target_date: date | None) -> list[FlightRecord]:
    if target_date is None:
        raise ImportFormatError("Không xác định được ngày của thư mục để đọc VFC")
    workbook = load_workbook(path, read_only=True, data_only=True)
    records: list[FlightRecord] = []
    try:
        expected_sheet = f"{target_date.day:02d}"
        matching_sheets = [sheet for sheet in workbook.worksheets if sheet.title.strip().zfill(2) == expected_sheet]
        if not matching_sheets:
            raise ImportFormatError(f"Không tìm thấy sheet {expected_sheet} trong file VFC")
        for sheet in matching_sheets:
            current_date: date | None = None
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                marker = _date(row[0] if row else None)
                if marker:
                    current_date = marker
                    continue
                if current_date != target_date or not row or not _text(row[0]) or _text(row[0]).upper() == "FLT. NO":
                    continue
                # Một số sheet có thêm cột đăng ký tàu bay trước số hiệu chuyến.
                offset = 1 if len(row) >= 6 and "-" in _text(row[2]) else 0
                eta_plus = row[5 + offset] if len(row) > 5 + offset else None
                records.append(_record(current_date, row[offset], row[1 + offset], row[2 + offset], row[3 + offset], row[4 + offset], "VFC", path, row_number, prefix="VFC", eta_plus=eta_plus))
    finally:
        workbook.close()
    return records


def parse_khh_file(path: Path, target_date: date | None) -> list[FlightRecord]:
    name = path.name.upper()
    operator = name[:3]
    if name.startswith("BAV"):
        return _parse_standard(path, target_date, operator, "FLT", operator)
    if name.startswith("VAG"):
        return _parse_standard(path, target_date, operator, "REG", remove_hyphen=True)
    if name.startswith("VJC"):
        return _parse_standard(path, target_date, operator, "FLT", operator)
    if name.startswith("VNA"):
        return _parse_vna(path, target_date)
    if name.startswith("PIC"):
        return _parse_pic(path, target_date)
    if name.startswith("VFC"):
        return _parse_vfc(path, target_date)
    raise ImportFormatError(f"Chưa hỗ trợ định dạng file: {path.name}")


def parse_khh_folder(folder: Path, files: list[Path]) -> tuple[list[FlightRecord], list[str]]:
    target_date = _target_date_from_path(folder)
    all_records: list[FlightRecord] = []
    details: list[str] = []
    for path in files:
        records = parse_khh_file(path, target_date)
        if not records:
            raise ImportFormatError(f"{path.name}: không tìm thấy chuyến bay phù hợp")
        all_records.extend(records)
        details.append(f"{path.name}: {len(records)} dòng")
    unique: dict[tuple[object, ...], FlightRecord] = {}
    for record in all_records:
        unique.setdefault(record.key, record)
    if len(unique) != len(all_records):
        details.append(f"Loại {len(all_records) - len(unique)} dòng trùng trong các file")
    return list(unique.values()), details


def supports_folder(folder: Path) -> bool:
    return folder.name.strip().upper() == "KHH"

