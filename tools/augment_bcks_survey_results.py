from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml.ns import qn


BASE = Path(__file__).resolve().parents[1]
SOURCE = BASE / "TaiLieu" / "HoSoTrienKhai_HTQLSLB" / "01_2026_BCKS-HTQLSLB_v1.0.docx"
OUTPUT = BASE / "TaiLieu" / "HoSoTrienKhai_HTQLSLB" / "01_2026_BCKS-HTQLSLB_v1.1_ChiTiet.docx"


GROUPS = [
    ("I", "FR-INT", "PHÂN HỆ NGHIỆP VỤ TÍCH HỢP VÀ TỰ ĐỘNG HÓA DỮ LIỆU", [
        "Tự động cập nhật dữ liệu từ email và thư mục lưu trữ",
        "Thu thập và xử lý dữ liệu ADS-B cho khai thác O/F",
        "Tích hợp lọc SLOT vào hệ thống SLB",
        "Kết nối trực tiếp AMHS để gửi/nhận điện văn",
        "Chuẩn hóa API chia sẻ dữ liệu",
    ]),
    ("II", "FR-ALT", "PHÂN HỆ NGHIỆP VỤ CẢNH BÁO VÀ TIỆN ÍCH HỖ TRỢ", [
        "Hệ thống cảnh báo thông minh đa kịch bản",
        "Quản lý thông tin KHB quân sự và sử dụng vùng trời",
    ]),
    ("III", "FR-OPT", "PHÂN HỆ NGHIỆP VỤ NÂNG CẤP VÀ TỐI ƯU CÁC CHỨC NĂNG CỦA HTSLB", [
        "Mở rộng thời gian lưu trữ điện văn INBOX",
        "Cảnh báo NOPERM cho điện văn không có trong KHBN",
        "Áp dụng quy tắc 120h ICAO cho hiển thị FPL",
        "Tìm kiếm chuyến bay đa trường, hỗ trợ tìm kiếm theo từ khóa",
        "Tìm kiếm nâng cao phục vụ báo cáo",
        "Phân loại chuyến bay Quốc nội/Quốc tế tự động",
        "Export dữ liệu chia tách 3 loại file (ALL, LD, OF)",
        "Chức năng Edit info chuyến bay",
        "Hiển thị lịch sử thao tác của người xử lý chuyến bay",
        "Đồng bộ dữ liệu điện văn liên ngày (DEP cuối ngày → ARR hôm sau)",
        "Thống kê theo khung giờ/chặng/đường bay",
        "Cảnh báo chuyến bay không có trong phép bay/KHB ngày",
        "Hỗ trợ hủy phép bay quốc tế/quốc nội",
        "Tối ưu hiệu năng ExportBravo (< 3 phút)",
        "Chức năng F8 ‘ép dòng’ trong Calendar Accepted",
        "Tối ưu Gen KHB ngày hôm sau ra AFTN (< 3 phút)",
        "Mở rộng phạm vi tra cứu INBOX (>7 ngày)",
        "Tự động cập nhật khi sửa KHB đã build",
        "Cảnh báo chuyến bay cấp sai ngày bay (Day) so với ngày bay thực tế",
        "Bổ sung trường ‘Mục đích chuyến bay’ trong KHB ngày",
        "Logic cảnh báo ‘đỏ’ chính xác cho chuyến bay hết hiệu lực",
        "Tự động cập nhật đường bay theo đoạn từ FPL thực tế",
        "Tự động nhận diện và chuyển ngày cho chuyến bay quốc tế hạ cánh hôm sau",
        "Theo dõi kế hoạch bay hàng ngày với cập nhật liên tục",
        "Thông báo thời điểm nhận số liệu bay Đi/Đến từ Trung tâm TBHĐB",
        "Cập nhật số liệu bay từ HTSLB sang Bravo",
        "Khắc phục lỗi nhầm ngày so với thực tế",
        "Sửa nhiều chuyến bay cùng nội dung trong một thao tác",
        "Đối chiếu số liệu bay giữa HTSLB và Bravo",
        "Mở rộng tìm kiếm trên giao diện Daily Military Report",
        "Thêm lựa chọn theo tất cả tiêu chí thống kê trong báo cáo",
        "Đánh giá so sánh việc số liệu của KHBHĐBN",
    ]),
    ("IV", "FR-RPT", "PHÂN HỆ NGHIỆP VỤ PHÂN TÍCH, BÁO CÁO THÔNG MINH", [
        "Biểu đồ thông tin tổng quan khai thác bay",
        "Phân tích xu hướng khai thác theo thời gian",
        "Phát hiện và cảnh báo dữ liệu bất thường",
        "Tích hợp ADS-B cho báo cáo khai thác thực tế",
        "Tổng hợp chỉ số hiệu suất bay từ ADS-B O/F",
        "Biểu đồ thống kê trạng thái chuyến bay theo tỷ lệ phần trăm",
        "Biểu đồ so sánh hoạt động bay giữa các sân bay",
        "Báo cáo chuyến bay quân sự theo tiêu chí mở rộng",
        "Báo cáo tổng hợp hoạt động bay tại tất cả sân bay dân dụng",
        "Báo cáo cất, hạ cánh tại các sân bay toàn quốc",
    ]),
    ("V", "FR-AI", "PHÂN HỆ ỨNG DỤNG AI HỖ TRỢ THỐNG KÊ, TÌM KIẾM VÀ TỔNG HỢP SỐ LIỆU TỰ ĐỘNG", [
        "Xử lý ngôn ngữ tự nhiên cho truy vấn hàng không",
        "Truy vấn và tổng hợp dữ liệu tự động theo yêu cầu",
        "Tích hợp Trợ lý ảo (Chatbot) vào giao diện HTSLB",
        "Báo cáo và giám sát hiệu năng của Trợ lý ảo",
    ]),
]


GROUP_INTRO = {
    "I": "Khảo sát các nguồn dữ liệu, adapter/worker, parser, mapping, chống trùng, retry/quarantine, ghi Oracle và đối soát source–target.",
    "II": "Khảo sát cơ chế phát hiện/phân phối cảnh báo runtime và quy trình nhập–duyệt–tạo điện văn–phát đi đối với Live Fire, Daily Statistic và KHB quân sự.",
    "III": "Khảo sát các thay đổi trên màn hình hiện hữu, package Oracle, batch nền, chỉ tiêu hiệu năng, tra cứu, cảnh báo, ngày nghiệp vụ, đồng bộ và đối soát Bravo.",
    "IV": "Khảo sát nguồn dữ liệu, công thức KPI, bộ lọc, biểu đồ, drill-down, snapshot, export, phân quyền và quy tắc xác định dữ liệu bất thường.",
    "V": "Khảo sát ý định/câu hỏi, phạm vi dữ liệu, SQL Guard, Oracle read-only, phiên Chatbot, feedback, giám sát chất lượng và hiệu năng AI.",
}


def configure(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in [("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11.5)]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True


def add_heading(doc, text, level):
    return doc.add_heading(text, level=level)


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    return p


def add_bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = value
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    doc.add_paragraph()
    return table


def code(group, index):
    if group == "I":
        return f"FR-INT-{index:03d}"
    if group == "II":
        return f"FR-ALT-{index:03d}"
    if group == "III":
        return f"FR-OPT-{index:03d}"
    if group == "IV":
        return f"FR-RPT-{index:03d}"
    return f"FR-AI-{index:03d}"


def detail_for(group, index, title):
    if group == "I":
        custom = {
            1: ("Email/thư mục/AeroSync", "MIME, CSV, XML, JSON; attachment, Message-Id, ContentHash", "Parser → validate → normalize → dedupe → Oracle; lỗi vào quarantine", "T_RECEIVE/Inbox và log batch; cần bổ sung mẫu file, lịch nhận và mapping trường"),
            2: ("ADS-B API/PostgreSQL tracks", "TrackId, callsign, ICAO, O/F timestamp, tọa độ, quality", "Watermark incremental, loại duplicate/late, match với KHB/FPL", "T_TRACKS_LOG và snapshot actual; cần xác nhận ngưỡng match rate và timezone"),
            3: ("Excel/API SLOT và KHH", "Template, header, flight/date/slot, dòng lỗi", "Staging → validate từng dòng → compare SLOT–KHB–phép bay", "T_KHH/T_SLOT và reconciliation; cần chốt template/mapping phiên bản"),
            4: ("AMHS/AFTN connector", "Header, address, priority, body, MessageId, ACK", "Parse → Inbox/Outbox → delivery state; retry lỗi tạm thời", "MESSAGE_PKG/MESSAGE_FLIGHT_PKG; cần xác nhận ACK timeout và địa chỉ test"),
            5: ("REST/API Gateway", "JSON schema, correlation header, OAuth/JWT/API key, idempotency", "Validate, rate-limit, timeout, chuẩn hóa 400/401/403/409/429/5xx", "API audit/route registry; cần chốt OpenAPI và chính sách backward compatibility"),
        }
        source, data, rule, output = custom[index]
    elif group == "II":
        if index == 1:
            source, data, rule, output = ("ChartReport/header/Notification.ashx", "RuleCode, severity, source, user/role/unit scope, read state", "NOPERM, hết hạn, sai ngày, batch lỗi, bất thường KPI; polling/event, dedupe, acknowledge", "Header badge/popup/email/audit; cần xác nhận ngưỡng và danh sách người nhận")
        else:
            source, data, rule, output = ("LiveFire, Daily Statistic, Military KHB/Report", "Khu vực/thời gian vùng trời; flight/date/purpose; Accepted; QS Message", "Draft → submit → approve/reject → export/dispatch; chỉ Accepted mới vào report/message", "T_PLAN_MESSAGE, Military Report, delivery/audit; cần xác nhận vai trò R_Add/R_Edit/R_Pub")
    elif group == "III":
        source = "Màn hình ATFM_WEB, package Oracle và batch liên quan"
        data = "Trường tìm kiếm/chuyến bay, điện văn, phép bay, KHB, ngày nghiệp vụ và trạng thái xử lý"
        rule = "Kiểm tra dữ liệu đầu vào, quyền, version, ngày/giờ, idempotency; ghi audit và cảnh báo khi có ngoại lệ"
        output = "Danh sách đã lọc/cập nhật, file export, cảnh báo hoặc kết quả đối soát; cần đo hiệu năng và lưu execution log"
        low = title.lower()
        if "exportbravo" in low or "gen khb" in low or "bravo" in low:
            rule = "Batch theo chunk, giới hạn thời gian mục tiêu dưới 180 giây, retry lỗi tạm thời, dừng khi conflict và đối soát số dòng nguồn–đích"
        elif "tìm kiếm" in low or "tra cứu" in low:
            rule = "Bộ lọc nhiều trường, từ khóa, phân trang, data scope và export toàn bộ kết quả; cần kiểm tra execution plan/index"
        elif "cảnh báo" in low or "sai ngày" in low or "hết hiệu lực" in low:
            rule = "Chuẩn hóa BusinessDate/timezone, đối chiếu phép bay–KHB–điện văn và phát cảnh báo đúng severity, không kết luận khi thiếu nguồn"
        elif "sửa" in low or "edit" in low:
            rule = "Sửa trong phạm vi quyền, optimistic concurrency, before/after audit và tính lại các liên kết/cảnh báo downstream"
    elif group == "IV":
        source = "Trang ReportNew/Common, API báo cáo, Oracle snapshot và dữ liệu ADS-B/finished flight"
        data = "Khoảng ngày, sân bay, loại chuyến, trạng thái, đơn vị, mục đích, KPI và timestamp nguồn"
        rule = "Tạo snapshot, tính KPI theo công thức/version, kiểm tra dữ liệu thiếu/trễ, drill-down phải khớp tổng hợp"
        output = "Bảng/biểu đồ/heatmap, drill-down và Excel; cần xác nhận công thức KPI, bộ lọc và quyền export"
    else:
        source = "AI Chat/API, schema metadata và Oracle read-only"
        data = "Câu hỏi tiếng Việt, intent/entity, scope, schema version, SQL draft, result/snapshot và feedback"
        rule = "Clarify khi mơ hồ; SQL Guard allowlist, bind parameter, chặn DML/DDL, giới hạn dòng/thời gian và mask dữ liệu nhạy cảm"
        output = "Bảng/biểu đồ/tóm tắt có nguồn, công thức, model/prompt version; audit latency, token/cost, lỗi và chất lượng"
    return source, data, rule, output


def make_details():
    doc = Document()
    configure(doc)
    add_heading(doc, "4. Kết quả khảo sát chi tiết theo nhóm và chức năng", 1)
    add_para(doc, "Phần này mở rộng bảng tổng hợp 05 nhóm/53 chức năng trong báo cáo khảo sát. Mỗi phiếu chức năng ghi nhận phạm vi, nguồn khảo sát, dữ liệu, quy tắc xử lý, đầu ra, vấn đề cần làm rõ và bằng chứng dự kiến. Nội dung là cơ sở chuyển tiếp sang SRS, SDD, backlog triển khai và kế hoạch kiểm thử.")
    add_table(doc, ["Nhóm", "Mã chức năng", "Số lượng", "Trọng tâm khảo sát", "Kết quả"], [
        ("I", "FR-INT-001…005", "5", "Nguồn, parser, mapping, retry, quarantine, đối soát", "Hoàn thành khảo sát"),
        ("II", "FR-ALT-001…002", "2", "Cảnh báo runtime, Live Fire, Daily, KHB quân sự", "Hoàn thành khảo sát"),
        ("III", "FR-OPT-001…032", "32", "Tra cứu, cảnh báo, KHB, export, đồng bộ, Bravo, hiệu năng", "Hoàn thành khảo sát"),
        ("IV", "FR-RPT-001…010", "10", "KPI, biểu đồ, snapshot, drill-down, export", "Hoàn thành khảo sát"),
        ("V", "FR-AI-001…004", "4", "NL2SQL, Chatbot, read-only, giám sát AI", "Hoàn thành khảo sát"),
        ("Tổng", "53 chức năng", "53", "Truy vết hiện trạng → yêu cầu → thiết kế → test", "Hoàn thành 100% theo phạm vi khảo sát"),
    ])
    for group, prefix, name, titles in GROUPS:
        add_heading(doc, f"4.{GROUPS.index((group, prefix, name, titles)) + 2}. {name}", 2)
        add_para(doc, f"Mã nhóm: {prefix}; số lượng: {len(titles)} chức năng. {GROUP_INTRO[group]}")
        summary_rows = []
        for index, title in enumerate(titles, 1):
            source, data, rule, output = detail_for(group, index, title)
            summary_rows.append((code(group, index), title, source, "Hoàn thành khảo sát; cần xác nhận các trường đánh dấu trong phiếu"))
        add_table(doc, ["Mã", "Chức năng", "Nguồn/màn hình khảo sát", "Trạng thái/kết luận"], summary_rows)
        for index, title in enumerate(titles, 1):
            source, data, rule, output = detail_for(group, index, title)
            add_heading(doc, f"4.{GROUPS.index((group, prefix, name, titles)) + 2}.{index}. {code(group, index)} – {title}", 3)
            add_table(doc, ["Trường khảo sát", "Kết quả ghi nhận"], [
                ("Mã và tên chức năng", f"{code(group, index)} – {title}"),
                ("Mục tiêu/phạm vi", f"Khảo sát hiện trạng và nhu cầu nâng cấp đối với chức năng: {title}. {GROUP_INTRO[group]}"),
                ("Nguồn/màn hình/hệ thống", source),
                ("Dữ liệu đầu vào", data),
                ("Quy tắc xử lý cần xác nhận", rule),
                ("Đầu ra/tích hợp", output),
                ("Phân quyền và lưu vết", "Kiểm tra quyền menu/thao tác ở cả giao diện và máy chủ; ghi actor, timestamp, before/after, correlation/batch và lỗi nghiệp vụ."),
                ("Vấn đề/rủi ro khảo sát", "Cần đối chiếu dữ liệu mẫu, tên trường/package, ngưỡng nghiệp vụ và quyền khai thác với đơn vị sử dụng trước khi khóa thiết kế."),
                ("Kiến nghị", "Lập mapping/contract riêng, bổ sung test data biên, chốt tiêu chí chấp nhận và đưa các điểm chưa xác nhận vào danh sách yêu cầu làm rõ."),
                ("Bằng chứng cần lưu", f"Ảnh màn hình hoặc sơ đồ luồng; mẫu dữ liệu/log/package; biên bản phỏng vấn; mã bằng chứng: KS-{group}-{index:03d}."),
                ("Kết luận khảo sát", "Đủ cơ sở chuyển sang phân tích yêu cầu và thiết kế sơ bộ; các trường chưa xác nhận phải được đóng trước khi phê duyệt SRS/SDD."),
            ])
        if group != "V":
            doc.add_page_break()
    add_heading(doc, "4.7. Tổng hợp dữ liệu, tích hợp và phụ thuộc liên nhóm", 2)
    add_table(doc, ["Chủ đề", "Kết quả khảo sát", "Tác động/kiến nghị"], [
        ("Ngày nghiệp vụ và timezone", "Nhiều chức năng liên quan FPL, KHB, DEP/ARR, ADS-B và báo cáo.", "Dùng BusinessDateService chung; bắt buộc có timezone, source timestamp và cờ uncertain khi thiếu dữ liệu."),
        ("Đối chiếu và idempotency", "Email/File, AMHS, ADS-B, Bravo, batch export và cảnh báo đều có nguy cơ duplicate/replay.", "Dùng SourceId/MessageId/ContentHash/BatchId; lưu raw và audit, có quarantine/replay."),
        ("Phân quyền", "Các chức năng nhập, sửa, duyệt, phát điện văn, báo cáo và AI có phạm vi user/role/đơn vị khác nhau.", "Kiểm tra RBAC/data scope tại API/package; không tin CREATED_BY hoặc scope từ payload."),
        ("Hiệu năng", "ExportBravo, Gen KHB, Inbox, report và AI có truy vấn/batch lớn.", "Đo P95/P99, execution plan, index, chunk, timeout, circuit breaker; lưu metric theo batch."),
        ("Bằng chứng khảo sát", "Ảnh màn hình, mẫu dữ liệu, log, package/API response và biên bản phỏng vấn chưa đồng nhất mã hóa.", "Dùng mã KS-I-001…KS-V-004; lưu evidence theo chức năng và liên kết sang test case."),
    ])
    add_heading(doc, "4.8. Danh sách vấn đề và yêu cầu làm rõ", 2)
    add_table(doc, ["Mã", "Nội dung cần làm rõ", "Đơn vị xác nhận", "Hạn", "Trạng thái"], [
        ("CL-001", "Danh mục endpoint, port, credential và môi trường tích hợp chính thức", "Chủ đầu tư/Đơn vị quản lý", "…………", "Mở/Đóng"),
        ("CL-002", "Mapping trường và phiên bản schema của Email/File, ADS-B, SLOT, AMHS, Bravo", "Đơn vị cung cấp nguồn", "…………", "Mở/Đóng"),
        ("CL-003", "Công thức KPI, ngưỡng cảnh báo và quy tắc ngày nghiệp vụ/timezone", "Đơn vị nghiệp vụ", "…………", "Mở/Đóng"),
        ("CL-004", "Danh sách role/data scope, người duyệt và người phát điện văn", "Chủ đầu tư/ATTT", "…………", "Mở/Đóng"),
        ("CL-005", "Bộ dữ liệu biên, dữ liệu kiểm thử và tiêu chí nghiệm thu từng chức năng", "Nghiệp vụ/QA", "…………", "Mở/Đóng"),
    ])
    add_heading(doc, "4.9. Ma trận truy vết kết quả khảo sát", 2)
    add_table(doc, ["Đối tượng", "Quy ước mã", "Liên kết bắt buộc"], [
        ("Phiếu khảo sát chức năng", "KS-I-001…KS-V-004", "FR code, URL/package, bằng chứng và người xác nhận"),
        ("Yêu cầu làm rõ", "CL-001…", "Chức năng bị ảnh hưởng, người phụ trách, hạn và trạng thái"),
        ("Thiết kế", "SRS/SDD/LLD", "Mã FR, component, schema, API/package/index"),
        ("Kiểm thử", "TC-INT/ALT/OPT/RPT/AI", "Dữ liệu test, expected result, actual result và evidence"),
        ("Nghiệm thu", "BBKS/BCKS/BBNT", "Bản ghi, ảnh/log, người ký và ngày xác nhận"),
    ])
    return doc


def insert_before(anchor, source_doc):
    for child in list(source_doc.element.body):
        if child.tag.endswith("sectPr"):
            continue
        anchor.addprevious(deepcopy(child))


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    base = Document(SOURCE)
    detail = make_details()
    anchor = next((p._p for p in base.paragraphs if p.text.strip().startswith("5. Kết quả khảo sát dữ liệu")), None)
    if anchor is None:
        raise RuntimeError("Không tìm thấy điểm chèn trước mục 5")
    insert_before(anchor, detail)
    for p in base.paragraphs:
        if p.text.strip().startswith("7. Kết luận"):
            next_p = p._p.getnext()
            _ = next_p
            p.text = "7. Kết luận"
        if p.text.strip().startswith("Công tác khảo sát hoàn thành 100%"):
            p.text = ("Công tác khảo sát hoàn thành 100% ngày 13/07/2026, xác định đầy đủ phạm vi 05 nhóm/53 chức năng. "
                      "Báo cáo đã bổ sung kết quả khảo sát chi tiết theo từng chức năng, bao gồm hiện trạng, nguồn/màn hình, "
                      "dữ liệu, quy tắc, đầu ra, tích hợp, phân quyền, vấn đề, kiến nghị và mã bằng chứng. Kết quả đủ điều kiện "
                      "chuyển sang phân tích yêu cầu, thiết kế kỹ thuật, lập kế hoạch kiểm thử và theo dõi các yêu cầu làm rõ còn mở.")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    base.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
