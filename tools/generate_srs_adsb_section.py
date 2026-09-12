#!/usr/bin/env python3
"""Create a standalone ADS-B SRS using the current AeroSync SRS as template."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "TableGrid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
        for run in table.rows[0].cells[index].paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    return table


def add_picture(document: Document, image_path: Path, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Cm(16.2))
    cap = document.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def clear_document_body(document: Document) -> None:
    """Remove template content while preserving styles, section, headers and footer."""
    body = document._element.body
    for child in list(body):
        if not child.tag.endswith("}sectPr"):
            body.remove(child)


def add_cover(document: Document) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(130)
    run = title.add_run("TÀI LIỆU ĐẶC TẢ YÊU CẦU PHẦN MỀM")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(20)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("PHÂN HỆ BÁO CÁO VÀ GIÁM SÁT ADS-B")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)

    info = document.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_before = Pt(48)
    info.add_run(
        "Phạm vi: Báo cáo hiệu suất, bản đồ theo dõi chuyến bay ADS-B\n"
        "Mô-đun: Web, Flight Tracking API và TracksSync/T_TRACKS_LOG\n"
        "Phiên bản tài liệu: 1.0\n"
        "Ngày cập nhật: 12/08/2026"
    )
    document.add_page_break()


def build(input_path: Path, output_path: Path, report_image: Path, map_image: Path) -> None:
    document = Document(str(input_path))
    clear_document_body(document)
    add_cover(document)

    document.add_heading("1. Thông tin chung", level=1)
    add_table(
        document,
        ["Nội dung", "Giá trị"],
        [
            ["Tên tài liệu", "Đặc tả yêu cầu phân hệ Báo cáo và giám sát ADS-B"],
            ["Mã tài liệu", "SRS-ADSB-001"],
            ["Hệ thống", "VATM ATFM Web / AeroSync"],
            ["Phiên bản", "1.0"],
            ["Ngày cập nhật", "12/08/2026"],
            ["Phạm vi", "AdsBPerformanceReport, FlightTrackingMap, Flight Tracking API và TracksSync/T_TRACKS_LOG"],
        ],
        [5.0, 11.2],
    )
    document.add_paragraph(
        "Tài liệu mô tả nhóm chức năng khai thác dữ liệu giám sát ADS-B của hệ thống VATM, "
        "bao gồm báo cáo hiệu suất lịch sử, bản đồ theo dõi gần thời gian thực, dịch vụ Flight Tracking API "
        "và công cụ đồng bộ dữ liệu vào bảng Oracle T_TRACKS_LOG."
    )

    document.add_heading("2. Phạm vi và kiến trúc tổng thể", level=1)
    document.add_paragraph(
        "Giải pháp tách việc truy cập PostgreSQL ADS-B khỏi IIS. Flight Tracking API chỉ đọc dữ liệu vị trí "
        "để phục vụ bản đồ; TracksSync thực hiện phân vùng FIR, đối chiếu kế hoạch bay và ghi lịch sử vào Oracle. "
        "Báo cáo hiệu suất sử dụng dữ liệu lịch sử T_TRACKS_LOG, trong khi bản đồ dùng dữ liệu vị trí mới nhất từ API."
    )
    add_table(
        document,
        ["Thành phần", "Vai trò", "Nguồn/đích dữ liệu"],
        [
            ["AdsBPerformanceReport.aspx", "Tổng hợp KPI, xu hướng và chi tiết ADS-B theo thời gian", "API AdsBPerformance và Oracle T_TRACKS_LOG"],
            ["FlightTrackingMap.aspx", "Hiển thị vị trí chuyến bay trong FIR, tự cập nhật 5 giây", "Flight Tracking API, Oracle T_DAY_FLIGHTS_GOINGON, GeoJSON offline"],
            ["Flight Tracking API", "Cung cấp vị trí mới nhất theo callsign qua HTTP JSON", "Đọc PostgreSQL public.tracks"],
            ["TracksSync", "Phân vùng FIR, làm giàu dữ liệu và lưu lịch sử", "PostgreSQL public.tracks → Oracle T_TRACKS_LOG"],
        ],
        [4.0, 6.0, 6.2],
    )
    document.add_paragraph(
        "Luồng dữ liệu chính: PostgreSQL public.tracks → Flight Tracking API → FlightTrackingMap.aspx; đồng thời "
        "PostgreSQL public.tracks → TracksSync → Oracle T_TRACKS_LOG → AdsBPerformanceReport.aspx."
    )

    document.add_heading("3. Báo cáo hiệu suất ADS-B (AdsBPerformanceReport.aspx)", level=1)
    document.add_paragraph(
        "Trang báo cáo cho phép người dùng nghiệp vụ ATFM theo dõi lưu lượng LD, O/F và các chuyến chưa xác định "
        "đầy đủ thông tin dựa trên dữ liệu TRACKS_LOG. Người dùng phải đăng nhập; quyền menu được kiểm tra khi URL có Menu_ID."
    )
    add_picture(document, report_image, "Hình 1. Giao diện báo cáo tổng hợp chỉ số hiệu suất bay từ ADS-B")
    add_table(
        document,
        ["Nhóm", "Yêu cầu chức năng"],
        [
            ["Bộ lọc", "Chọn từ ngày, đến ngày, PERMTYPE và hãng khai thác; mặc định từ đầu tháng đến ngày hiện tại; ngày kết thúc được tính bao gồm toàn ngày."],
            ["KPI", "Hiển thị tổng bản ghi, số và tỷ lệ LD, số và tỷ lệ O/F, chuyến bay khác và số hãng khai thác."],
            ["Xu hướng", "Vẽ biểu đồ theo ngày cho ba nhóm LD, O/F và OTHER; tự co giãn theo vùng hiển thị."],
            ["Chi tiết", "Hiển thị bảng callsign, hãng, phép bay, sân bay đi/đến, ETD/ETA, FIR, ngày và thời gian cập nhật."],
            ["Phân trang", "Cho phép 50, 100, 200 hoặc 500 dòng trên mỗi trang."],
            ["Cuối ngày", "Mở báo cáo cuối ngày trong cửa sổ modal, bổ sung TIME_IN và TIME_OUT; đóng bằng nút đóng, Escape hoặc bấm vùng nền."],
        ],
        [4.0, 12.2],
    )
    document.add_paragraph("Luồng xử lý:")
    add_bullets(
        document,
        [
            "Khởi tạo khoảng ngày mặc định và tải danh sách hãng theo khoảng ngày/PERMTYPE.",
            "Gọi API AdsBPerformance để nhận dữ liệu tổng hợp, sau đó cập nhật KPI, biểu đồ và bảng chi tiết.",
            "Khi người dùng áp dụng bộ lọc, hệ thống tải lại danh sách hãng và dữ liệu báo cáo.",
            "Báo cáo cuối ngày gọi Page Method của ASP.NET và truy vấn Oracle bằng bind parameter.",
            "Dữ liệu động được mã hóa trước khi hiển thị; lỗi HTTP, mã nghiệp vụ khác 00 và dữ liệu rỗng đều có thông báo rõ ràng.",
        ],
    )
    document.add_paragraph(
        "Quy tắc phân loại OTHER: PERMTYPE là OTHER, thiếu sân bay đi, thiếu sân bay đến hoặc không đối chiếu được hãng khai thác. "
        "STATUS=1 hiển thị VVHN, STATUS=2 hiển thị VVHM; giá trị khác hiển thị Không xác định."
    )

    document.add_heading("4. Bản đồ theo dõi chuyến bay ADS-B (FlightTrackingMap.aspx)", level=1)
    document.add_paragraph(
        "Trang bản đồ hiển thị gần thời gian thực vị trí chuyến bay trong FIR Hà Nội và FIR Hồ Chí Minh. "
        "Bản đồ sử dụng dữ liệu nền GeoJSON offline, không phụ thuộc dịch vụ bản đồ Internet."
    )
    add_picture(document, map_image, "Hình 2. Giao diện bản đồ theo dõi chuyến bay ADS-B")
    add_table(
        document,
        ["Nhóm", "Yêu cầu chức năng"],
        [
            ["Cập nhật", "Tải dữ liệu ngay khi khởi tạo, tự cập nhật mỗi 5 giây và cho phép cập nhật thủ công; không chạy chồng các lượt cập nhật."],
            ["Bản đồ", "Hiển thị nền Đông Nam Á, ranh giới VVHN/VVHM, lưới tọa độ; hỗ trợ phóng to, thu nhỏ, kéo và đặt lại khung nhìn."],
            ["Marker", "Chỉ hiển thị vị trí trong FIR; màu theo LD/O/F/OTHER và xoay biểu tượng theo heading."],
            ["Tổng hợp", "Hiển thị tổng chuyến đang hiển thị và số lượng từng nhóm LD, O/F, OTHER."],
            ["Chi tiết", "Khi chọn marker, hiển thị callsign, Flight ID, OPER_ID, PERMTYPE, tọa độ và thời điểm cập nhật."],
            ["Lỗi", "Khi API/Oracle lỗi, hiển thị thông báo và giữ marker cũ; nếu không tải được GeoJSON thì không khởi động vòng tự cập nhật."],
        ],
        [4.0, 12.2],
    )
    document.add_paragraph("Luồng xử lý:")
    add_bullets(
        document,
        [
            "Đọc vietnam-fir-VVHM-VVHN.geojson và southeast-asia-basemap.geojson từ App_Data để dựng bản đồ.",
            "Page Method GetFlights gọi server-to-server tới Flight Tracking API với ngày UTC hiện tại.",
            "Máy chủ đối chiếu metadata từ Oracle T_DAY_FLIGHTS_GOINGON theo callsign và ngày bay.",
            "Trình duyệt kiểm tra điểm nằm trong polygon FIR, dựng marker, popup và số liệu tổng hợp.",
            "Với nhiều kế hoạch cùng callsign, chỉ chọn bản ghi duy nhất có ETD–ETA chứa thời điểm ADS-B; trường hợp không xác định duy nhất bị loại.",
        ],
    )

    document.add_heading("5. Tool Flight Tracking API", level=1)
    document.add_paragraph(
        "Flight Tracking API là ứng dụng Python độc lập sử dụng Flask, Waitress, psycopg và Flasgger. Tool chạy trên máy "
        "có quyền truy cập PostgreSQL, giúp IIS không cần cài driver hoặc lưu thông tin đăng nhập PostgreSQL."
    )
    add_table(
        document,
        ["Endpoint", "Kết quả/yêu cầu"],
        [
            ["GET /", "Chuyển hướng tới /swagger/."],
            ["GET /swagger/", "Giao diện tài liệu và thử API."],
            ["GET /openapi.json", "Đặc tả OpenAPI của dịch vụ."],
            ["GET /health/live", "Trả HTTP 200 và trạng thái ok khi tiến trình đang hoạt động."],
            ["GET /health/ready", "Trả HTTP 200 khi kết nối PostgreSQL sẵn sàng, ngược lại HTTP 503."],
            ["GET /api/v1/tracks?date=yyyy-MM-dd", "Trả vị trí mới nhất của mỗi callsign trong ngày hợp lệ; ngày mặc định là ngày UTC hiện tại."],
        ],
        [6.0, 10.2],
    )
    document.add_paragraph(
        "Kết quả tracks gồm day, serverTimeUtc, count và flights; mỗi flight có flightIdCurrent, callsign, updatedAtUtc, "
        "latitude, longitude và heading. Callsign được trim, viết hoa và lấy phần trước dấu gạch nối; bản ghi thiếu callsign, "
        "thời gian hoặc tọa độ bị loại; mỗi callsign chỉ giữ vị trí mới nhất."
    )
    document.add_paragraph("Cấu hình và vận hành:")
    add_bullets(
        document,
        [
            "Cấu hình PostgreSQL: host, port, database, username, password, sslmode và connect_timeout.",
            "Cấu hình máy chủ: host, port (mặc định 5088), số thread; giới hạn lookback mặc định 7 ngày, hợp lệ 0–366 ngày.",
            "Hỗ trợ biến môi trường FLIGHT_API_PG_* và FLIGHT_API_HOST/PORT; CLI hỗ trợ --config, --init-config, --check, --host và --port.",
            "Build bằng PyInstaller thành EXE windowed và ZIP; gói ZIP chỉ chứa config mẫu để tránh đóng gói nhầm mật khẩu.",
            "Web.config sử dụng FlightTrackingApiBaseUrl và FlightTrackingApiTimeoutSeconds (1–120 giây, hiện dùng 10 giây).",
        ],
    )

    document.add_heading("6. Tool đồng bộ và ghi dữ liệu T_TRACKS_LOG", level=1)
    document.add_paragraph(
        "TracksSync đọc dữ liệu tăng dần từ PostgreSQL public.tracks theo watermark, phân vùng FIR bằng GeoJSON, "
        "đối chiếu kế hoạch bay Oracle và MERGE dữ liệu vào T_TRACKS_LOG theo khóa CALLSIGN + DATE. Tool cung cấp GUI, CLI và chế độ Auto."
    )
    add_table(
        document,
        ["Bước", "Mô tả"],
        [
            ["1. Đọc nguồn", "Đọc bản ghi có updated_at_utc lớn hơn watermark; --full cho phép backfill toàn bộ; chỉ nhận dòng có tọa độ."],
            ["2. Phân vùng FIR", "STATUS=1 cho VVHN hoặc điểm biên chung; STATUS=2 cho VVHM; điểm ngoài hai FIR bị loại."],
            ["3. Khử trùng", "Theo CALLSIGN + DATE, giữ tọa độ/trạng thái/thời gian mới nhất; TIME_IN là lần đầu vào FIR, TIME_OUT là lần đầu đổi FIR sau đó."],
            ["4. Đối chiếu", "Tìm T_DAY_FLIGHTS_GOINGON theo callsign/ngày; thiếu kế hoạch hoặc tuyến thì PERMTYPE=OTHER; xử lý ETD–ETA qua nửa đêm."],
            ["5. MERGE", "Cập nhật vị trí mới nhất, giữ TIME_IN sớm nhất, không tạo trùng khóa; có thể nâng OTHER khi metadata đầy đủ xuất hiện."],
            ["6. Hoàn tất", "Chỉ ghi watermark sau khi đồng bộ thành công; verify thống kê đúng, sai, thiếu, thừa và trùng khóa."],
        ],
        [4.0, 12.2],
    )
    add_table(
        document,
        ["Cột T_TRACKS_LOG", "Kiểu/ý nghĩa"],
        [
            ["TRLOG_ID", "NUMBER, khóa chính, cấp tự động bởi sequence/trigger."],
            ["CALLSIGN, DATE", "VARCHAR2; khóa nghiệp vụ duy nhất UX_TRACKS_LOG_CALL_DATE."],
            ["FROM_AIRP, TO_AIRP, ETD, ETA", "Thông tin kế hoạch bay được đối chiếu từ Oracle."],
            ["STATUS", "NUMBER(1): 1=VVHN, 2=VVHM."],
            ["PERMTYPE", "LD, O/F hoặc OTHER."],
            ["UPDATED_AT_UTC, TIME_IN, TIME_OUT", "Chuỗi thời gian UTC định dạng yyyy-MM-dd HH:mm:ss."],
            ["LAT, LON", "NUMBER NOT NULL đối với bảng được tạo mới; lưu vị trí ADS-B mới nhất."],
        ],
        [6.0, 10.2],
    )
    document.add_paragraph("Chế độ vận hành:")
    add_bullets(
        document,
        [
            "check: kiểm tra cấu hình/kết nối; sync: đồng bộ và cập nhật watermark; verify: đối chiếu kết quả; all: check → sync → verify.",
            "GUI có các nút thao tác và chế độ Auto theo chu kỳ cấu hình (mặc định 10 giây); Stop chờ lượt DB hiện tại kết thúc an toàn.",
            "Tool tự bảo đảm schema, bổ sung cột/sequence/trigger/index còn thiếu. Nếu có dữ liệu trùng CALLSIGN + DATE, tool dừng để quản trị viên xử lý trước.",
            "Lỗi trong một lượt Auto được ghi log và lượt kế tiếp vẫn tiếp tục; lỗi trước khi hoàn tất không làm tiến watermark.",
        ],
    )

    document.add_heading("7. Yêu cầu chức năng tổng hợp", level=1)
    add_table(
        document,
        ["Mã", "Yêu cầu", "Mức độ"],
        [
            ["FR-ADSB-01", "Hệ thống phải cung cấp báo cáo ADS-B theo khoảng ngày, loại phép bay và hãng khai thác.", "Bắt buộc"],
            ["FR-ADSB-02", "Hệ thống phải tổng hợp KPI và xu hướng LD/O/F/OTHER nhất quán với bảng chi tiết.", "Bắt buộc"],
            ["FR-ADSB-03", "Hệ thống phải cung cấp báo cáo cuối ngày có TIME_IN/TIME_OUT.", "Bắt buộc"],
            ["FR-ADSB-04", "Bản đồ phải hiển thị các chuyến trong VVHN/VVHM và tự cập nhật mỗi 5 giây.", "Bắt buộc"],
            ["FR-ADSB-05", "Marker phải phản ánh vị trí, heading, PERMTYPE và metadata kế hoạch bay.", "Bắt buộc"],
            ["FR-ADSB-06", "API phải cung cấp health, Swagger/OpenAPI và vị trí mới nhất theo callsign.", "Bắt buộc"],
            ["FR-ADSB-07", "TracksSync phải MERGE idempotent theo CALLSIGN + DATE và chỉ tiến watermark sau thành công.", "Bắt buộc"],
            ["FR-ADSB-08", "Tool phải hỗ trợ kiểm tra, đồng bộ, xác minh, backfill và chạy tự động.", "Bắt buộc"],
        ],
        [3.0, 10.7, 2.5],
    )

    document.add_heading("8. Yêu cầu phi chức năng, bảo mật và vận hành", level=1)
    add_bullets(
        document,
        [
            "Hai trang web phải yêu cầu đăng nhập, tương thích giao diện responsive và mã hóa nội dung động trước khi hiển thị.",
            "Bản đồ phải hoạt động với dữ liệu nền offline; vòng cập nhật không được chạy chồng và lỗi tạm thời không được làm treo giao diện.",
            "Flight Tracking API chỉ được mở trong mạng nội bộ, firewall allowlist IP máy IIS; không công bố cổng 5088 ra Internet.",
            "Tài khoản PostgreSQL của API chỉ cần quyền đọc. TracksSync sử dụng tài khoản tối thiểu đủ SELECT nguồn và SELECT/INSERT/UPDATE/DDL cần thiết tại Oracle.",
            "Mật khẩu phải nằm trong file cấu hình local có ACL hệ điều hành hoặc biến môi trường; không commit vào Git và không đóng gói vào ZIP phát hành.",
            "Production phải lựa chọn sslmode PostgreSQL phù hợp; Swagger và API hiện không có API key nên phải được bảo vệ bằng phân đoạn mạng/firewall.",
            "Cần giám sát health/live, health/ready, dung lượng T_TRACKS_LOG, tuổi watermark, số lỗi Auto và độ trễ updated_at_utc.",
        ],
    )

    document.add_heading("9. Tiêu chí nghiệm thu", level=1)
    add_bullets(
        document,
        [
            "Báo cáo tải đúng khoảng ngày, bộ lọc, KPI, biểu đồ, phân trang và modal cuối ngày; dữ liệu rỗng/lỗi có thông báo.",
            "Bản đồ dựng đúng nền/biên FIR, chỉ hiển thị điểm trong VVHN/VVHM, tổng các nhóm bằng tổng marker và cập nhật ổn định mỗi 5 giây.",
            "API live trả 200; ready trả 200 khi PostgreSQL sẵn sàng và 503 khi mất kết nối; ngày sai/ngoài giới hạn trả 400; lỗi nguồn trả 503.",
            "API trả đúng một vị trí mới nhất cho mỗi callsign và không lộ thông tin kết nối PostgreSQL cho trình duyệt.",
            "TracksSync phân vùng đúng FIR/biên, xử lý OTHER và ETD–ETA qua nửa đêm, không tạo trùng CALLSIGN + DATE và bảo toàn watermark khi lỗi.",
            "Schema có đủ bảng, cột, sequence, trigger và unique index; chạy lại bước bảo đảm schema không gây sai lệch dữ liệu.",
            "Bộ unit test Flight Tracking API (6 ca) và TracksSync (18 ca) phải đạt; kiểm thử tích hợp môi trường thật phải xác nhận kết nối PostgreSQL, Oracle và IIS.",
        ],
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--report-image", type=Path, required=True)
    parser.add_argument("--map-image", type=Path, required=True)
    args = parser.parse_args()
    for path in (args.input, args.report_image, args.map_image):
        if not path.is_file():
            raise FileNotFoundError(path)
    build(args.input, args.output, args.report_image, args.map_image)


if __name__ == "__main__":
    main()
