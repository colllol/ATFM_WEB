#!/usr/bin/env python3
"""Build a Word guide for tools/TracksSyncPython only."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt


def clear_body(doc):
    for child in list(doc._element.body):
        if not child.tag.endswith("}sectPr"):
            doc._element.body.remove(child)


def toc(paragraph):
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    items = [("fldChar", "begin"), ("instrText", ' TOC \\o "1-2" \\h \\z \\u '),
             ("fldChar", "separate"), ("t", "Nhấn Ctrl+A rồi F9 trong Word để cập nhật mục lục."),
             ("fldChar", "end")]
    for tag, value in items:
        element = OxmlElement("w:" + tag)
        if tag == "fldChar":
            element.set(namespace + "fldCharType", value)
        elif tag == "instrText":
            element.set(namespace + "space", "preserve")
            element.text = value
        else:
            element.text = value
        run = OxmlElement("w:r")
        run.append(element)
        paragraph._p.append(run)


def table(doc, headers, rows):
    result = doc.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    for i, text in enumerate(headers):
        cell = result.rows[0].cells[i]
        cell.paragraphs[0].add_run(text).bold = True
        shade = OxmlElement("w:shd")
        shade.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill", "D9EAF7")
        cell._tc.get_or_add_tcPr().append(shade)
    for values in rows:
        cells = result.add_row().cells
        for i, text in enumerate(values):
            cells[i].paragraphs[0].add_run(text)


def bullets(doc, values):
    for value in values:
        doc.add_paragraph(value, style="List Bullet")


def build(template, screenshot, output):
    doc = Document(template)
    clear_body(doc)
    doc.core_properties.title = "Hướng dẫn sử dụng TracksSyncPython"
    doc.core_properties.subject = "Đồng bộ dữ liệu ADS-B"
    doc.core_properties.author = "VATM ATFM"

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(130)
    run = title.add_run("TÀI LIỆU HƯỚNG DẪN SỬ DỤNG\nCÔNG CỤ TRACKSSYNCPYTHON")
    run.bold, run.font.name, run.font.size = True, "Arial", Pt(22)
    details = doc.add_paragraph("Hệ thống Quản lý luồng không lưu Việt Nam - ATFM Web\nVị trí mã nguồn: tools\\TracksSyncPython\\\nPhiên bản: 1.1 | Ngày cập nhật: 19/09/2026")
    details.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()
    doc.add_heading("MỤC LỤC", 1)
    toc(doc.add_paragraph())
    doc.add_page_break()

    doc.add_heading("Giới thiệu", 1)
    doc.add_paragraph("Tài liệu này chỉ hướng dẫn tools\\TracksSyncPython. Công cụ đọc dữ liệu vị trí tàu bay từ PostgreSQL public.tracks, lọc theo FIR VVHN/VVHM, đối chiếu Oracle và ghi vào ATFM.T_TRACKS_LOG.")
    table(doc, ["Thành phần", "Chức năng"], [
        ["main.py", "Mã nguồn chính; hỗ trợ giao diện Windows và dòng lệnh."],
        ["TracksSync.sample.json", "Cấu hình mẫu PostgreSQL, Oracle, FIR và watermark."],
        ["ATFM-TracksSync.exe", "Bản giao diện Windows sau khi build."],
        ["ATFM-TracksSync-CLI.exe", "Bản CLI dùng cho Task Scheduler."],
    ])

    doc.add_heading("1. Chuẩn bị và cấu hình", 1)
    bullets(doc, [
        "Máy chạy phải kết nối được PostgreSQL, Oracle và file GeoJSON ranh giới FIR VVHN/VVHM.",
        "PostgreSQL cần quyền đọc public.tracks; Oracle cần đọc T_DAY_FLIGHTS_GOINGON và ghi T_TRACKS_LOG.",
        "Không đưa TracksSync.local.json hoặc file watermark lên Git.",
        "Không chạy chồng nhiều phiên TracksSync trên cùng T_TRACKS_LOG hoặc watermark.",
    ])
    doc.add_paragraph("Thứ tự tìm cấu hình: TRACKS_CONFIG, TracksSync.local.json cạnh mã nguồn/EXE, rồi prjApplication\\App_Data\\TracksSync.local.json.")
    table(doc, ["Nhóm", "Nội dung"], [
        ["postgres", "host, port, database, username, password; nguồn public.tracks."],
        ["oracle", "host, port, service_name, tài khoản và T_DAY_FLIGHTS_GOINGON."],
        ["fir", "geojson_path tuyệt đối đến file vùng VVHN/VVHM."],
        ["watermark_path", "Mốc dữ liệu đã đồng bộ; để trống để tool tự đặt cạnh chương trình."],
    ])

    doc.add_heading("2. Luồng xử lý dữ liệu", 1)
    doc.add_paragraph("Luồng chuẩn: public.tracks -> lọc FIR VVHN/VVHM -> chuẩn hóa callsign -> T_DAY_FLIGHTS_GOINGON -> T_TRACKS_LOG -> watermark.")
    table(doc, ["Quy tắc", "Cách xử lý"], [
        ["Callsign", "Lấy phần trước dấu gạch nối của flight_id_current, bỏ khoảng trắng và chuyển chữ hoa."],
        ["FIR", "VVHN hoặc biên chung có STATUS=1; VVHM có STATUS=2; ngoài FIR bị loại."],
        ["Đối chiếu", "Ghép CALLSIGN + ngày; nếu trùng, chọn duy nhất khoảng ETD-ETA chứa thời điểm track."],
        ["Không ghép được", "Vẫn ghi bản ghi đủ điều kiện với PERMTYPE=OTHER."],
        ["Bản ghi cũ", "Cập nhật tọa độ/thời gian mới nhất và giữ TIME_IN/TIME_OUT đã ghi."],
        ["Watermark", "Chỉ cập nhật sau sync thành công để lượt sau chỉ đọc dữ liệu mới hơn."],
    ])

    doc.add_heading("3. Vận hành bằng giao diện Windows", 1)
    doc.add_paragraph("Mở tools\\dist\\ATFM-TracksSync.exe hoặc chạy python .\\tools\\TracksSyncPython\\main.py.")
    table(doc, ["Nút", "Thao tác"], [
        ["Kiểm tra đối chiếu", "Đọc và đối chiếu dữ liệu; không ghi dữ liệu nghiệp vụ hoặc watermark."],
        ["Ghi T_TRACKS_LOG", "Ghi mới/cập nhật T_TRACKS_LOG và cập nhật watermark."],
        ["Xác minh kết quả", "Thống kê đúng, sai, thiếu, thừa và khóa trùng sau khi ghi."],
        ["Auto", "Lặp check -> sync -> verify theo chu kỳ đã nhập."],
        ["Stop", "Dừng Auto sau khi lượt thao tác cơ sở dữ liệu hiện tại hoàn tất an toàn."],
    ])
    doc.add_paragraph("Quy trình khuyến nghị: Kiểm tra đối chiếu -> đọc log -> Ghi T_TRACKS_LOG -> Xác minh kết quả. Chỉ bật Auto sau khi một lượt thủ công hoàn tất bình thường.")

    doc.add_heading("4. Vận hành bằng dòng lệnh", 1)
    table(doc, ["Lệnh", "Công dụng"], [
        ["--mode check", "Đọc, phân vùng và đối chiếu; không ghi log nghiệp vụ."],
        ["--mode sync", "Ghi/cập nhật T_TRACKS_LOG và watermark."],
        ["--mode verify", "Kiểm tra dữ liệu dự kiến với Oracle; không cập nhật watermark."],
        ["--mode all", "Chạy tuần tự check, sync và verify."],
        ["--full", "Bỏ watermark để quét lại toàn bộ; dùng khi backfill hoặc sửa watermark."],
    ])
    doc.add_paragraph("Ví dụ: ATFM-TracksSync-CLI.exe --mode all. Trước sync --full hoặc all --full trên môi trường thật, dừng Auto, sao lưu/đánh giá dữ liệu và chạy check --full trước.")

    doc.add_heading("5. Giao diện công cụ", 1)
    doc.add_paragraph("Giao diện TracksSyncPython hiển thị các nút Kiểm tra đối chiếu, Ghi T_TRACKS_LOG, Xác minh kết quả, Auto, Stop, trường Chu kỳ và vùng log vận hành.")
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(screenshot), width=Cm(15.8))
    caption = doc.add_paragraph("Hình 1. Giao diện thực tế của công cụ ATFM FIR Tracks Logger (TracksSyncPython).")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("6. Kiểm tra và xử lý sự cố", 1)
    table(doc, ["Hiện tượng", "Cách xử lý"], [
        ["Không kết nối PostgreSQL/Oracle", "Kiểm tra host, port, service, tài khoản, firewall và quyền cơ sở dữ liệu."],
        ["Không có dữ liệu mới", "Kiểm tra watermark, updated_at_utc, tọa độ nguồn và FIR; chạy check --full để chẩn đoán."],
        ["Nhiều bản ghi bị bỏ qua", "Rà soát chuyến bay trùng CALLSIGN + ngày và khoảng ETD-ETA."],
        ["Watermark sai", "Dừng Auto, sao lưu watermark, chạy check --full trước khi đồng bộ lại."],
        ["Có sai/thiếu/trùng", "Không bật Auto; đọc log, sửa nguyên nhân rồi chạy lại check -> sync -> verify."],
    ])
    doc.add_heading("7. Checklist vận hành", 1)
    bullets(doc, ["Đúng file cấu hình, kết nối PostgreSQL/Oracle và file FIR GeoJSON.", "Log không có traceback; sai, thiếu và khóa trùng bằng 0.", "Watermark chỉ tăng sau sync/all thành công.", "Không chạy song song nhiều Auto hoặc Task Scheduler trên cùng watermark.", "Sao lưu/đánh giá dữ liệu trước khi dùng --full trên môi trường thật."])
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main():
    parser = argparse.ArgumentParser(description="Tạo hướng dẫn sử dụng TracksSyncPython.")
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--screenshot", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.template, args.screenshot, args.output)


if __name__ == "__main__":
    main()
