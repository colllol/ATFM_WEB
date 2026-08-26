from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Cm, Pt
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT_DIR = Path(__file__).resolve().parents[1] / "TaiLieu" / "BienBanNghiemThu_202608"
PROJECT = "Cập nhật, hiệu chỉnh cơ sở dữ liệu, phần mềm hệ thống số liệu điều hành bay"
PACKAGE = "Gói thầu số 1: Cập nhật, hiệu chỉnh cơ sở dữ liệu, phần mềm hệ thống số liệu điều hành bay"


def set_cell_shading(cell, fill="F2F2F2"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            tag = "w:%s" % edge
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for key in ["val", "sz", "space", "color"]:
                if key in kwargs[edge]:
                    element.set(qn("w:%s" % key), str(kwargs[edge][key]))


def configure(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True


def para(doc, text="", bold=False, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    return p


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        set_cell_shading(c, "D9EAF7")
        for r in c.paragraphs[0].runs:
            r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()
    return t


def add_signature_table(doc, role_count=6):
    rows = [(f"{i}.", "Ông/Bà: …………………………………", "Chức vụ: ………………………", "Ký, ghi rõ họ tên") for i in range(1, role_count + 1)]
    add_table(doc, ["STT", "Đại diện", "Chức vụ", "Xác nhận"], rows)


def add_image_placeholder(doc, code, caption, instruction):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    set_cell_shading(cell, "F7F7F7")
    set_cell_border(cell, top={"val": "single", "sz": 12, "color": "808080"}, left={"val": "single", "sz": 12, "color": "808080"}, bottom={"val": "single", "sz": 12, "color": "808080"}, right={"val": "single", "sz": 12, "color": "808080"})
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run(f"{code}\n{caption}\n\n{instruction}")
    r.italic = True
    r.font.size = Pt(11)
    para(doc, f"Mã bằng chứng: {code}. Ảnh phải ghi ngày/giờ chụp, màn hình hoặc máy chủ, người thực hiện và được lưu cùng hồ sơ nghiệm thu.")


def add_image_evidence(doc, path, code, caption):
    if not path.exists():
        add_image_placeholder(doc, code, caption, "CHƯA CÓ ẢNH TƯƠNG ỨNG – BỔ SUNG ẢNH THỰC TẾ")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(15.5))
    para(doc, f"Hình {code}. {caption}. Nguồn ảnh: giao diện Report New đã lưu trong hồ sơ kỹ thuật; cần ghi bổ sung ngày/giờ chụp, URL, môi trường và người chụp trước khi ký nghiệm thu.", align=WD_ALIGN_PARAGRAPH.CENTER)


def common_header(doc, title, date_line, subtitle=""):
    para(doc, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", True, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "Độc lập - Tự do - Hạnh phúc", True, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "________________", align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, date_line, align=WD_ALIGN_PARAGRAPH.RIGHT)
    para(doc, title, True, WD_ALIGN_PARAGRAPH.CENTER)
    if subtitle:
        para(doc, subtitle, True, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, f"DỰ ÁN: {PROJECT}", True)
    para(doc, f"GÓI THẦU: {PACKAGE}", True)


def technical_installation():
    doc = Document()
    configure(doc)
    common_header(doc, "BIÊN BẢN NGHIỆM THU KỸ THUẬT", "………, ngày 07 tháng 08 năm 2026", "CÀI ĐẶT PHẦN MỀM TRƯỚC KHI KIỂM THỬ")
    heading(doc, "I. Đối tượng nghiệm thu", 1)
    para(doc, "Cài đặt phần mềm hệ thống, phần mềm nội bộ và cấu hình kỹ thuật cho môi trường ATFM_WEB trước khi thực hiện kiểm thử hệ thống. Phạm vi gồm mã ứng dụng, runtime/IIS, cơ sở dữ liệu Oracle, cấu hình kết nối, tài khoản/phân quyền, ghi nhật ký, đồng bộ thời gian, sao lưu và giám sát.")
    heading(doc, "II. Thành phần trực tiếp nghiệm thu", 1)
    add_signature_table(doc)
    heading(doc, "III. Thời gian, địa điểm", 1)
    add_table(doc, ["Nội dung", "Thông tin"], [
        ("Thời gian", "Ngày 07/08/2026"),
        ("Địa điểm", "………………………………………………………………"),
        ("Môi trường", "Máy chủ/VM triển khai ATFM_WEB: …………………………………"),
    ])
    heading(doc, "IV. Nội dung và kết quả nghiệm thu", 1)
    add_table(doc, ["STT", "Hạng mục", "Kết quả/ghi nhận", "Trạng thái"], [
        ("1", "IIS, website và runtime", "Đã cài đặt, cấu hình binding, application pool và kiểm tra khởi động.", "Đạt/Không đạt"),
        ("2", "Ứng dụng ATFM_WEB", "Đã triển khai đúng gói build/checksum: …………………………………", "Đạt/Không đạt"),
        ("3", "Oracle/package/script", "Đã chạy script triển khai, kiểm tra package/schema và kết nối ứng dụng.", "Đạt/Không đạt"),
        ("4", "Tài khoản và phân quyền", "Đã tạo/kiểm tra tài khoản kỹ thuật, role và data scope theo phê duyệt.", "Đạt/Không đạt"),
        ("5", "Kết nối tích hợp", "Đã kiểm tra kết nối AMHS/AFTN, ADS-B, Bravo, API/Email theo phạm vi môi trường.", "Đạt/Không đạt"),
        ("6", "Logging, backup, monitoring", "Đã kiểm tra log, lịch sao lưu, cảnh báo và phương án rollback.", "Đạt/Không đạt"),
    ])
    add_image_placeholder(doc, "IMG-INSTALL-01", "Ảnh sau cài đặt: máy chủ/IIS và trạng thái dịch vụ", "CHÈN ẢNH CHỤP THỰC TẾ TẠI ĐÂY")
    add_image_placeholder(doc, "IMG-INSTALL-02", "Ảnh sau cài đặt: màn hình ATFM_WEB hoặc trang kiểm tra kết nối", "CHÈN ẢNH CHỤP THỰC TẾ TẠI ĐÂY")
    heading(doc, "V. Kết luận", 1)
    para(doc, "Căn cứ kết quả kiểm tra và các bằng chứng kèm theo, các bên thống nhất: môi trường cài đặt phần mềm trước kiểm thử đạt/không đạt điều kiện chuyển sang giai đoạn chạy thử. Các tồn tại (nếu có): ........................................................................................................................")
    heading(doc, "VI. Xác nhận và chữ ký", 1)
    add_signature_table(doc)
    return doc


def trial_run():
    doc = Document()
    configure(doc)
    common_header(doc, "BIÊN BẢN NGHIỆM THU CHẠY THỬ", "………, ngày 03 tháng 09 năm 2026", "CHẠY THỬ HỆ THỐNG VÀ CÁC TÍNH NĂNG NGHIỆP VỤ")
    heading(doc, "I. Phạm vi và thời gian chạy thử", 1)
    para(doc, "Chạy thử trên môi trường đã nghiệm thu cài đặt, bao gồm kiểm tra luồng tích hợp, cảnh báo, điện văn, phê duyệt, KHB, báo cáo và AI theo kế hoạch kiểm thử được phê duyệt.")
    add_table(doc, ["Nội dung", "Thông tin"], [
        ("Thời gian", "Từ ngày 09/08/2026 đến ngày 03/09/2026"),
        ("Địa điểm/môi trường", "………………………………………………………………"),
        ("Phiên bản chạy thử", "………………………………… Build/Release: ………………………"),
        ("Tài liệu đối chiếu", "SRS/SDD, kế hoạch kiểm thử, bộ test case và nhật ký lỗi"),
    ])
    heading(doc, "II. Thành phần tham gia", 1)
    add_signature_table(doc)
    heading(doc, "III. Kết quả chạy thử", 1)
    add_table(doc, ["STT", "Nhóm/tính năng", "Phạm vi kiểm tra", "Kết quả", "Bằng chứng"], [
        ("1", "Tích hợp và tự động hóa", "Email/File, ADS-B, SLOT/KHH, AMHS/AFTN, API Gateway", "Đạt/Không đạt", "TC-INT-…"),
        ("2", "Cảnh báo và tiện ích", "Notification runtime, Live Fire, Daily Statistic, KHB quân sự", "Đạt/Không đạt", "TC-ALT-…"),
        ("3", "Nâng cấp và tối ưu", "Tra cứu, batch, cảnh báo, đồng bộ, đối soát", "Đạt/Không đạt", "TC-OPT-…"),
        ("4", "Báo cáo và phân tích", "Dashboard, KPI, ADS-B, anomaly, sân bay/quân sự", "Đạt/Không đạt", "TC-RPT-…"),
        ("5", "AI hỗ trợ", "NL2SQL, truy vấn, chatbot, giám sát", "Đạt/Không đạt", "TC-AI-…"),
        ("6", "Phân quyền và an toàn", "RBAC, session, audit, XSS/CSRF, backup/restore", "Đạt/Không đạt", "TC-SEC-…"),
    ])
    heading(doc, "IV. Danh sách vấn đề và xử lý", 1)
    add_table(doc, ["Mã lỗi/ý kiến", "Mô tả", "Mức độ", "Biện pháp", "Trạng thái"], [
        ("", "", "", "", "Đóng/Mở"),
        ("", "", "", "", "Đóng/Mở"),
        ("", "", "", "", "Đóng/Mở"),
    ])
    heading(doc, "V. Hình ảnh tính năng trong quá trình chạy thử", 1)
    for code, caption, image_name in [
        ("IMG-TRIAL-01", "Màn hình tổng quan/trạng thái khai thác", "image1.png"),
        ("IMG-TRIAL-02", "Màn hình phân tích xu hướng hoặc bất thường", "image3.png"),
        ("IMG-TRIAL-03", "Màn hình báo cáo quân sự/dân dụng", "image5.png"),
        ("IMG-TRIAL-04", "Màn hình biểu đồ sân bay và cất hạ cánh", "image8.png"),
    ]:
        add_image_evidence(doc, OUT_DIR / "HinhAnh_ThamChieu" / image_name, code, caption)
    heading(doc, "VI. Kết luận nghiệm thu chạy thử", 1)
    para(doc, "Các bên xác nhận kết quả chạy thử trong khoảng thời gian nêu trên: đạt/đạt có điều kiện/không đạt. Điều kiện hoặc tồn tại cần theo dõi: ........................................................................................................................")
    heading(doc, "VII. Xác nhận và chữ ký", 1)
    add_signature_table(doc)
    return doc


def training_acceptance():
    doc = Document()
    configure(doc)
    common_header(doc, "BIÊN BẢN NGHIỆM THU DỊCH VỤ ĐÀO TẠO", "………, ngày 04 tháng 09 năm 2026", "ĐÀO TẠO, HƯỚNG DẪN SỬ DỤNG VÀ VẬN HÀNH")
    heading(doc, "I. Nội dung dịch vụ đào tạo", 1)
    para(doc, "Nghiệm thu dịch vụ đào tạo, hướng dẫn sử dụng và vận hành Hệ thống số liệu điều hành bay cho người khai thác, người duyệt, cán bộ quản trị và cán bộ hỗ trợ. Nội dung đào tạo bám theo tài liệu hướng dẫn sử dụng, quản trị và các chức năng đã chạy thử.")
    add_table(doc, ["Nội dung", "Thông tin"], [
        ("Ngày đào tạo/nghiệm thu", "04/09/2026"),
        ("Địa điểm/hình thức", "………………………………………………………………"),
        ("Tài liệu đào tạo", "Mã tài liệu: ……………………… Phiên bản: ……………………"),
        ("Giảng viên/đơn vị đào tạo", "………………………………………………………………"),
    ])
    heading(doc, "II. Nội dung đã đào tạo", 1)
    add_table(doc, ["STT", "Chuyên đề", "Kết quả"], [
        ("1", "Đăng nhập, phiên làm việc, phân quyền và an toàn thông tin", "Đã đào tạo/Chưa đào tạo"),
        ("2", "Nhập, sửa, duyệt và phát điện văn", "Đã đào tạo/Chưa đào tạo"),
        ("3", "KHB ngày, KHB quân sự, Daily Statistic và báo cáo", "Đã đào tạo/Chưa đào tạo"),
        ("4", "Cảnh báo runtime, tra cứu, đối soát và xử lý lỗi", "Đã đào tạo/Chưa đào tạo"),
        ("5", "Sao lưu, phục hồi, giám sát và hỗ trợ vận hành", "Đã đào tạo/Chưa đào tạo"),
    ])
    heading(doc, "III. Danh sách học viên và chữ ký xác nhận", 1)
    para(doc, "Danh sách dưới đây phải được điền theo danh sách tham dự thực tế; mỗi học viên ký xác nhận sau khi hoàn thành nội dung đào tạo.")
    rows = [(str(i), "………………………………", "Đơn vị: ……………………", "Vai trò: ……………", "", "") for i in range(1, 16)]
    add_table(doc, ["STT", "Họ và tên học viên", "Đơn vị", "Vai trò", "Ký buổi học", "Ký nghiệm thu"], rows)
    heading(doc, "IV. Hình ảnh đào tạo", 1)
    add_image_placeholder(doc, "IMG-TRAIN-01", "Ảnh lớp/phiên đào tạo và giảng viên", "CHÈN ẢNH THỰC TẾ TẠI ĐÂY")
    add_image_placeholder(doc, "IMG-TRAIN-02", "Ảnh học viên thực hành trên phần mềm", "CHÈN ẢNH THỰC TẾ TẠI ĐÂY")
    heading(doc, "V. Kết luận", 1)
    para(doc, "Các bên xác nhận dịch vụ đào tạo đã được thực hiện đầy đủ/đầy đủ có điều kiện/không đạt theo nội dung trên. Tài liệu, danh sách học viên và chữ ký kèm theo là một phần không tách rời của biên bản.")
    heading(doc, "VI. Xác nhận và chữ ký các bên", 1)
    add_signature_table(doc)
    return doc


def product_acceptance():
    doc = Document()
    configure(doc)
    common_header(doc, "BIÊN BẢN NGHIỆM THU SẢN PHẨM", "………, ngày 05 tháng 09 năm 2026", "NGHIỆM THU SẢN PHẨM/HẠNG MỤC CÔNG VIỆC CỦA DỰ ÁN")
    heading(doc, "I. Đối tượng nghiệm thu", 1)
    para(doc, "Sản phẩm cập nhật, hiệu chỉnh cơ sở dữ liệu và phần mềm nội bộ Hệ thống số liệu điều hành bay thuộc Gói thầu số 1, sau khi hoàn thành cài đặt, chạy thử, đào tạo và khắc phục các tồn tại theo hồ sơ được phê duyệt.")
    para(doc, "Phạm vi gồm 05 nhóm: tích hợp và tự động hóa dữ liệu; cảnh báo và tiện ích hỗ trợ; nâng cấp và tối ưu HTSLB; báo cáo và phân tích; AI hỗ trợ thống kê, tìm kiếm và tổng hợp. Danh mục chi tiết đối chiếu theo Phụ lục 01 và SRS/SDD phiên bản được phê duyệt.")
    heading(doc, "II. Thành phần tham gia nghiệm thu", 1)
    add_signature_table(doc)
    heading(doc, "III. Hồ sơ và căn cứ nghiệm thu", 1)
    add_table(doc, ["STT", "Hồ sơ/căn cứ", "Mã/số, ngày, phiên bản", "Tình trạng"], [
        ("1", "Hợp đồng/gói thầu và phụ lục", "………………………………", "Đủ/Thiếu"),
        ("2", "SRS và SDD/thiết kế chi tiết", "………………………………", "Đủ/Thiếu"),
        ("3", "Biên bản nghiệm thu kỹ thuật cài đặt", "07/08/2026", "Đính kèm"),
        ("4", "Biên bản nghiệm thu chạy thử", "03/09/2026", "Đính kèm"),
        ("5", "Biên bản nghiệm thu dịch vụ đào tạo", "04/09/2026", "Đính kèm"),
        ("6", "Báo cáo kiểm thử/ma trận truy vết", "………………………………", "Đủ/Thiếu"),
    ])
    heading(doc, "IV. Kết quả nghiệm thu sản phẩm", 1)
    add_table(doc, ["Nhóm sản phẩm", "Kết quả đối chiếu", "Bằng chứng", "Kết luận"], [
        ("I. Tích hợp và tự động hóa dữ liệu", "Đối chiếu FR-INT-001…005 và luồng nguồn–đích", "Test/log/reconcile", "Đạt/Không đạt"),
        ("II. Cảnh báo và tiện ích hỗ trợ", "Đối chiếu FR-ALT-001/002, Live Fire, Daily, Military", "Test/ảnh/audit", "Đạt/Không đạt"),
        ("III. Nâng cấp và tối ưu HTSLB", "Đối chiếu FR-OPT-001…032", "Test/performance/export", "Đạt/Không đạt"),
        ("IV. Báo cáo và phân tích", "Đối chiếu FR-RPT-001…010", "Snapshot/chart/export", "Đạt/Không đạt"),
        ("V. AI hỗ trợ", "Đối chiếu FR-AI-001…004", "SQL guard/result/audit", "Đạt/Không đạt"),
    ])
    add_image_evidence(doc, OUT_DIR / "HinhAnh_ThamChieu" / "image2.png", "IMG-PRODUCT-01", "Ảnh tổng quan sản phẩm sau hoàn thiện")
    add_image_evidence(doc, OUT_DIR / "HinhAnh_ThamChieu" / "image7.png", "IMG-PRODUCT-02", "Ảnh màn hình đại diện các nhóm chức năng")
    heading(doc, "V. Kết luận và kiến nghị", 1)
    para(doc, "Căn cứ hồ sơ, kết quả kiểm thử, chạy thử, đào tạo và đối chiếu sản phẩm, các bên thống nhất nghiệm thu sản phẩm/hạng mục: đạt/đạt có điều kiện/không đạt. Kiến nghị và điều kiện bảo hành/hỗ trợ: ........................................................................................................................")
    heading(doc, "VI. Xác nhận và chữ ký", 1)
    add_signature_table(doc)
    return doc


def save(doc, name):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / name
    doc.save(path)
    return path


def main():
    paths = [
        save(technical_installation(), "01_BBNT_KyThuat_CaiDat_20260807.docx"),
        save(trial_run(), "02_BBNT_ChayThu_Den_20260903.docx"),
        save(training_acceptance(), "03_BBNT_DaoTao_20260904.docx"),
        save(product_acceptance(), "04_BBNT_SanPhamDuAn_20260905.docx"),
    ]
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
