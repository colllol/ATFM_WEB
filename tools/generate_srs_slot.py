#!/usr/bin/env python3
"""Create a standalone SLOT SRS from the current AeroSync Word template."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


def clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if not child.tag.endswith("}sectPr"):
            body.remove(child)


def add_cover(document: Document) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(130)
    run = title.add_run("TÀI LIỆU ĐẶC TẢ YÊU CẦU PHẦN MỀM")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(20)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("PHÂN HỆ QUẢN LÝ VÀ ĐỐI CHIẾU SLOT")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)
    info = document.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_before = Pt(48)
    info.add_run(
        "Phạm vi: Import, tra cứu và đối chiếu dữ liệu SLOT/KHH/Phép bay\n"
        "Mô-đun: ATFM Excel Importer, Web SLOTS và API\n"
        "Phiên bản tài liệu: 1.0\n"
        "Ngày cập nhật: 12/08/2026"
    )
    document.add_page_break()


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "TableGrid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
        for run in table.rows[0].cells[index].paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    return table


def add_picture(document: Document, path: Path, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Cm(16.2))
    cap = document.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build(template: Path, output: Path, images: dict[str, Path]) -> None:
    document = Document(str(template))
    clear_body(document)
    add_cover(document)

    document.add_heading("1. Thông tin chung", level=1)
    add_table(
        document,
        ["Nội dung", "Giá trị"],
        [
            ["Tên tài liệu", "Đặc tả yêu cầu phân hệ Quản lý và đối chiếu SLOT"],
            ["Mã tài liệu", "SRS-SLOT-001"],
            ["Hệ thống", "VATM ATFM Web"],
            ["Phiên bản", "1.0"],
            ["Ngày cập nhật", "12/08/2026"],
            ["Phạm vi", "ATFM Excel Importer; SlotAero; KHH; SlotComparison; SlotResults"],
        ],
        [5.0, 11.2],
    )
    document.add_paragraph(
        "Tài liệu mô tả quy trình tiếp nhận dữ liệu kế hoạch hãng và SLOT sân bay từ Excel, lưu vào Oracle, "
        "tra cứu trên web, đối chiếu với dữ liệu phép bay và khai thác kết quả KQ1–KQ4."
    )

    document.add_heading("2. Phạm vi và kiến trúc tổng thể", level=1)
    add_table(
        document,
        ["Thành phần", "Vai trò", "Nguồn/đích"],
        [
            ["ATFM Excel Importer", "Quét, kiểm tra và import Excel KHH/SLOT", "Excel → Oracle T_KHH, T_SLOT_AERO"],
            ["SlotAero.aspx", "Tra cứu dữ liệu SLOT sân bay", "POST api/SlotAero/GetData"],
            ["KHH.aspx", "Tra cứu kế hoạch chuyến bay hãng", "POST api/Khh/GetData"],
            ["SlotComparison.aspx", "Đối chiếu SLOT – Phép bay – KHH và lưu KQ1–KQ4", "Oracle T_SLOT_AERO, T_DAY_FLIGHTS, T_KHH"],
            ["SlotResults.aspx", "Tra cứu kết quả đã lưu và xuất Excel", "API SlotComparison/GetBootstrap, GetResults"],
        ],
        [4.0, 6.5, 5.7],
    )
    document.add_paragraph(
        "Luồng chính: tệp Excel KHH/SLOT → ATFM Excel Importer → Oracle → API/Web tra cứu → "
        "SlotComparison đối chiếu với phép bay → lưu kết quả → SlotResults tra cứu và xuất Excel."
    )
    document.add_paragraph("Tác nhân:")
    add_bullets(
        document,
        [
            "Nhân viên nhập liệu: lựa chọn thư mục, kiểm tra và chạy import Excel.",
            "Nhân viên khai thác SLOT: tra cứu dữ liệu SLOT/KHH và thực hiện đối chiếu.",
            "Nhân viên báo cáo: tra cứu KQ1–KQ4 và xuất Excel.",
            "Quản trị viên: cấu hình kết nối, phân quyền menu, triển khai schema/API và kiểm soát bản phát hành tool.",
        ],
    )

    document.add_heading("3. Tool ATFM Excel Importer", level=1)
    document.add_paragraph(
        "Gói triển khai tại C:/Users/Admin/Downloads/QLB/dist gồm bản release one-file ATFM_Excel_Importer.exe "
        "và bản debug one-folder. Ứng dụng desktop Tkinter chạy trên Windows, đọc Excel bằng openpyxl và ghi Oracle "
        "qua python-oracledb; tool không gọi HTTP API."
    )
    add_table(
        document,
        ["Bước", "Xử lý"],
        [
            ["1. Chọn dữ liệu", "Người dùng chọn thư mục NGAY ddMM; scanner chỉ nhận nhóm KHH và SLOT CHK/SLOT."],
            ["2. Kiểm tra", "Quét .xlsx/.xlsm, bỏ file tạm ~$; hiển thị nhóm/file và cho phép chọn, đổi thứ tự hàng đợi."],
            ["3. Kết nối", "Nhập host, port, service, user và password Oracle; password không được ghi ra log."],
            ["4. Phân tích", "Đọc workbook read-only/data-only, chuẩn hóa ngày, giờ, callsign, sân bay và dữ liệu hãng."],
            ["5. Import", "Worker nền import tuần tự; mỗi file dùng savepoint/transaction riêng, commit khi thành công và rollback khi lỗi."],
            ["6. Tổng kết", "Hiển thị tiến độ, số file thành công/thất bại, số dòng mới/đã có; hủy tại ranh giới file/thư mục."],
        ],
        [4.0, 12.2],
    )
    document.add_heading("3.1. Dữ liệu SLOT", level=2)
    add_bullets(
        document,
        [
            "Tên file theo dạng SLOT_<ICAO>_<ngày>.xlsx; sân bay đi lấy từ tên file, ngày bay lấy từ thư mục NGAY ddMM và năm hệ thống.",
            "Header A:E bắt buộc gồm ETD, CARRIER, AC, SEATS, TO; cột F là callsign; ETD trống được kế thừa từ dòng trước.",
            "Chuẩn hóa ETD HH:MM và kiểm tra 00:00–23:59; ghi ETD_ETA, CARRIE, AC, SEATS, FROM_AIRP, TO_AIRP, CALLSIGN, FLIGHT_DATE, AERO.",
            "T_SLOT_AERO được tạo/migration khi cần; bản ghi đã tồn tại theo toàn bộ chín trường nghiệp vụ được bỏ qua.",
        ],
    )
    document.add_heading("3.2. Dữ liệu KHH", level=2)
    add_bullets(
        document,
        [
            "Hỗ trợ cấu trúc hãng BAV, PIC, VAG, VFC, VJC và VNA; chỉ nhận dữ liệu thuộc ngày của thư mục.",
            "Ghi các trường Date, CALLSIGN, From, To, ETD, ETA, OPER vào T_KHH.",
            "Khử trùng trong nguồn KHH theo bảy trường và bỏ qua bản ghi đã có trong cơ sở dữ liệu.",
            "File sai header, tên, ngày, giờ hoặc độ dài trường phải báo rõ và không để dữ liệu dở dang của file đó.",
        ],
    )
    document.add_heading("3.3. Cấu hình và giới hạn hiện tại", level=2)
    add_bullets(
        document,
        [
            "Mặc định host 172.29.187.90, port 1521, service PDBORCL, user ATFM; password nhập trên GUI hoặc ATFM_DB_PASSWORD.",
            "Log chỉ hiển thị trong GUI, chưa lưu file bền vững, chưa có timestamp/correlation ID/audit run.",
            "EXE hiện chưa ký Authenticode và chưa có FileVersion/ProductVersion; release cần bổ sung version, chữ ký và manifest SHA-256.",
            "Năm dữ liệu SLOT suy ra từ năm hiện tại của máy; cần kiểm soát khi import lịch sử hoặc thời điểm chuyển năm.",
            "Không nên chạy nhiều instance đồng thời vì bảng chưa có cơ chế MERGE/unique constraint bao phủ mọi trường hợp trùng nội-file.",
        ],
    )

    document.add_page_break()
    document.add_heading("4. Trang SLOT sân bay (SLOTS/SlotAero.aspx)", level=1)
    document.add_paragraph(
        "Trang chỉ đọc dùng control SlotsTable để tra cứu T_SLOT_AERO qua POST api/SlotAero/GetData; "
        "không import, thêm, sửa, xóa hoặc xuất file."
    )
    add_picture(document, images["slot_aero"], "Hình 1. Giao diện toàn màn hình trang SLOT sân bay")
    add_table(
        document,
        ["Nhóm", "Yêu cầu"],
        [
            ["Bộ lọc", "Tìm kiếm tự do, ngày bay yyyy-MM-dd, sân bay tối đa 8 ký tự và số dòng 100–8000."],
            ["Dữ liệu", "ETD_ETA, CARRIE, AC, SEATS, FROM_AIRP, TO_AIRP, CALLSIGN, FLIGHT_DATE, AERO, ID."],
            ["Hiển thị", "Metadata cột động, STT toàn cục, header và ba cột trình bày đầu cố định, cuộn ngang/dọc."],
            ["Điều hướng", "Tìm kiếm về trang 1; đổi page size về trang 1; làm mới xóa bộ lọc; trang trước/sau theo TotalPages."],
            ["Lỗi", "HTTP/API/envelope lỗi hiển thị thông báo đã encode, xóa bảng và khóa nút phân trang."],
        ],
        [4.0, 12.2],
    )

    document.add_page_break()
    document.add_heading("5. Trang Kế hoạch chuyến bay (SLOTS/KHH.aspx)", level=1)
    document.add_paragraph(
        "Trang chỉ đọc dùng cùng control SlotsTable, nguồn KHH và POST api/Khh/GetData. Giao diện và cơ chế "
        "phân trang/lỗi thống nhất với SlotAero."
    )
    add_picture(document, images["khh"], "Hình 2. Giao diện toàn màn hình trang Kế hoạch chuyến bay")
    add_table(
        document,
        ["Nhóm", "Yêu cầu"],
        [
            ["Bộ lọc", "Tìm kiếm tự do, ngày áp lên cột Date, sân bay áp lên From/To, số dòng 100–8000."],
            ["Dữ liệu", "ID, Date, CALLSIGN, From, To, ETD, ETA và OPER."],
            ["Hiển thị", "Bảng metadata động, STT liên tục giữa các trang, sticky header/cột và giá trị đầy đủ trong tooltip."],
            ["Validation", "Ngày rỗng hợp lệ; ngày nhập phải đúng yyyy-MM-dd; sân bay được trim và chuyển chữ hoa."],
            ["An toàn", "Nguồn dữ liệu nằm trong allow-list T_KHH/T_SLOT_AERO; mọi dữ liệu động được HTML/attribute encode."],
        ],
        [4.0, 12.2],
    )

    document.add_page_break()
    document.add_heading("6. Trang thực hiện đối chiếu (SLOTS/SlotComparison.aspx)", level=1)
    document.add_paragraph(
        "Trang đối chiếu ba nguồn T_SLOT_AERO, T_DAY_FLIGHTS và T_KHH theo ngày/OPER, phân loại sai lệch KQ1–KQ4 "
        "và lưu kết quả Oracle trong transaction."
    )
    add_picture(document, images["comparison"], "Hình 3. Giao diện toàn màn hình trang Đối chiếu SLOT – Phép bay – Kế hoạch hãng")
    add_table(
        document,
        ["Loại", "Quy tắc"],
        [
            ["KQ1", "Có SLOT nhưng không có Phép bay."],
            ["KQ2", "Có Phép bay nhưng không có SLOT."],
            ["KQ3", "Có Kế hoạch hãng nhưng thiếu SLOT hoặc thiếu Phép bay."],
            ["KQ4", "Có đồng thời SLOT và Phép bay nhưng không có Kế hoạch hãng."],
        ],
        [4.0, 12.2],
    )
    document.add_paragraph("Quy tắc và luồng xử lý:")
    add_bullets(
        document,
        [
            "Ngày mặc định là ngày mới nhất đồng thời có trong T_KHH và T_SLOT_AERO; OPER mặc định ALL; phân trang 100 dòng.",
            "Chuẩn hóa chữ hoa, callsign A-Z0-9 và tiền tố ICAO, sân bay IATA→ICAO, ETD bốn ký tự; loại record thiếu callsign/tuyến.",
            "Khóa so khớp hiện tại là CALLSIGN|FROM_AIRP|TO_AIRP; ETD không tham gia khóa; bản ghi trùng nguồn giữ dòng đầu tiên.",
            "Khi bấm KQ, hệ thống tải ba nguồn, tính sai lệch, sắp xếp, tạo/cập nhật run, xóa kết quả cũ cùng phạm vi và array-bind kết quả mới.",
            "Mọi thao tác xóa/chèn nằm trong một transaction; lỗi phải rollback và hiển thị thông báo được encode.",
        ],
    )

    document.add_page_break()
    document.add_heading("7. Trang kết quả đối chiếu (SLOTS/SlotResults.aspx)", level=1)
    document.add_paragraph(
        "Trang chỉ đọc kết quả KQ1–KQ4 đã lưu, lọc theo ngày/OPER, phân trang server-side và xuất Excel theo biểu mẫu BM.QLL-ĐCKSPK."
    )
    add_picture(document, images["results"], "Hình 4. Giao diện toàn màn hình trang Kết quả đối chiếu")
    add_table(
        document,
        ["Chức năng", "Yêu cầu"],
        [
            ["Khởi tạo", "POST api/SlotComparison/GetBootstrap để nhận DefaultDate và Operators."],
            ["Tra cứu", "POST api/SlotComparison/GetResults với ngày, OPER, KQ, PageIndex và PageSize=100."],
            ["Chuyển KQ", "Chỉ đọc loại kết quả khác và về trang 1; không chạy lại đối chiếu, không ghi/xóa Oracle."],
            ["Tổng hợp", "Hiển thị số lượng KQ1–KQ4, bảng tám cột, tổng dòng và trang hiện tại/tổng trang."],
            ["Xuất Excel", "Đọc tuần tự tối đa 8.000 dòng/lô, xuất toàn bộ KQ đang chọn thành {KQ}_{ddMMyy}_{OPER}.xls UTF-8 BOM."],
            ["Lỗi", "Báo rõ thiếu API URL, HTTP ngoài 2xx, JSON/envelope sai, Code khác 00 hoặc ListValue null; timeout 120 giây."],
        ],
        [4.0, 12.2],
    )

    document.add_heading("8. Yêu cầu dữ liệu và quy tắc nghiệp vụ", level=1)
    add_table(
        document,
        ["Đối tượng", "Yêu cầu chính"],
        [
            ["T_KHH", "Date, CALLSIGN, From, To, ETD, ETA, OPER; dữ liệu chuẩn hóa từ kế hoạch hãng."],
            ["T_SLOT_AERO", "ETD_ETA, CARRIE, AC, SEATS, FROM_AIRP, TO_AIRP, CALLSIGN, FLIGHT_DATE, AERO, ID."],
            ["T_DAY_FLIGHTS", "Nguồn phép bay dùng khi đối chiếu: FLIGHTDATE, FLIGHTNBR, FROM_AIRP, TO_AIRP, ETD, OPER_ID."],
            ["M_OPER/M_AERO", "Danh mục chuẩn hóa IATA/ICAO hãng và sân bay; VNA được ánh xạ thành HVN trong đối chiếu."],
            ["T_SLOT_COMPARE_RUN", "Phiên đối chiếu duy nhất theo ngày, lưu người thực hiện và thời điểm cập nhật."],
            ["T_SLOT_COMPARE_RESULT", "Chi tiết KQ1–KQ4, liên kết RUN_ID và hỗ trợ lọc theo RESULT_TYPE/OPER."],
        ],
        [5.0, 11.2],
    )

    document.add_heading("9. Yêu cầu chức năng tổng hợp", level=1)
    add_table(
        document,
        ["Mã", "Yêu cầu", "Mức độ"],
        [
            ["FR-SLOT-01", "Tool phải quét, kiểm tra và import Excel KHH/SLOT vào đúng bảng Oracle.", "Bắt buộc"],
            ["FR-SLOT-02", "Mỗi file import phải có transaction độc lập và rollback khi lỗi.", "Bắt buộc"],
            ["FR-SLOT-03", "SlotAero/KHH phải hỗ trợ lọc và phân trang dữ liệu API ở chế độ chỉ đọc.", "Bắt buộc"],
            ["FR-SLOT-04", "Hệ thống phải tính và lưu đúng bốn loại KQ1–KQ4 theo ngày/OPER.", "Bắt buộc"],
            ["FR-SLOT-05", "Chạy lại cùng ngày/OPER/KQ phải thay thế kết quả cũ, không cộng dồn.", "Bắt buộc"],
            ["FR-SLOT-06", "SlotResults phải tra cứu mà không thay đổi dữ liệu và xuất đủ kết quả ra Excel.", "Bắt buộc"],
            ["FR-SLOT-07", "Giao diện phải hiển thị rõ trạng thái rỗng, lỗi API/DB và biên phân trang.", "Bắt buộc"],
        ],
        [3.0, 10.7, 2.5],
    )

    document.add_heading("10. Yêu cầu phi chức năng, bảo mật và vận hành", level=1)
    add_bullets(
        document,
        [
            "Web yêu cầu Forms Authentication và người dùng được cấp Menu_ID tương ứng; cần bổ sung kiểm tra quyền riêng cho thao tác chạy/ghi đối chiếu.",
            "API và Oracle phải dùng bind parameters/allow-list nguồn; dữ liệu, thông báo và nội dung Excel phải được encode.",
            "Tài khoản import chỉ được cấp quyền tối thiểu; quyền CREATE/ALTER schema tách khỏi tài khoản vận hành thường xuyên.",
            "Không ghi password vào log/file; cấu hình môi trường phải được bảo vệ bằng ACL; kết nối production cần chính sách TLS/wallet/timeout phù hợp.",
            "Bản release tool phải có version resource, chữ ký số, hash manifest, antivirus scan và quy trình build tái lập; chốt dependency oracledb thống nhất.",
            "Tra cứu dùng phân trang server-side; export lớn cần giới hạn tổng dòng hoặc xử lý bất đồng bộ để tránh giữ request/bộ nhớ quá lâu.",
            "Giao diện responsive, bảng cuộn, sticky header/cột; lỗi API/DB không được làm treo trang hoặc để giao dịch dở dang.",
            "Cần giám sát số file/dòng import, lỗi theo run, độ mới dữ liệu T_KHH/T_SLOT_AERO, thời gian đối chiếu và số lượng KQ bất thường.",
        ],
    )

    document.add_heading("11. Tiêu chí nghiệm thu", level=1)
    add_bullets(
        document,
        [
            "Tool chỉ nhận đúng thư mục KHH/SLOT CHK và file .xlsx/.xlsm, bỏ file ~$; file sai phải báo rõ và không ghi dở dữ liệu.",
            "Import lại cùng bộ dữ liệu không phát sinh dòng mới ngoài các duplicate nội-file đã được nghiệp vụ chấp thuận; tổng kết GUI khớp Oracle.",
            "Lỗi một file rollback file đó và cho phép tiếp tục file sau; lỗi kết nối/schema dừng an toàn; hủy không làm hỏng transaction.",
            "Dữ liệu vừa import hiển thị đúng trên SlotAero/KHH với bộ lọc ngày, sân bay, keyword, page size và STT chính xác.",
            "KQ1–KQ4 đúng quy tắc; chạy lại thay thế đúng phạm vi; lỗi ghi rollback toàn bộ và lưu đúng người thực hiện.",
            "SlotResults chuyển KQ/phân trang mà không ghi dữ liệu; tổng hợp khớp API và Excel đủ toàn bộ dòng/tám cột, đúng tên tệp/tiếng Việt.",
            "Người chưa đăng nhập hoặc không có quyền menu không truy cập được; dữ liệu đặc biệt không tạo HTML/script; bí mật không xuất hiện trong log.",
            "Bốn trang web hiển thị ổn định trên desktop/mobile; API lỗi, dữ liệu rỗng và cấu hình thiếu đều có thông báo rõ ràng.",
        ],
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--slot-aero", type=Path, required=True)
    parser.add_argument("--khh", type=Path, required=True)
    parser.add_argument("--comparison", type=Path, required=True)
    parser.add_argument("--results", type=Path, required=True)
    args = parser.parse_args()
    paths = [args.template, args.slot_aero, args.khh, args.comparison, args.results]
    for path in paths:
        if not path.is_file():
            raise FileNotFoundError(path)
    build(
        args.template,
        args.output,
        {"slot_aero": args.slot_aero, "khh": args.khh, "comparison": args.comparison, "results": args.results},
    )


if __name__ == "__main__":
    main()
