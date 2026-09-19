from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT = Path(__file__).resolve().parents[1] / "TaiLieu" / "HoSoTrienKhai_HTQLSLB" / "Mau_05_BBNT_CongTacKhaoSat_HTQLSLB.docx"
PROJECT = "Cập nhật, hiệu chỉnh cơ sở dữ liệu, phần mềm Hệ thống số liệu điều hành bay"
PACKAGE = "Gói thầu số 1: Cập nhật, hiệu chỉnh cơ sở dữ liệu, phần mềm Hệ thống số liệu điều hành bay"


def shade(cell, value="D9EAF7"):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), value)


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(5)
    for name, size in [("Heading 1", 14), ("Heading 2", 12)]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True


def paragraph(doc, text="", bold=False, align=None, italic=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    return p


def table(doc, headers, rows):
    result = doc.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    for index, title in enumerate(headers):
        cell = result.rows[0].cells[index]
        cell.text = title
        shade(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = result.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()
    return result


def signatures(doc):
    table(doc, ["ĐẠI DIỆN CHỦ ĐẦU TƯ/ĐƠN VỊ SỬ DỤNG", "ĐẠI DIỆN NHÀ THẦU TRIỂN KHAI", "ĐẠI DIỆN CÁC BÊN LIÊN QUAN (NẾU CÓ)"], [
        ("\n\n\n\n(Ký, ghi rõ họ tên, chức vụ)", "\n\n\n\n(Ký, ghi rõ họ tên, chức vụ, đóng dấu)", "\n\n\n\n(Ký, ghi rõ họ tên, chức vụ)"),
    ])


def main():
    doc = Document()
    configure(doc)
    paragraph(doc, "Mẫu số 5 – Đề xuất áp dụng cho nghiệm thu công tác khảo sát", True, WD_ALIGN_PARAGRAPH.RIGHT)
    paragraph(doc, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", True, WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, "Độc lập - Tự do - Hạnh phúc", True, WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, "________________", align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, "………, ngày …… tháng …… năm ………", align=WD_ALIGN_PARAGRAPH.RIGHT)
    paragraph(doc, "BIÊN BẢN NGHIỆM THU CÔNG TÁC KHẢO SÁT", True, WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, "HIỆN TRẠNG, YÊU CẦU NGHIỆP VỤ VÀ ĐIỀU KIỆN TRIỂN KHAI", True, WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, f"DỰ ÁN: {PROJECT}", True)
    paragraph(doc, f"GÓI THẦU: {PACKAGE}", True)
    paragraph(doc, "Ghi chú: Các trường đặt trong dấu […] hoặc bằng dòng chấm được điền, đối chiếu và xác nhận bằng hồ sơ thực tế trước khi ký.", italic=True)

    doc.add_heading("I. Đối tượng nghiệm thu", level=1)
    paragraph(doc, "Công tác khảo sát hiện trạng, quy trình nghiệp vụ, dữ liệu, phần mềm, hạ tầng, tích hợp, tài khoản/phân quyền và điều kiện vận hành của Hệ thống số liệu điều hành bay; kết quả khảo sát là đầu vào để lập tài liệu yêu cầu phần mềm, thiết kế kỹ thuật, kế hoạch triển khai và kế hoạch kiểm thử.")
    table(doc, ["Nội dung", "Phạm vi cần đối chiếu/xác nhận"], [
        ("Phạm vi nghiệp vụ", "05 nhóm chức năng/53 chức năng theo danh mục phạm vi được phê duyệt; quy trình hiện tại, tác nhân, trạng thái, ngoại lệ và tiêu chí chấp nhận."),
        ("Dữ liệu", "Nguồn, cấu trúc, chất lượng, tần suất, mã định danh, ngày nghiệp vụ/timezone, quyền truy cập, lưu trữ và đối soát."),
        ("Tích hợp", "Email/File, ADS-B, SLOT, AMHS/AFTN, Bravo, API/Gateway; endpoint, định dạng, xác thực, lỗi, retry, log và trách nhiệm đầu mối."),
        ("Hạ tầng/vận hành", "Máy chủ, IIS/runtime, Oracle/package, tài khoản, mạng, backup/restore, log, giám sát, bảo mật và phương án rollback."),
    ])

    doc.add_heading("II. Thành phần trực tiếp nghiệm thu", level=1)
    table(doc, ["STT", "Thành phần", "Ông/Bà", "Chức vụ", "Đơn vị"], [
        ("1", "Đại diện chủ đầu tư", "………………………………", "……………………", "……………………"),
        ("2", "Đại diện đơn vị quản lý, sử dụng", "………………………………", "……………………", "……………………"),
        ("3", "Đại diện đơn vị giám sát công tác triển khai (nếu có)", "………………………………", "……………………", "……………………"),
        ("4", "Đại diện tổ chức/cá nhân thiết kế chi tiết (nếu có)", "………………………………", "……………………", "……………………"),
        ("5", "Đại diện nhà thầu triển khai", "………………………………", "……………………", "……………………"),
        ("6", "Đại diện đơn vị/bên liên quan", "………………………………", "……………………", "……………………"),
    ])

    doc.add_heading("III. Thời gian, địa điểm và căn cứ nghiệm thu", level=1)
    table(doc, ["Nội dung", "Thông tin"], [
        ("Thời gian khảo sát", "Từ ngày ……/……/……… đến ngày ……/……/………"),
        ("Thời gian nghiệm thu", "…… giờ …… ngày ……/……/………"),
        ("Địa điểm/hình thức", "…………………………………………………………………………………………………………"),
        ("Hợp đồng/văn bản giao nhiệm vụ", "Số: …………………………………… ngày: ………………………"),
        ("Kế hoạch khảo sát", "Mã/số: …………………………………… phiên bản: ………………………"),
        ("Tài liệu đầu ra", "Báo cáo khảo sát; biên bản/phỏng vấn; danh mục phạm vi; danh mục dữ liệu–tích hợp; danh sách yêu cầu làm rõ."),
    ])

    doc.add_heading("IV. Đánh giá công tác khảo sát", level=1)
    doc.add_heading("4.1. Phương pháp và bằng chứng khảo sát", level=2)
    table(doc, ["Phương pháp", "Kết quả/bằng chứng cần kèm"], [
        ("Phỏng vấn/xác nhận nghiệp vụ", "Danh sách người tham gia, biên bản/ý kiến xác nhận, quy trình hiện trạng và vấn đề cần làm rõ."),
        ("Quan sát/thao tác trên hệ thống", "Ảnh màn hình có ngày giờ, URL/môi trường, tài khoản/role đã che thông tin nhạy cảm và kịch bản thao tác."),
        ("Rà soát tài liệu/dữ liệu mẫu", "Mẫu file/điện văn/API, mapping trường, báo cáo, biểu mẫu, package/schema và quy ước mã dữ liệu."),
        ("Kiểm tra kỹ thuật", "Thông tin môi trường, kết nối, endpoint, log, phiên bản, quyền truy cập, backup/restore và các giới hạn đã xác nhận."),
    ])

    doc.add_heading("4.2. Kết quả khảo sát theo nhóm chức năng", level=2)
    table(doc, ["STT", "Nhóm chức năng", "Số chức năng", "Kết quả khảo sát cần xác nhận", "Kết luận"], [
        ("I", "Tích hợp và tự động hóa dữ liệu", "05", "Nguồn/đích, mapping, parser, idempotency, retry/quarantine, đối soát và audit.", "Đạt/Đạt có điều kiện/Không đạt"),
        ("II", "Cảnh báo và tiện ích hỗ trợ", "02", "Kịch bản/mức độ cảnh báo, phạm vi nhận, Live Fire, Daily Statistic, KHB quân sự, duyệt/phát điện văn.", "Đạt/Đạt có điều kiện/Không đạt"),
        ("III", "Nâng cấp và tối ưu HTSLB", "32", "Tra cứu, điện văn, ngày nghiệp vụ, cảnh báo, KHB, export, Bravo, hiệu năng, lịch sử và đối soát.", "Đạt/Đạt có điều kiện/Không đạt"),
        ("IV", "Phân tích, báo cáo thông minh", "10", "Nguồn dữ liệu, KPI/công thức, bộ lọc, biểu đồ, drill-down, export, dữ liệu thiếu/trễ.", "Đạt/Đạt có điều kiện/Không đạt"),
        ("V", "AI hỗ trợ thống kê, tìm kiếm và tổng hợp", "04", "Intent/entity, SQL Guard, Oracle read-only, scope, fallback, feedback, giám sát hiệu năng/chất lượng.", "Đạt/Đạt có điều kiện/Không đạt"),
        ("Tổng", "05 nhóm", "53", "Đối chiếu với báo cáo khảo sát chi tiết và danh mục phạm vi.", "……………………"),
    ])

    doc.add_heading("4.3. Kết quả khảo sát dữ liệu, tích hợp và hạ tầng", level=2)
    table(doc, ["Nội dung", "Kết quả/đánh giá", "Tồn tại hoặc yêu cầu làm rõ"], [
        ("Dữ liệu nghiệp vụ", "Đã/Chưa xác định nguồn chuẩn, trường khóa, chất lượng, lịch cập nhật và dữ liệu mẫu.", "………………………………………………………………"),
        ("Email/File/ADS-B/SLOT", "Đã/Chưa xác định định dạng, mapping, phương thức nhận, tần suất, lỗi và đối soát.", "………………………………………………………………"),
        ("AMHS/AFTN/Bravo/API", "Đã/Chưa xác định endpoint, xác thực, schema, mã lỗi, ACK, retry, idempotency và log.", "………………………………………………………………"),
        ("Hạ tầng/Oracle/IIS", "Đã/Chưa xác định môi trường, component, account, package, backup/restore, monitoring và rollback.", "………………………………………………………………"),
        ("Bảo mật/phân quyền", "Đã/Chưa xác định RBAC/data scope, tài khoản kỹ thuật, retention, masking, audit và quy trình cấp quyền.", "………………………………………………………………"),
    ])

    doc.add_heading("4.4. Danh mục hồ sơ, sản phẩm và bằng chứng", level=2)
    table(doc, ["STT", "Hồ sơ/bằng chứng", "Mã/số/phiên bản", "Tình trạng"], [
        ("1", "Báo cáo kết quả khảo sát", "………………………………", "Đính kèm/Chưa có"),
        ("2", "Biên bản khảo sát, phỏng vấn hoặc xác nhận nghiệp vụ", "………………………………", "Đính kèm/Chưa có"),
        ("3", "Danh mục 05 nhóm/53 chức năng và phiếu khảo sát chi tiết", "………………………………", "Đính kèm/Chưa có"),
        ("4", "Danh mục dữ liệu, tích hợp, API/package và mapping", "………………………………", "Đính kèm/Chưa có"),
        ("5", "Ảnh màn hình, mẫu dữ liệu, log và sơ đồ kiến trúc/ngữ cảnh", "………………………………", "Đính kèm/Chưa có"),
        ("6", "Danh sách vấn đề/yêu cầu làm rõ và kế hoạch xử lý", "………………………………", "Đính kèm/Chưa có"),
    ])

    doc.add_heading("4.5. Tồn tại, sai khác và biện pháp xử lý", level=2)
    table(doc, ["STT", "Tồn tại/sai khác", "Ảnh hưởng", "Biện pháp/đầu mối", "Hạn xử lý", "Trạng thái"], [
        ("1", "", "", "", "", "Mở/Đóng"),
        ("2", "", "", "", "", "Mở/Đóng"),
        ("3", "", "", "", "", "Mở/Đóng"),
    ])

    doc.add_heading("V. Kết luận", level=1)
    paragraph(doc, "Căn cứ hồ sơ, bằng chứng và kết quả đánh giá nêu trên, các bên thống nhất nghiệm thu công tác khảo sát ở mức: ☐ Đạt; ☐ Đạt có điều kiện; ☐ Không đạt.")
    paragraph(doc, "Nếu đạt có điều kiện, các nội dung chưa hoàn tất phải được lập thành danh sách yêu cầu làm rõ, xác định đầu mối và hạn xử lý; chỉ được chuyển sang phê duyệt SRS/SDD, kế hoạch triển khai hoặc kế hoạch kiểm thử sau khi các điều kiện liên quan được xử lý/xác nhận theo thẩm quyền.")
    paragraph(doc, "Ý kiến khác của các bên (nếu có): ................................................................................................................................................................................")
    signatures(doc)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
