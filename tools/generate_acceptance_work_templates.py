from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT_DIR = Path(__file__).resolve().parents[1] / "TaiLieu" / "HoSoTrienKhai_HTQLSLB"
PROJECT = "Cập nhật, hiệu chỉnh cơ sở dữ liệu, phần mềm Hệ thống số liệu điều hành bay"
PACKAGE = "Gói thầu số 1: Cập nhật, hiệu chỉnh cơ sở dữ liệu, phần mềm Hệ thống số liệu điều hành bay"


def shade(cell, value="D9EAF7"):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), value)


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(5)
    for name, size in [("Heading 1", 14), ("Heading 2", 12)]:
        heading = doc.styles[name]
        heading.font.name = "Times New Roman"
        heading._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        heading.font.size = Pt(size)
        heading.font.bold = True


def paragraph(doc, text="", bold=False, align=None, italic=False):
    item = doc.add_paragraph()
    if align is not None:
        item.alignment = align
    run = item.add_run(text)
    run.bold = bold
    run.italic = italic
    return item


def table(doc, headers, rows):
    result = doc.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    for index, value in enumerate(headers):
        cell = result.rows[0].cells[index]
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = result.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()
    return result


def participants(doc):
    table(doc, ["STT", "Thành phần", "Ông/Bà", "Chức vụ", "Đơn vị"], [
        ("1", "Đại diện chủ đầu tư", "………………………………", "……………………", "……………………"),
        ("2", "Đại diện đơn vị quản lý, sử dụng", "………………………………", "……………………", "……………………"),
        ("3", "Đại diện đơn vị giám sát (nếu có)", "………………………………", "……………………", "……………………"),
        ("4", "Đại diện nhà thầu triển khai", "………………………………", "……………………", "……………………"),
        ("5", "Đại diện đơn vị/bên liên quan", "………………………………", "……………………", "……………………"),
    ])


def signatures(doc):
    table(doc, ["ĐẠI DIỆN CHỦ ĐẦU TƯ/ĐƠN VỊ SỬ DỤNG", "ĐẠI DIỆN NHÀ THẦU TRIỂN KHAI", "ĐẠI DIỆN CÁC BÊN LIÊN QUAN (NẾU CÓ)"], [
        ("\n\n\n\n(Ký, ghi rõ họ tên, chức vụ)", "\n\n\n\n(Ký, ghi rõ họ tên, chức vụ, đóng dấu)", "\n\n\n\n(Ký, ghi rõ họ tên, chức vụ)"),
    ])


def header(doc, title, subtitle):
    paragraph(doc, "Mẫu số 5 – Đề xuất áp dụng", True, WD_ALIGN_PARAGRAPH.RIGHT)
    paragraph(doc, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", True, WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, "Độc lập - Tự do - Hạnh phúc", True, WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, "________________", align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, "………, ngày …… tháng …… năm ………", align=WD_ALIGN_PARAGRAPH.RIGHT)
    paragraph(doc, title, True, WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, subtitle, True, WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, f"DỰ ÁN: {PROJECT}", True)
    paragraph(doc, f"GÓI THẦU: {PACKAGE}", True)
    paragraph(doc, "Ghi chú: Điền thông tin và đính kèm bằng chứng thực tế trước khi ký; không dùng thông tin mẫu thay cho số liệu, phiên bản hoặc chữ ký chính thức.", italic=True)


def common_start(doc, object_text):
    doc.add_heading("I. Đối tượng nghiệm thu", level=1)
    paragraph(doc, object_text)
    doc.add_heading("II. Thành phần trực tiếp nghiệm thu", level=1)
    participants(doc)
    doc.add_heading("III. Thời gian, địa điểm và căn cứ nghiệm thu", level=1)
    table(doc, ["Nội dung", "Thông tin"], [
        ("Thời gian thực hiện", "Từ ngày ……/……/……… đến ngày ……/……/………"),
        ("Thời gian nghiệm thu", "…… giờ …… ngày ……/……/………"),
        ("Địa điểm/môi trường", "…………………………………………………………………………………………………………"),
        ("Hợp đồng/văn bản giao nhiệm vụ", "Số: …………………………………… ngày: ………………………"),
        ("Kế hoạch/thiết kế/tài liệu đối chiếu", "Mã/số: …………………………………… phiên bản: ………………………"),
    ])


def finish(doc, conclusion):
    doc.add_heading("V. Kết luận", level=1)
    paragraph(doc, "Căn cứ hồ sơ, bằng chứng và kết quả đánh giá nêu trên, các bên thống nhất nghiệm thu ở mức: ☐ Đạt; ☐ Đạt có điều kiện; ☐ Không đạt.")
    paragraph(doc, conclusion)
    paragraph(doc, "Ý kiến khác của các bên (nếu có): ................................................................................................................................................................................")
    signatures(doc)


def commercial_software():
    doc = Document()
    configure(doc)
    header(doc, "BIÊN BẢN NGHIỆM THU CÔNG TÁC CÀI ĐẶT", "PHẦN MỀM THƯƠNG MẠI")
    common_start(doc, "Công tác cài đặt, cấu hình, kích hoạt bản quyền và kiểm tra vận hành phần mềm thương mại/phần mềm nền được sử dụng cho Hệ thống số liệu điều hành bay. Đối tượng bao gồm phần mềm hệ điều hành, cơ sở dữ liệu, middleware, runtime, công cụ giám sát, bảo mật hoặc thành phần thương mại khác thuộc phạm vi hợp đồng.")
    doc.add_heading("IV. Đánh giá công tác đã thực hiện", level=1)
    table(doc, ["STT", "Phần mềm/thành phần", "Nhà cung cấp", "Phiên bản/bản quyền", "Máy chủ/vị trí", "Kết quả"], [
        ("1", "………………………………", "……………………", "License key/Subscription: ……………………", "……………………", "Đạt/Không đạt"),
        ("2", "………………………………", "……………………", "License key/Subscription: ……………………", "……………………", "Đạt/Không đạt"),
        ("3", "………………………………", "……………………", "License key/Subscription: ……………………", "……………………", "Đạt/Không đạt"),
    ])
    doc.add_heading("4.1. Tiêu chí kiểm tra cài đặt", level=2)
    table(doc, ["Nhóm kiểm tra", "Yêu cầu", "Bằng chứng"], [
        ("Nguồn gốc/bản quyền", "Phần mềm, license và quyền sử dụng hợp lệ, đúng phạm vi hợp đồng; không ghi lộ license key trong biên bản.", "Hóa đơn/chứng chỉ/license record đã che thông tin nhạy cảm."),
        ("Cài đặt/cấu hình", "Đúng phiên bản, checksum/gói cài đặt, hướng dẫn nhà sản xuất và cấu hình môi trường đã phê duyệt.", "Install log, ảnh màn hình, file cấu hình đã che secret."),
        ("Tương thích", "Tương thích với hệ điều hành, IIS/runtime, Oracle, mạng và thành phần ứng dụng liên quan.", "Kết quả health check, service status và test kết nối."),
        ("Bảo mật/vận hành", "Cập nhật bản vá, phân quyền tối thiểu, log, backup cấu hình, giám sát và quy trình gia hạn bản quyền.", "Phiếu kiểm tra/vulnerability record, backup ID, monitoring log."),
        ("Hoàn nguyên", "Có điểm khôi phục/gói uninstall hoặc rollback được phê duyệt, không làm ảnh hưởng dữ liệu nghiệp vụ.", "Rollback runbook và người phê duyệt."),
    ])
    doc.add_heading("4.2. Tồn tại và biện pháp xử lý", level=2)
    table(doc, ["STT", "Tồn tại/sai khác", "Ảnh hưởng", "Biện pháp/đầu mối", "Hạn", "Trạng thái"], [("1", "", "", "", "", "Mở/Đóng"), ("2", "", "", "", "", "Mở/Đóng")])
    finish(doc, "Nếu đạt có điều kiện, các vấn đề về bản quyền, tương thích, bản vá, kết nối hoặc backup/rollback phải được đóng trước khi đưa phần mềm vào môi trường chạy thử/vận hành.")
    return doc


def internal_software():
    doc = Document()
    configure(doc)
    header(doc, "BIÊN BẢN NGHIỆM THU CÔNG TÁC CÀI ĐẶT", "PHẦN MỀM NỘI BỘ")
    common_start(doc, "Công tác triển khai phần mềm nội bộ ATFM_WEB, các dịch vụ ứng dụng, worker tích hợp, package/script Oracle, cấu hình IIS/API Gateway và các thành phần hỗ trợ theo thiết kế kỹ thuật được phê duyệt.")
    doc.add_heading("IV. Đánh giá công tác đã thực hiện", level=1)
    table(doc, ["STT", "Thành phần", "Phiên bản/build", "Máy chủ/zone", "Nội dung kiểm tra", "Kết quả"], [
        ("1", "ATFM_WEB/IIS", "Build: …………… Checksum: ……………", "Application Zone", "Website, app pool, binding, config, health URL và quyền thư mục.", "Đạt/Không đạt"),
        ("2", "Oracle/schema/package", "Script: ……………", "Data Zone", "Schema/package VALID, migration, service account, audit, backup trước/sau triển khai.", "Đạt/Không đạt"),
        ("3", "Integration Workers", "Worker/job: ……………", "Application Zone", "Scheduler/queue/retry/quarantine, service status và log.", "Đạt/Không đạt"),
        ("4", "AI Service/API", "Service/model: ……………", "Application Zone", "Endpoint, SQL Guard, Oracle read-only, timeout, health và audit.", "Đạt/Không đạt/Không áp dụng"),
        ("5", "Gateway/Reverse Proxy", "Config: ……………", "DMZ", "HTTPS/certificate, route, header, access/error log và allowlist.", "Đạt/Không đạt"),
    ])
    doc.add_heading("4.1. Ma trận kết nối sau cài đặt", level=2)
    table(doc, ["Nguồn", "Đích", "Kiểm tra", "Kết quả"], [
        ("Người dùng/API Consumer", "Gateway/ATFM_WEB", "HTTPS, certificate, xác thực, phân quyền và URL hoạt động.", "Đạt/Không đạt"),
        ("ATFM_WEB/Worker", "Oracle ATFM/Audit", "Kết nối service account, package smoke test, audit log.", "Đạt/Không đạt"),
        ("Worker", "ADS-B/SLOT/AMHS-AFTN/Bravo", "Endpoint, schema/metadata, timeout, retry và trạng thái lỗi.", "Đạt/Không đạt/Chờ nguồn"),
        ("ATFM_WEB", "Email/File/Reports", "Đọc/ghi thử, export, quyền truy cập và retention.", "Đạt/Không đạt"),
    ])
    doc.add_heading("4.2. Bảo mật, sao lưu và hoàn nguyên", level=2)
    table(doc, ["Nội dung", "Kết quả/bằng chứng"], [
        ("RBAC/secret", "Tài khoản tách biệt, quyền tối thiểu; không đưa password/token vào biên bản hoặc log."),
        ("Audit/monitoring", "Đã kiểm tra log IIS/gateway/worker/Oracle, đồng bộ thời gian và cảnh báo dịch vụ."),
        ("Backup/rollback", "Có backup ID, gói build/script rollback, thứ tự hoàn nguyên và người phê duyệt."),
    ])
    finish(doc, "Chỉ chuyển sang chạy thử khi các thành phần cài đặt, kết nối thiết yếu, audit, backup và phương án rollback được xác nhận đạt hoặc các điều kiện tồn tại được phê duyệt bằng văn bản.")
    return doc


def trial_run():
    doc = Document()
    configure(doc)
    header(doc, "BIÊN BẢN NGHIỆM THU CÔNG TÁC CHẠY THỬ", "HỆ THỐNG SỐ LIỆU ĐIỀU HÀNH BAY")
    common_start(doc, "Công tác chạy thử hệ thống trên môi trường đã cài đặt, bao gồm kiểm thử chức năng, tích hợp, dữ liệu, phân quyền, báo cáo, cảnh báo, hiệu năng và các yêu cầu an toàn/lưu vết theo kế hoạch kiểm thử được phê duyệt.")
    doc.add_heading("IV. Đánh giá công tác đã thực hiện", level=1)
    table(doc, ["Nhóm kiểm thử", "Phạm vi", "Kết quả cần xác nhận", "Bằng chứng"], [
        ("Chức năng nghiệp vụ", "05 nhóm/53 chức năng theo SRS", "Kết quả test case, expected/actual, ảnh màn hình và người xác nhận.", "TC/ảnh/log"),
        ("Tích hợp/dữ liệu", "Email/File, ADS-B, SLOT, AMHS/AFTN, Bravo, API", "Kết nối, mapping, đối soát, retry, duplicate/quarantine và lỗi nguồn.", "Batch/reconcile/log"),
        ("Phân quyền/an toàn", "Role, data scope, session, audit, XSS/CSRF, secret", "Thao tác đúng quyền; từ chối thao tác sai quyền; audit đầy đủ.", "Security test/audit"),
        ("Hiệu năng/vận hành", "ExportBravo, Gen KHB, Inbox, báo cáo, AI và batch", "SLA, timeout, tải, retry, monitoring, backup/restore và rollback.", "APM/query/log"),
    ])
    doc.add_heading("4.1. Tổng hợp kết quả test", level=2)
    table(doc, ["Chỉ tiêu", "Số lượng", "Kết quả"], [
        ("Tổng test case lập kế hoạch", "……………", ""),
        ("Đạt", "……………", ""),
        ("Không đạt", "……………", ""),
        ("Chưa thực hiện/blocked", "……………", ""),
        ("Lỗi nghiêm trọng/cản trở", "……………", ""),
        ("Lỗi còn mở được chấp nhận", "……………", ""),
    ])
    doc.add_heading("4.2. Danh sách lỗi, tồn tại và kiểm thử lại", level=2)
    table(doc, ["Mã lỗi", "Mô tả", "Mức độ", "Cách xử lý", "Kết quả test lại", "Trạng thái"], [
        ("", "", "Critical/High/Medium/Low", "", "Đạt/Không đạt", "Mở/Đóng"),
        ("", "", "Critical/High/Medium/Low", "", "Đạt/Không đạt", "Mở/Đóng"),
        ("", "", "Critical/High/Medium/Low", "", "Đạt/Không đạt", "Mở/Đóng"),
    ])
    finish(doc, "Công tác chạy thử chỉ được xác nhận đạt khi không còn lỗi nghiêm trọng/cản trở; các lỗi còn mở phải được phân loại, đánh giá ảnh hưởng, phê duyệt chấp nhận và có kế hoạch khắc phục rõ ràng.")
    return doc


def training():
    doc = Document()
    configure(doc)
    header(doc, "BIÊN BẢN NGHIỆM THU CÔNG TÁC ĐÀO TẠO, HUẤN LUYỆN", "SỬ DỤNG VÀ VẬN HÀNH HỆ THỐNG")
    common_start(doc, "Công tác đào tạo, huấn luyện người sử dụng, người duyệt, cán bộ khai thác và quản trị viên về sử dụng, vận hành, giám sát và hỗ trợ Hệ thống số liệu điều hành bay.")
    doc.add_heading("IV. Đánh giá công tác đã thực hiện", level=1)
    table(doc, ["STT", "Chuyên đề", "Đối tượng", "Kết quả đào tạo"], [
        ("1", "Đăng nhập, phiên làm việc, phân quyền và an toàn thông tin", "Người dùng/Quản trị", "Đạt/Chưa đạt"),
        ("2", "Điện văn, cảnh báo, Live Fire, Daily Statistic và KHB quân sự", "Khai thác/Người duyệt", "Đạt/Chưa đạt"),
        ("3", "Tra cứu, phép bay/KHB, đồng bộ, export và đối soát Bravo", "Khai thác/Quản trị", "Đạt/Chưa đạt"),
        ("4", "Báo cáo, biểu đồ, xuất dữ liệu và AI hỗ trợ", "Khai thác/Lãnh đạo", "Đạt/Chưa đạt"),
        ("5", "Giám sát, log, sao lưu/khôi phục, xử lý sự cố và đầu mối hỗ trợ", "Quản trị/Hỗ trợ", "Đạt/Chưa đạt"),
    ])
    doc.add_heading("4.1. Danh sách học viên và xác nhận", level=2)
    table(doc, ["STT", "Họ và tên", "Đơn vị/vị trí", "Nhóm đào tạo", "Kết quả", "Chữ ký"], [
        (str(i), "………………………………", "………………………………", "……………………", "Đạt/Cần bổ sung", "") for i in range(1, 16)
    ])
    doc.add_heading("4.2. Hồ sơ và bằng chứng đào tạo", level=2)
    table(doc, ["Nội dung", "Bằng chứng cần kèm"], [
        ("Kế hoạch/chương trình đào tạo", "Mã tài liệu, lịch đào tạo, giảng viên, đối tượng và phiên bản phần mềm/tài liệu."),
        ("Tài liệu đào tạo", "Hướng dẫn sử dụng, hướng dẫn quản trị/vận hành, slide, bài thực hành và tài liệu bàn giao."),
        ("Kết quả thực hành/đánh giá", "Bài tập, câu hỏi, kết quả thao tác, ý kiến học viên và nội dung cần hỗ trợ thêm."),
        ("Hình ảnh/điểm danh", "Ảnh lớp học/phiên trực tuyến có thời gian; danh sách học viên ký trực tiếp hoặc xác nhận điện tử hợp lệ."),
    ])
    doc.add_heading("4.3. Nội dung cần bổ sung sau đào tạo", level=2)
    table(doc, ["STT", "Nội dung", "Học viên/đơn vị", "Hình thức hỗ trợ", "Hạn", "Trạng thái"], [("1", "", "", "", "", "Mở/Đóng"), ("2", "", "", "", "", "Mở/Đóng")])
    finish(doc, "Công tác đào tạo, huấn luyện được nghiệm thu khi đã thực hiện đủ chuyên đề theo kế hoạch, có danh sách học viên xác nhận, tài liệu bàn giao và kế hoạch hỗ trợ cho các nội dung cần bổ sung.")
    return doc


def save(document, name):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / name
    document.save(path)
    print(path)


def main():
    save(commercial_software(), "Mau_05_02_BBNT_CaiDatPhanMemThuongMai_HTQLSLB.docx")
    save(internal_software(), "Mau_05_03_BBNT_CaiDatPhanMemNoiBo_HTQLSLB.docx")
    save(trial_run(), "Mau_05_04_BBNT_ChayThu_HTQLSLB.docx")
    save(training(), "Mau_05_05_BBNT_DaoTaoHuanLuyen_HTQLSLB.docx")


if __name__ == "__main__":
    main()
