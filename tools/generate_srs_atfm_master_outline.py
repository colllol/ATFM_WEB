import os
import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = Path(os.environ.get(
    "SRS_ATFM_MASTER_OUTPUT",
    ROOT / "TaiLieu" / "SRS_ATFM_WEB_Khung_Suon.docx",
))
AEROSYNC_SOURCE = ROOT / "TaiLieu" / "SRS_VATM_AeroSync_Hien_Tai.docx"
INTEGRATION_SOURCES = [
    ("5.4", "FR-INT-002", "Thu thập và xử lý dữ liệu ADS-B cho khai thác O/F", "SRS_VATM_ADS-B.docx"),
    ("5.5", "FR-INT-003", "Tích hợp lọc SLOT vào hệ thống SLB", "SRS_VATM_SLOT.docx"),
    ("5.6", "FR-INT-004", "Kết nối trực tiếp AMHS để gửi/nhận điện văn", "SRS_ATFM_AMHS_004.docx"),
    ("5.7", "FR-INT-005", "Chuẩn hóa API chia sẻ dữ liệu", "SRS_ATFM_API_Gateway_005.docx"),
]
ALT002_SOURCE = ROOT / "TaiLieu" / "SRS_Live_Fire_Message.docx"
AI_NL2SQL_SOURCE = "SRS_NL2SQL_Vanna_Oracle_Hien_Tai.docx"


SECTIONS = [
    ("I", "PHÂN HỆ TÍCH HỢP VÀ TỰ ĐỘNG HÓA DỮ LIỆU", [
        "Tự động cập nhật dữ liệu từ email và thư mục lưu trữ",
        "Thu thập và xử lý dữ liệu ADS-B cho khai thác O/F",
        "Tích hợp lọc SLOT vào hệ thống SLB",
        "Kết nối trực tiếp AMHS để gửi/nhận điện văn",
        "Chuẩn hóa API chia sẻ dữ liệu",
    ]),
    ("II", "PHÂN HỆ CẢNH BÁO VÀ TIỆN ÍCH HỖ TRỢ", [
        "Hệ thống cảnh báo thông minh đa kịch bản",
        "Quản lý thông tin KHB quân sự và sử dụng vùng trời",
    ]),
    ("III", "PHÂN HỆ NÂNG CẤP VÀ TỐI ƯU CHỨC NĂNG HTSLB", [
        "Mở rộng thời gian lưu trữ điện văn INBOX",
        "Cảnh báo NOPERM cho điện văn không có trong KHBN",
        "Áp dụng quy tắc 120 giờ ICAO cho hiển thị FPL",
        "Tìm kiếm chuyến bay đa trường và theo từ khóa",
        "Tìm kiếm nâng cao phục vụ báo cáo",
        "Phân loại chuyến bay Quốc nội/Quốc tế tự động",
        "Export dữ liệu chia tách ba loại file ALL, LD, OF",
        "Chức năng Edit info chuyến bay",
        "Hiển thị lịch sử thao tác của người xử lý chuyến bay",
        "Đồng bộ dữ liệu điện văn liên ngày",
        "Thống kê theo khung giờ, chặng và đường bay",
        "Cảnh báo chuyến bay không có trong phép bay/KHB ngày",
        "Hỗ trợ hủy phép bay quốc tế/quốc nội",
        "Tối ưu hiệu năng ExportBravo dưới 3 phút",
        "Chức năng F8 ép dòng trong Calendar Accepted",
        "Tối ưu Gen KHB ngày hôm sau ra AFTN dưới 3 phút",
        "Mở rộng phạm vi tra cứu INBOX trên 7 ngày",
        "Tự động cập nhật khi sửa KHB đã build",
        "Cảnh báo chuyến bay cấp sai ngày bay so với thực tế",
        "Bổ sung trường Mục đích chuyến bay trong KHB ngày",
        "Logic cảnh báo đỏ chính xác cho chuyến bay hết hiệu lực",
        "Tự động cập nhật đường bay theo đoạn từ FPL thực tế",
        "Tự động nhận diện và chuyển ngày cho chuyến bay quốc tế hạ cánh hôm sau",
        "Theo dõi kế hoạch bay hằng ngày với cập nhật liên tục",
        "Thông báo thời điểm nhận số liệu bay Đi/Đến từ Trung tâm TBHĐB",
        "Cập nhật số liệu bay từ HTSLB sang Bravo",
        "Khắc phục lỗi nhầm ngày so với thực tế",
        "Sửa nhiều chuyến bay cùng nội dung trong một thao tác",
        "Đối chiếu số liệu bay giữa HTSLB và Bravo",
        "Mở rộng tìm kiếm trên giao diện Daily Military Report",
        "Thêm lựa chọn theo tất cả tiêu chí thống kê trong báo cáo",
        "Đánh giá, so sánh số liệu của KHBHĐBN",
    ]),
    ("IV", "PHÂN HỆ PHÂN TÍCH VÀ BÁO CÁO THÔNG MINH", [
        "Biểu đồ thông tin tổng quan khai thác bay",
        "Phân tích xu hướng khai thác theo thời gian",
        "Phát hiện và cảnh báo dữ liệu bất thường",
        "Tích hợp ADS-B cho báo cáo khai thác thực tế",
        "Tổng hợp chỉ số hiệu suất bay từ ADS-B O/F",
        "Biểu đồ thống kê trạng thái chuyến bay theo tỷ lệ phần trăm",
        "Biểu đồ so sánh hoạt động bay giữa các sân bay",
        "Báo cáo chuyến bay quân sự theo tiêu chí mở rộng",
        "Báo cáo tổng hợp hoạt động bay tại tất cả sân bay dân dụng",
        "Báo cáo cất, hạ cánh tại các sân bay toàn quốc",
    ]),
    ("V", "PHÂN HỆ AI HỖ TRỢ THỐNG KÊ, TÌM KIẾM VÀ TỔNG HỢP", [
        "Xử lý ngôn ngữ tự nhiên cho truy vấn hàng không",
        "Truy vấn và tổng hợp dữ liệu tự động theo yêu cầu",
        "Tích hợp Trợ lý ảo (Chatbot) vào giao diện HTSLB",
        "Báo cáo và giám sát hiệu năng của Trợ lý ảo",
    ]),
]


def font(style, size, bold=False):
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.bold = bold


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    item = OxmlElement("w:shd")
    item.set(qn("w:fill"), fill)
    props.append(item)


def repeat_header(row):
    props = row._tr.get_or_add_trPr()
    item = OxmlElement("w:tblHeader")
    item.set(qn("w:val"), "true")
    props.append(item)


def table(doc, headers, rows):
    result = doc.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    repeat_header(result.rows[0])
    for index, value in enumerate(headers):
        cell = result.rows[0].cells[index]
        cell.text = value
        shade(cell, "D9EAF7")
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for values in rows:
        cells = result.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()
    return result


def placeholder(doc, text="[Sẽ đặc tả tại bước tiếp theo]"):
    paragraph = doc.add_paragraph(text)
    paragraph.runs[0].italic = True
    paragraph.runs[0].font.color.rgb = None


def source_rows(source, table_index):
    source_table = source.tables[table_index - 1]
    return [[cell.text.strip() for cell in row.cells] for row in source_table.rows]


def source_table(doc, source, table_index):
    rows = source_rows(source, table_index)
    table(doc, rows[0], rows[1:])


def iter_blocks(document):
    parent = document.element.body
    for child in parent.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def clean_source_heading(text):
    return re.sub(r"^\d+(?:\.\d+)*\.\s*", "", text.strip())


def add_integration_source(doc, section_number, requirement_id, title, filename):
    path = ROOT / "TaiLieu" / filename
    if not path.exists():
        raise FileNotFoundError(f"Thiếu tài liệu nguồn {requirement_id}: {path}")
    source = Document(path)
    doc.add_heading(f"{section_number}. Đặc tả chi tiết {requirement_id} – {title}", level=2)
    table(doc, ["Thuộc tính", "Nội dung"], [
        ("Mã yêu cầu tổng thể", requirement_id),
        ("Tên yêu cầu", title),
        ("Tài liệu nguồn", filename),
        ("Trạng thái", "Đã tích hợp đặc tả hiện trạng"),
    ])

    skipped_cover_lines = 0
    source_heading_seen = False
    for block in iter_blocks(source):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            style = block.style.name if block.style else "Normal"
            if not source_heading_seen and not style.startswith("Heading"):
                skipped_cover_lines += 1
                if skipped_cover_lines <= 3:
                    continue
            if style.startswith("Heading"):
                source_heading_seen = True
                level_match = re.search(r"(\d+)$", style)
                source_level = int(level_match.group(1)) if level_match else 1
                target_level = 3 if source_level == 1 else 3
                doc.add_heading(clean_source_heading(text), level=target_level)
            elif style.startswith("List Bullet"):
                doc.add_paragraph(text, style="List Bullet")
            elif style.startswith("List Number"):
                doc.add_paragraph(text, style="List Number")
            elif style == "Caption":
                paragraph = doc.add_paragraph(text)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if paragraph.runs:
                    paragraph.runs[0].italic = True
            else:
                doc.add_paragraph(text)
        else:
            rows = [[cell.text.strip() for cell in row.cells] for row in block.rows]
            if rows:
                table(doc, rows[0], rows[1:])

    doc.add_heading(f"{section_number}.1. Liên kết kiểm thử và truy vết", level=3)
    table(doc, ["Yêu cầu tổng thể", "Yêu cầu chi tiết", "Nguồn bằng chứng"], [
        (requirement_id, "Các FR/BR/NFR trong đặc tả nguồn", "Test case, log, ảnh màn hình, API/DB và biên bản nghiệm thu"),
    ])


def add_ai_nl2sql_specification(doc):
    doc.add_heading("9.3. Ánh xạ yêu cầu Trợ lý AI NL2SQL/Vanna Oracle", level=2)
    doc.add_paragraph(
        "Đặc tả hiện trạng của Hệ thống Trợ lý AI ATFM được tích hợp vào phân hệ AI. "
        "Các yêu cầu chi tiết FR-NL, BR-NL và NFR-NL trong tài liệu nguồn là cơ sở "
        "thiết kế test case, thu thập bằng chứng và nghiệm thu cho các yêu cầu tổng thể dưới đây."
    )
    table(doc, ["Mã tổng thể", "Phạm vi ánh xạ", "Nhóm yêu cầu nguồn"], [
        ("FR-AI-001", "Tiếp nhận tiếng Việt/giọng nói, nhận biết ý định, bổ sung ngữ cảnh schema và sinh Oracle SQL", "FR-NL-01…17; BR-NL-01…12"),
        ("FR-AI-002", "Thực thi truy vấn đọc an toàn, chuẩn hóa dữ liệu, tổng hợp và trình bày bảng/biểu đồ", "FR-NL-09…23; BR-NL-03…15"),
        ("FR-AI-003", "Tích hợp giao diện chat, API thời gian thực, phiên hội thoại, phản hồi và huấn luyện có kiểm duyệt", "FR-NL-01…08; FR-NL-24…39"),
        ("FR-AI-004", "Báo cáo hoạt động, audit, giám sát hiệu năng, sức khỏe dịch vụ và vận hành", "FR-NL-40…44; NFR-NL-01…16"),
    ])
    add_integration_source(
        doc,
        "9.4",
        "FR-AI-001…FR-AI-004",
        "Trợ lý AI chuyển câu hỏi tiếng Việt thành Oracle SQL, thực thi và trình bày kết quả",
        AI_NL2SQL_SOURCE,
    )


def add_opt001_003_specification(doc):
    doc.add_heading("7.3. Đặc tả nhóm sửa đổi, bổ sung chức năng hiện hữu", level=2)
    doc.add_paragraph(
        "FR-OPT-001, FR-OPT-002 và FR-OPT-003 là yêu cầu nâng cấp trên luồng INBOX/FPL hiện hữu, "
        "không tạo phân hệ nghiệp vụ độc lập. Việc triển khai phải bảo toàn dữ liệu, URL, phân quyền "
        "và các thao tác đang sử dụng; thay đổi cơ sở dữ liệu phải có phương án sao lưu và rollback."
    )
    table(doc, ["Mã", "Loại thay đổi", "Chức năng hiện hữu bị tác động", "Kết quả mong đợi"], [
        ("FR-OPT-001", "Mở rộng chính sách lưu giữ", "Tiếp nhận, lưu trữ và khai thác điện văn INBOX", "Điện văn được giữ đủ thời hạn cấu hình và vẫn truy xuất được"),
        ("FR-OPT-002", "Bổ sung cảnh báo nghiệp vụ", "Phân tích/đối chiếu điện văn FPL với KHB ngày", "Nhận diện và cảnh báo NOPERM có giải thích"),
        ("FR-OPT-003", "Sửa quy tắc thời gian", "Ghép nối và hiển thị FPL", "Xử lý đúng cửa sổ 120 giờ, kể cả khác ngày"),
    ])

    doc.add_heading("7.4. FR-OPT-001 – Mở rộng thời gian lưu trữ điện văn INBOX", level=2)
    table(doc, ["Thuộc tính", "Nội dung"], [
        ("Mục tiêu", "Ngăn điện văn bị loại bỏ sớm và bảo đảm tra cứu, đối soát, kiểm tra sau khai thác."),
        ("Phạm vi sửa đổi", "Dữ liệu điện văn, tệp/log nguồn và bản ghi liên kết cần cho việc tái hiện nội dung INBOX."),
        ("Không bao gồm", "Không đồng nghĩa tự động mở toàn bộ khoảng ngày trên giao diện; phạm vi tìm kiếm được quản lý bởi yêu cầu tra cứu riêng."),
        ("Tham số cần phê duyệt", "INBOX_RETENTION_DAYS – số ngày lưu tối thiểu; giá trị production do đơn vị nghiệp vụ phê duyệt."),
    ])
    table(doc, ["Mã chi tiết", "Yêu cầu"], [
        ("FR-OPT-001.01", "Hệ thống phải lưu điện văn INBOX và dữ liệu liên quan tối thiểu bằng INBOX_RETENTION_DAYS tính từ thời điểm nhận."),
        ("FR-OPT-001.02", "Thời hạn lưu phải cấu hình được mà không sửa mã nguồn và chỉ người có quyền quản trị mới được thay đổi."),
        ("FR-OPT-001.03", "Tác vụ dọn dữ liệu không được xóa bản ghi còn trong thời hạn lưu hoặc làm mất liên kết giữa điện văn, nội dung và log nguồn."),
        ("FR-OPT-001.04", "Điện văn còn thời hạn phải xem được nội dung và thông tin nhận bằng chức năng INBOX hiện hữu theo đúng phân quyền."),
        ("FR-OPT-001.05", "Khi dọn dữ liệu hết hạn, hệ thống phải ghi nhật ký số lượng, khoảng thời gian, người/tác vụ thực hiện và kết quả."),
        ("FR-OPT-001.06", "Nâng cấp không được xóa dữ liệu lịch sử hiện có; trước thay đổi cấu trúc/chính sách phải có sao lưu và phương án rollback."),
    ])
    table(doc, ["Mã kiểm thử", "Tình huống", "Kết quả mong đợi"], [
        ("TC-OPT-001-01", "Điện văn có tuổi nhỏ hơn hoặc bằng thời hạn cấu hình", "Không bị dọn và xem được tại INBOX"),
        ("TC-OPT-001-02", "Điện văn quá hạn và chạy tác vụ dọn", "Được xử lý theo chính sách, có log đầy đủ, không tạo bản ghi mồ côi"),
        ("TC-OPT-001-03", "Thay đổi thời hạn bởi tài khoản không có quyền", "Bị từ chối và có dấu vết an toàn"),
        ("TC-OPT-001-04", "Nâng cấp trên dữ liệu lịch sử", "Không mất dữ liệu; có thể rollback theo phương án đã duyệt"),
    ])

    doc.add_heading("7.5. FR-OPT-002 – Cảnh báo NOPERM cho điện văn không có trong KHB ngày", level=2)
    table(doc, ["Thuộc tính", "Nội dung"], [
        ("Mục tiêu", "Cảnh báo người khai thác khi điện văn chuyến bay không tìm thấy KHB ngày/phép bay phù hợp."),
        ("Nguồn vào", "Điện văn đã phân tích và dữ liệu KHB ngày/phép bay có hiệu lực."),
        ("Kết quả", "Trạng thái NOPERM hiển thị nổi bật nhưng không tự thay thế quyết định nghiệp vụ của người khai thác."),
    ])
    table(doc, ["Mã chi tiết", "Yêu cầu"], [
        ("FR-OPT-002.01", "Sau khi phân tích điện văn, hệ thống phải đối chiếu chuyến bay với KHB ngày/phép bay theo số hiệu, ngày bay, sân bay đi/đến và điều kiện hiệu lực áp dụng."),
        ("FR-OPT-002.02", "Chỉ gắn NOPERM khi không có bản ghi phù hợp sau khi đã chuẩn hóa số hiệu, ngày giờ và mã sân bay."),
        ("FR-OPT-002.03", "Cảnh báo phải hiển thị rõ tại danh sách/chi tiết liên quan và cung cấp lý do không khớp để người dùng kiểm tra."),
        ("FR-OPT-002.04", "Hệ thống phải tính lại trạng thái khi điện văn, KHB ngày hoặc phép bay liên quan được thêm/sửa/duyệt."),
        ("FR-OPT-002.05", "Không được cảnh báo NOPERM nếu tồn tại ít nhất một bản ghi phù hợp và còn hiệu lực; không dùng dữ liệu nháp/chưa duyệt nếu quy trình không cho phép."),
        ("FR-OPT-002.06", "Việc xác nhận hoặc xử lý cảnh báo phải tuân theo phân quyền và được ghi nhật ký người dùng, thời gian, trạng thái trước/sau."),
        ("FR-OPT-002.07", "Lỗi dịch vụ đối chiếu hoặc thiếu dữ liệu bắt buộc phải hiển thị là 'Chưa xác định', không được kết luận NOPERM sai."),
    ])
    table(doc, ["Mã kiểm thử", "Tình huống", "Kết quả mong đợi"], [
        ("TC-OPT-002-01", "Điện văn khớp KHB ngày còn hiệu lực", "Không có cảnh báo NOPERM"),
        ("TC-OPT-002-02", "Không có KHB ngày/phép bay phù hợp", "Hiển thị NOPERM và lý do đối chiếu"),
        ("TC-OPT-002-03", "Bổ sung/duyệt KHB phù hợp sau khi đã cảnh báo", "Trạng thái được tính lại và bỏ cảnh báo"),
        ("TC-OPT-002-04", "Dịch vụ/dữ liệu đối chiếu lỗi", "Hiển thị Chưa xác định, không gắn NOPERM"),
    ])

    doc.add_heading("7.6. FR-OPT-003 – Áp dụng quy tắc 120 giờ ICAO cho hiển thị FPL", level=2)
    table(doc, ["Thuộc tính", "Nội dung"], [
        ("Mục tiêu", "Ghép nối và hiển thị FPL theo mốc thời gian khai thác thay vì chỉ theo ngày lịch, hỗ trợ FPL được nộp trước tối đa 120 giờ."),
        ("Mốc so sánh", "Thời điểm nhận/nộp FPL và EOBT hoặc thời điểm dự kiến khởi hành tương đương sau khi chuẩn hóa múi giờ."),
        ("Cửa sổ", "120 giờ; nên quản lý bằng tham số cấu hình nhưng giá trị nghiệp vụ mặc định là 120."),
    ])
    table(doc, ["Mã chi tiết", "Yêu cầu"], [
        ("FR-OPT-003.01", "Hệ thống phải cho phép ghép nối và hiển thị FPL khi EOBT nằm trong cửa sổ từ thời điểm nộp/nhận đến không quá 120 giờ sau đó."),
        ("FR-OPT-003.02", "Việc đối chiếu không được giới hạn bằng điều kiện cùng ngày; phải xử lý đúng trường hợp FPL và chuyến bay cách nhau nhiều ngày trong cửa sổ."),
        ("FR-OPT-003.03", "Các mốc thời gian phải được chuẩn hóa về cùng múi giờ trước khi tính chênh lệch và phải xử lý chính xác qua thời điểm đổi ngày/tháng/năm."),
        ("FR-OPT-003.04", "FPL vượt quá cửa sổ 120 giờ không được tự động gắn với chuyến bay; hệ thống phải ghi/hiển thị nguyên nhân loại trừ khi cần kiểm tra."),
        ("FR-OPT-003.05", "Tại đúng biên 120 giờ, FPL được coi là hợp lệ; lớn hơn 120 giờ là không hợp lệ theo quy tắc này."),
        ("FR-OPT-003.06", "Nếu thiếu hoặc sai định dạng mốc thời gian bắt buộc, hệ thống phải đưa về trạng thái Chưa xác định để xử lý thủ công, không tự suy đoán."),
        ("FR-OPT-003.07", "Thay đổi phải tương thích với các loại điện văn cập nhật/hủy liên quan và không làm nhân đôi bản ghi FPL đang hiển thị."),
    ])
    table(doc, ["Mã kiểm thử", "Tình huống", "Kết quả mong đợi"], [
        ("TC-OPT-003-01", "EOBT cách thời điểm nhận 119 giờ 59 phút", "FPL được ghép nối/hiển thị"),
        ("TC-OPT-003-02", "EOBT cách đúng 120 giờ", "FPL được ghép nối/hiển thị"),
        ("TC-OPT-003-03", "EOBT cách 120 giờ 01 phút", "Không tự động ghép; có lý do loại trừ"),
        ("TC-OPT-003-04", "Khoảng thời gian đi qua ngày/tháng/năm", "Chênh lệch được tính đúng sau chuẩn hóa"),
        ("TC-OPT-003-05", "Thiếu EOBT hoặc thời điểm nhận", "Trạng thái Chưa xác định, không ghép sai"),
    ])

    doc.add_heading("7.7. Quy tắc chung và truy vết kiểm định", level=2)
    table(doc, ["Mã", "Quy tắc"], [
        ("BR-OPT-001", "Mọi giá trị thời hạn/cửa sổ phải dùng thời gian máy chủ hoặc thời gian nghiệp vụ đã chuẩn hóa, không phụ thuộc đồng hồ trình duyệt."),
        ("BR-OPT-002", "Trạng thái cảnh báo là kết quả hỗ trợ khai thác; thao tác xác nhận/sửa dữ liệu phải theo quyền hiện hành."),
        ("BR-OPT-003", "Không được làm thay đổi nội dung điện văn gốc; dữ liệu chuẩn hóa và kết quả đối chiếu phải truy vết được về nguồn."),
        ("BR-OPT-004", "Các thay đổi phải có log, giám sát lỗi và kịch bản rollback trước khi triển khai production."),
    ])
    table(doc, ["Yêu cầu", "Thiết kế/thành phần", "Bằng chứng nghiệm thu"], [
        ("FR-OPT-001", "Kho điện văn, tác vụ retention, INBOX", "Cấu hình, log dọn dữ liệu, truy vấn DB, ảnh màn hình và TC-OPT-001-*"),
        ("FR-OPT-002", "Bộ phân tích/đối chiếu KHB ngày, giao diện cảnh báo", "Dữ liệu mẫu khớp/không khớp, log tính lại, ảnh và TC-OPT-002-*"),
        ("FR-OPT-003", "Bộ ghép FPL và xử lý thời gian", "Bộ dữ liệu biên 120 giờ, log đối chiếu, ảnh và TC-OPT-003-*"),
    ])


def add_opt006_010_specification(doc):
    doc.add_heading("7.8. FR-OPT-006 – Phân loại chuyến bay Quốc nội/Quốc tế tự động", level=2)
    table(doc, ["Thuộc tính", "Nội dung"], [
        ("Loại thay đổi", "Bổ sung thuộc tính phân loại và tự động xác định trên dữ liệu chuyến bay hiện hữu."),
        ("Mục tiêu", "Thống nhất cách lọc, thống kê và xuất dữ liệu Quốc nội/Quốc tế, giảm phân loại thủ công."),
        ("Nguyên tắc", "Phân loại theo quốc gia của sân bay đi/đến sau khi chuẩn hóa mã sân bay; không suy đoán khi thiếu danh mục."),
    ])
    table(doc, ["Mã chi tiết", "Yêu cầu"], [
        ("FR-OPT-006.01", "Hệ thống phải xác định Quốc nội khi cả sân bay đi và sân bay đến thuộc Việt Nam; xác định Quốc tế khi ít nhất một đầu thuộc quốc gia khác."),
        ("FR-OPT-006.02", "Việc phân loại phải sử dụng danh mục sân bay/quốc gia dùng chung và hỗ trợ mã ICAO/IATA sau bước chuẩn hóa."),
        ("FR-OPT-006.03", "Thiếu hoặc không nhận diện được một đầu sân bay phải trả trạng thái Chưa xác định, không mặc định Quốc nội hoặc Quốc tế."),
        ("FR-OPT-006.04", "Trạng thái phải được tính lại khi sân bay đi/đến hoặc danh mục sân bay thay đổi và được dùng thống nhất khi hiển thị, tìm kiếm, thống kê, export."),
        ("FR-OPT-006.05", "Dữ liệu lịch sử cần được phân loại lại bằng tác vụ có kiểm soát; kết quả phải ghi số lượng theo từng trạng thái và lỗi."),
    ])
    table(doc, ["Mã kiểm thử", "Dữ liệu", "Kết quả"], [
        ("TC-OPT-006-01", "VVNB–VVTS", "Quốc nội"),
        ("TC-OPT-006-02", "VVNB–WSSS", "Quốc tế"),
        ("TC-OPT-006-03", "Sân bay không có trong danh mục", "Chưa xác định, có lý do"),
    ])

    doc.add_heading("7.9. FR-OPT-007 – Export dữ liệu tách ba loại file ALL, LD, OF", level=2)
    table(doc, ["Thuộc tính", "Nội dung"], [
        ("Loại thay đổi", "Mở rộng chức năng export hiện hữu để tạo bộ ba đầu ra độc lập."),
        ("Quy ước", "ALL: toàn bộ dữ liệu hợp lệ; LD: chuyến bay hạ cánh/đến; OF: chuyến bay bay qua."),
    ])
    table(doc, ["Mã chi tiết", "Yêu cầu"], [
        ("FR-OPT-007.01", "Người dùng có quyền export phải chọn được ngày/khoảng dữ liệu và tạo ba file ALL, LD, OF trong một lần thao tác."),
        ("FR-OPT-007.02", "ALL phải chứa hợp của các nhóm dữ liệu thuộc phạm vi export và không nhân đôi cùng một bản ghi nghiệp vụ."),
        ("FR-OPT-007.03", "LD và OF phải được tách theo quy tắc loại hình bay đã phê duyệt; một bản ghi không được xuất sai nhóm."),
        ("FR-OPT-007.04", "Ba file phải dùng cùng phiên bản dữ liệu, cùng thời điểm chốt, cấu trúc cột/định dạng ngày giờ thống nhất và tên file nhận diện được loại, ngày, lần tạo."),
        ("FR-OPT-007.05", "Hệ thống phải hiển thị số bản ghi từng file, cảnh báo file rỗng và chỉ thông báo thành công khi các file bắt buộc được tạo đầy đủ."),
        ("FR-OPT-007.06", "Mỗi lần export phải ghi người thực hiện, tham số, thời gian, số dòng, tên/checksum file và lỗi nếu có; việc tải file tuân theo phân quyền."),
    ])
    table(doc, ["Mã kiểm thử", "Tình huống", "Kết quả"], [
        ("TC-OPT-007-01", "Dữ liệu có cả LD và OF", "Sinh đủ ALL, LD, OF; số liệu đối soát đúng"),
        ("TC-OPT-007-02", "Một nhóm không có dữ liệu", "Cảnh báo rõ file rỗng, không làm sai hai file còn lại"),
        ("TC-OPT-007-03", "Lỗi khi tạo một file bắt buộc", "Không báo hoàn tất toàn bộ; log chỉ rõ file lỗi"),
    ])

    doc.add_heading("7.10. FR-OPT-008 – Chức năng Edit info chuyến bay", level=2)
    table(doc, ["Thuộc tính", "Nội dung"], [
        ("Loại thay đổi", "Bổ sung/sửa màn hình cập nhật thông tin trên bản ghi chuyến bay hiện hữu."),
        ("Mục tiêu", "Cho phép hiệu chỉnh dữ liệu sai hoặc thiếu mà vẫn kiểm soát quyền, tính hợp lệ và truy vết."),
    ])
    table(doc, ["Mã chi tiết", "Yêu cầu"], [
        ("FR-OPT-008.01", "Chỉ người có quyền Edit mới được mở chế độ sửa và lưu; trường chỉ đọc/khóa theo trạng thái nghiệp vụ không được thay đổi."),
        ("FR-OPT-008.02", "Hệ thống phải hiển thị dữ liệu mới nhất và kiểm tra bản ghi chưa bị người khác thay đổi trước khi lưu."),
        ("FR-OPT-008.03", "Các trường số hiệu, ngày bay, sân bay, giờ, loại tàu bay, đăng ký và trường nghiệp vụ liên quan phải được kiểm tra bắt buộc, định dạng, danh mục và quan hệ logic."),
        ("FR-OPT-008.04", "Khi thay đổi trường khóa đối chiếu, hệ thống phải yêu cầu xác nhận và tính lại các dữ liệu dẫn xuất như phân loại, cảnh báo, ghép điện văn."),
        ("FR-OPT-008.05", "Lưu phải có tính nguyên tử; nếu một bước kiểm tra/cập nhật thất bại thì không được lưu một phần."),
        ("FR-OPT-008.06", "Mỗi thay đổi phải ghi giá trị trước/sau, người sửa, thời gian và lý do; nội dung điện văn nguồn không bị sửa trực tiếp."),
    ])
    table(doc, ["Mã kiểm thử", "Tình huống", "Kết quả"], [
        ("TC-OPT-008-01", "Sửa hợp lệ bởi người có quyền", "Lưu thành công và có lịch sử trước/sau"),
        ("TC-OPT-008-02", "Dữ liệu sai định dạng hoặc tài khoản không có quyền", "Từ chối, không thay đổi dữ liệu"),
        ("TC-OPT-008-03", "Bản ghi đã được người khác cập nhật", "Cảnh báo xung đột, yêu cầu tải lại"),
    ])

    doc.add_heading("7.11. FR-OPT-009 – Hiển thị lịch sử thao tác chuyến bay", level=2)
    table(doc, ["Thuộc tính", "Nội dung"], [
        ("Loại thay đổi", "Bổ sung giao diện khai thác nhật ký T_ACTIONHISTORY/lịch sử chuyên biệt theo chuyến bay."),
        ("Mục tiêu", "Truy vết đầy đủ ai đã làm gì, khi nào và dữ liệu thay đổi ra sao."),
    ])
    table(doc, ["Mã chi tiết", "Yêu cầu"], [
        ("FR-OPT-009.01", "Từ danh sách/chi tiết chuyến bay, người có quyền phải mở được lịch sử đúng bản ghi bằng định danh ổn định, không chỉ dựa vào số hiệu."),
        ("FR-OPT-009.02", "Lịch sử phải hiển thị thời gian, người dùng, thao tác, nguồn/màn hình, trạng thái kết quả, giá trị trước/sau hoặc mô tả thay đổi."),
        ("FR-OPT-009.03", "Các thao tác thêm, sửa, đổi ngày, duyệt, hủy, export/finish và xử lý liên quan phải được ghi nhận theo phạm vi áp dụng."),
        ("FR-OPT-009.04", "Danh sách phải sắp xếp mới nhất trước, hỗ trợ phân trang/lọc và dùng thời gian máy chủ thống nhất."),
        ("FR-OPT-009.05", "Nhật ký là dữ liệu chỉ đọc đối với người khai thác, không cho sửa/xóa qua giao diện và phải bảo vệ thông tin nhạy cảm."),
        ("FR-OPT-009.06", "Lỗi ghi lịch sử phải được giám sát; thao tác nghiệp vụ quan trọng không được báo hoàn tất nếu chính sách bắt buộc audit chưa được đáp ứng."),
    ])
    table(doc, ["Mã kiểm thử", "Tình huống", "Kết quả"], [
        ("TC-OPT-009-01", "Thêm, sửa và đổi ngày chuyến bay", "Có đủ sự kiện đúng thứ tự và giá trị trước/sau"),
        ("TC-OPT-009-02", "Hai chuyến cùng số hiệu khác ngày", "Lịch sử không bị trộn"),
        ("TC-OPT-009-03", "Người dùng khai thác thử sửa/xóa log", "Không được phép"),
    ])

    doc.add_heading("7.12. FR-OPT-010 – Đồng bộ dữ liệu điện văn liên ngày", level=2)
    table(doc, ["Thuộc tính", "Nội dung"], [
        ("Loại thay đổi", "Sửa cơ chế liên kết điện văn và chuyến bay khi ngày điện văn khác ngày khai thác thực tế."),
        ("Mục tiêu", "Không bỏ sót hoặc nhân đôi FPL/CHG/DLA/CNL và các điện văn liên quan tại thời điểm qua ngày."),
    ])
    table(doc, ["Mã chi tiết", "Yêu cầu"], [
        ("FR-OPT-010.01", "Hệ thống phải xác định chuyến bay mục tiêu bằng số hiệu, sân bay, thời gian và quan hệ điện văn trong cửa sổ liên ngày được cấu hình."),
        ("FR-OPT-010.02", "Điện văn nhận trước/sau thời điểm đổi ngày phải được liên kết với đúng ngày bay thực tế sau khi chuẩn hóa UTC/giờ nghiệp vụ."),
        ("FR-OPT-010.03", "Điện văn cập nhật/hủy phải áp dụng đúng bản FPL gốc và đúng thứ tự thời gian; không cập nhật nhầm chuyến cùng số hiệu ở ngày liền kề."),
        ("FR-OPT-010.04", "Đồng bộ phải idempotent: xử lý lại cùng điện văn không tạo bản ghi hoặc hiệu ứng trùng; phải dùng định danh/hash nguồn để chống trùng."),
        ("FR-OPT-010.05", "Khi dữ liệu bị chuyển ngày, hệ thống phải cập nhật nhất quán bản ghi liên quan, giữ ngày cũ để truy vết và tính lại cảnh báo/phân loại cần thiết."),
        ("FR-OPT-010.06", "Trường hợp có nhiều ứng viên hoặc thiếu dữ liệu phải đưa vào Chưa xác định/xử lý thủ công, không tự ghép có rủi ro."),
        ("FR-OPT-010.07", "Mỗi lượt đồng bộ phải ghi điện văn nguồn, bản ghi đích, trạng thái trước/sau, quy tắc ghép, thời gian và lỗi; hỗ trợ chạy bù có kiểm soát."),
    ])
    table(doc, ["Mã kiểm thử", "Tình huống", "Kết quả"], [
        ("TC-OPT-010-01", "FPL nhận ngày D, chuyến bay ngày D+1 trong cửa sổ", "Liên kết đúng D+1, không nhân đôi"),
        ("TC-OPT-010-02", "CHG/DLA/CNL qua ngày", "Áp dụng đúng FPL gốc theo thứ tự"),
        ("TC-OPT-010-03", "Xử lý lại cùng điện văn", "Kết quả không đổi, không sinh bản ghi trùng"),
        ("TC-OPT-010-04", "Hai ứng viên cùng số hiệu ở ngày liền kề", "Không tự ghép; chuyển xử lý thủ công"),
    ])

    doc.add_heading("7.13. Ma trận truy vết FR-OPT-006…FR-OPT-010", level=2)
    table(doc, ["Yêu cầu", "Thành phần chịu tác động", "Bằng chứng kiểm định"], [
        ("FR-OPT-006", "Danh mục sân bay, dữ liệu chuyến bay, bộ lọc/báo cáo", "Bộ dữ liệu nội địa/quốc tế/không xác định và TC-OPT-006-*"),
        ("FR-OPT-007", "Luồng export và kho tệp", "Ba file mẫu, đối soát số dòng/checksum, log và TC-OPT-007-*"),
        ("FR-OPT-008", "Màn hình/DAL/API cập nhật chuyến bay", "Ảnh màn hình, DB trước/sau, phân quyền và TC-OPT-008-*"),
        ("FR-OPT-009", "T_ACTIONHISTORY và giao diện lịch sử", "Log theo định danh chuyến bay, phân quyền và TC-OPT-009-*"),
        ("FR-OPT-010", "Bộ phân tích/đồng bộ điện văn và dữ liệu chuyến bay", "Dữ liệu liên ngày, log ghép/chống trùng và TC-OPT-010-*"),
    ])


def add_alt002_specification(doc):
    if not ALT002_SOURCE.exists():
        raise FileNotFoundError(f"Thiếu tài liệu nguồn FR-ALT-002: {ALT002_SOURCE}")
    source = Document(ALT002_SOURCE)
    doc.add_heading(
        "6.4. Đặc tả chi tiết FR-ALT-002 – Quản lý thông tin KHB quân sự và sử dụng vùng trời",
        level=2,
    )
    table(doc, ["Thuộc tính", "Nội dung"], [
        ("Mã yêu cầu tổng thể", "FR-ALT-002"),
        ("Tên yêu cầu", "Quản lý thông tin KHB quân sự và sử dụng vùng trời"),
        ("Tài liệu nguồn", "SRS_Live_Fire_Message.docx"),
        ("Phạm vi quy trình", "Live Fire Message; Daily Statistic; KHB quân sự; QS Message; báo cáo khai thác"),
        ("Trạng thái", "Đã tích hợp đặc tả hiện trạng"),
    ])
    doc.add_paragraph(
        "Đặc tả này hợp nhất các quy trình nghiệp vụ có liên quan đến hoạt động bay quân sự "
        "và sử dụng vùng trời: lập/phê duyệt điện văn bắn đạn thật, thống kê chuyến bay hoàn "
        "thành, quản lý KHB quân sự, tạo QS Message, phát điện văn và cung cấp dữ liệu chỉ xem "
        "cho người khai thác."
    )

    skipped_cover_lines = 0
    source_heading_seen = False
    for block in iter_blocks(source):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            style = block.style.name if block.style else "Normal"
            if not source_heading_seen and not style.startswith("Heading"):
                skipped_cover_lines += 1
                if skipped_cover_lines <= 3:
                    continue
            if style.startswith("Heading"):
                source_heading_seen = True
                doc.add_heading(clean_source_heading(text), level=3)
            elif style.startswith("List Bullet"):
                doc.add_paragraph(text, style="List Bullet")
            elif style.startswith("List Number"):
                doc.add_paragraph(text, style="List Number")
            else:
                doc.add_paragraph(text)
        else:
            rows = [[cell.text.strip() for cell in row.cells] for row in block.rows]
            if rows:
                table(doc, rows[0], rows[1:])

    doc.add_heading("6.4.1. Liên kết kiểm thử và truy vết FR-ALT-002", level=3)
    table(doc, ["Nhóm nghiệp vụ", "Yêu cầu/test case nguồn", "Bằng chứng dự kiến"], [
        ("Live Fire Message", "FR-LF, BR-LF, NFR-LF, TC-LF", "Ảnh màn hình, API/package log, dữ liệu trạng thái"),
        ("Daily Statistic", "FR-DS, BR-DS, NFR-DS, TC-DS", "Danh sách trước/sau Accepted, export và audit"),
        ("KHB quân sự", "FR-MIL, BR-MIL, NFR-MIL, TC-MIL", "KHB nguồn, QS Message, trạng thái phát và Military Report"),
    ])


def add_alt001_specification(doc):
    doc.add_heading(
        "6.3. Đặc tả chi tiết FR-ALT-001 – Hệ thống cảnh báo thông minh đa kịch bản",
        level=2,
    )
    table(doc, ["Thuộc tính", "Nội dung"], [
        ("Mã yêu cầu", "FR-ALT-001"),
        ("Tên chức năng", "Hệ thống cảnh báo/thông báo cập nhật liên tục"),
        ("Điểm truy cập chính", "/Common/ChartReport.aspx"),
        ("Thành phần giao diện", "Ô thông báo trên ATFM_New.Master; trang /SLOTS/Notifications.aspx"),
        ("API nội bộ", "/Handlers/Notification.ashx"),
        ("Dữ liệu", "T_NOTIFICATION; T_NOTIFICATION_TARGET; T_NOTIFICATION_READ"),
        ("Gói Oracle", "NOTIFICATION_PKG"),
        ("Chu kỳ cập nhật giao diện", "30 giây và cập nhật theo sự kiện"),
    ])
    doc.add_paragraph(
        "Chức năng cung cấp thông báo gần thời gian thực cho người dùng đã đăng nhập. Ô "
        "thông báo nằm trên master page nên xuất hiện tại ChartReport và các trang dùng "
        "ATFM_New.Master. Hệ thống hợp nhất thông báo nội bộ và thông báo đồng bộ từ báo cáo "
        "email, phân phối theo toàn hệ thống hoặc từng người dùng, đồng thời quản lý trạng thái "
        "đã đọc độc lập cho mỗi tài khoản."
    )

    doc.add_heading("6.3.1. Tác nhân và thành phần", level=3)
    table(doc, ["Tác nhân/thành phần", "Trách nhiệm"], [
        ("Người dùng ATFM", "Xem badge, mở danh sách, đánh dấu đã đọc và xem toàn bộ thông báo."),
        ("Hệ thống nghiệp vụ", "Tạo thông báo chung hoặc chỉ định danh sách người nhận."),
        ("API Email", "Cung cấp email và trạng thái attachment để chuyển thành thông báo."),
        ("ATFM_New.Master", "Tải/cập nhật danh sách, badge và xử lý thao tác đọc."),
        ("Notification.ashx", "Xác thực phiên, đồng bộ email, gọi package và trả JSON."),
        ("NOTIFICATION_PKG", "Lọc theo người dùng, phân trang, ghi trạng thái đọc và tạo thông báo."),
        ("Oracle", "Lưu nội dung, đối tượng nhận và trạng thái đọc theo người dùng."),
    ])

    doc.add_heading("6.3.2. Luồng xử lý", level=3)
    table(doc, ["Bước", "Xử lý"], [
        ("1", "Người dùng đăng nhập và mở ChartReport hoặc trang dùng ATFM_New.Master."),
        ("2", "Giao diện gọi GET Notification.ashx để lấy thông báo chưa đọc và tổng số chưa đọc."),
        ("3", "Handler thử đồng bộ thông báo email nếu lần thử gần nhất đã cách ít nhất 15 giây."),
        ("4", "NOTIFICATION_PKG.GET_STATE lọc thông báo chung/cá nhân và loại thông báo người dùng đã đọc."),
        ("5", "Giao diện hiển thị tối đa 50 thông báo gần nhất và badge tổng số chưa đọc."),
        ("6", "Mỗi 30 giây hệ thống tải lại nếu trang đang hiển thị; request chồng được xếp chờ một lượt."),
        ("7", "Người dùng có thể đánh dấu một thông báo hoặc toàn bộ thông báo là đã đọc."),
        ("8", "Trang Xem tất cả cung cấp bộ lọc trạng thái, phân trang 100 dòng và tải lại thủ công."),
    ])

    doc.add_heading("6.3.3. Trạng thái và dữ liệu", level=3)
    table(doc, ["Đối tượng", "Trạng thái/thuộc tính", "Ý nghĩa"], [
        ("Thông báo", "TARGET_TYPE = 0", "Thông báo chung cho mọi người dùng."),
        ("Thông báo", "TARGET_TYPE = 1", "Chỉ hiển thị cho người dùng có trong T_NOTIFICATION_TARGET."),
        ("Người dùng/thông báo", "Chưa có T_NOTIFICATION_READ", "Thông báo chưa đọc."),
        ("Người dùng/thông báo", "Có T_NOTIFICATION_READ", "Thông báo đã đọc bởi người dùng đó."),
        ("Nguồn email", "SOURCE_TYPE = EMAIL_API", "Thông báo được đồng bộ từ API Email."),
        ("Khóa nguồn email", "SOURCE_KEY SHA-256", "Chống tạo trùng theo Message-ID hoặc ID báo cáo."),
    ])
    table(doc, ["Trường", "Yêu cầu"], [
        ("TITLE", "Bắt buộc, NVARCHAR2, tối đa 250 ký tự."),
        ("CONTENT", "Bắt buộc, NVARCHAR2, tối đa 2.000 ký tự."),
        ("DATETIME", "Thời điểm phát sinh/nhận thông báo; dùng sắp xếp mới nhất trước."),
        ("TARGET_TYPE", "Chỉ nhận 0 hoặc 1."),
        ("SOURCE_TYPE/SOURCE_KEY", "Nhận diện nguồn và bảo đảm idempotency khi đồng bộ ngoài."),
        ("NOTIFICATION_ID/USER_ID", "Khóa phân phối cá nhân và trạng thái đọc."),
    ])

    doc.add_heading("6.3.4. Yêu cầu chức năng", level=3)
    table(doc, ["Mã", "Yêu cầu"], [
        ("FR-NOTI-01", "Chỉ người dùng có identity và ATFM_CURRENT_USER hợp lệ, khớp tên đăng nhập, mới được gọi API thông báo."),
        ("FR-NOTI-02", "Khi tải trang, hệ thống phải tự động lấy trạng thái thông báo mà không cần người dùng thao tác."),
        ("FR-NOTI-03", "Hệ thống phải tự cập nhật thông báo mỗi 30 giây và khi nhận sự kiện atfm:notifications-changed."),
        ("FR-NOTI-04", "Không khởi tạo request cập nhật mới khi request trước đang chạy; phải thực hiện một lượt chờ sau khi request hiện tại kết thúc."),
        ("FR-NOTI-05", "Khi tab/trang bị ẩn, lượt cập nhật định kỳ không bắt buộc tải dữ liệu; thao tác mở popup phải tải cưỡng bức."),
        ("FR-NOTI-06", "Badge phải hiển thị tổng số chưa đọc, hiển thị 99+ khi tổng lớn hơn 99 và ẩn trạng thái nhấn mạnh khi bằng 0."),
        ("FR-NOTI-07", "Popup phải hiển thị tối đa 50 thông báo chưa đọc gần nhất gồm tiêu đề, nội dung và thời điểm."),
        ("FR-NOTI-08", "Nội dung động phải được gán bằng textContent để không thực thi HTML/script từ dữ liệu."),
        ("FR-NOTI-09", "Người dùng được đánh dấu từng thông báo là đã đọc; thông báo phải được loại khỏi popup và giảm badge."),
        ("FR-NOTI-10", "Người dùng được đánh dấu tất cả thông báo của mình là đã đọc; không ảnh hưởng trạng thái của người dùng khác."),
        ("FR-NOTI-11", "Trang Xem tất cả phải lọc Tất cả/Chưa đọc/Đã đọc, phân trang 100 dòng và hiển thị tổng số."),
        ("FR-NOTI-12", "Hệ thống phải hỗ trợ tạo thông báo chung hoặc thông báo cho danh sách user ID hợp lệ."),
        ("FR-NOTI-13", "Thông báo cá nhân chỉ được trả cho người dùng thuộc danh sách đích."),
        ("FR-NOTI-14", "Khi GET state/list, handler phải thử đồng bộ nguồn Email API với thời gian tối thiểu 15 giây giữa hai lần thử trong cùng tiến trình."),
        ("FR-NOTI-15", "Các dòng attachment của cùng email phải được gộp theo Message-ID và chọn dòng có thông tin xử lý hữu ích nhất."),
        ("FR-NOTI-16", "Đồng bộ email phải MERGE theo SOURCE_TYPE/SOURCE_KEY để cập nhật nội dung mà không tạo bản ghi trùng."),
        ("FR-NOTI-17", "Thông báo email phải gồm người gửi, attachment, thời gian, trạng thái xử lý, trạng thái xác nhận và lỗi."),
        ("FR-NOTI-18", "Lỗi Email API không được làm gián đoạn việc trả các thông báo nội bộ đang có."),
        ("FR-NOTI-19", "API phải trả lỗi 401 cho phiên không hợp lệ, 403 cho POST không phải XMLHttpRequest, 405 cho phương thức sai và 400 cho tham số/thao tác sai."),
        ("FR-NOTI-20", "Khi tải lần đầu thất bại, giao diện phải báo Không thể tải thông báo; lỗi các lần sau không được xóa dữ liệu đã hiển thị thành công."),
    ])

    doc.add_heading("6.3.5. Quy tắc nghiệp vụ", level=3)
    table(doc, ["Mã", "Quy tắc"], [
        ("BR-NOTI-01", "Trạng thái đọc được quản lý riêng theo cặp thông báo–người dùng."),
        ("BR-NOTI-02", "Thông báo chung áp dụng cho mọi người; thông báo cá nhân phải có ít nhất một user ID hợp lệ."),
        ("BR-NOTI-03", "Người dùng không được đánh dấu đọc thông báo không thuộc phạm vi nhận của mình."),
        ("BR-NOTI-04", "TITLE và CONTENT là bắt buộc khi tạo thông báo."),
        ("BR-NOTI-05", "Danh sách và popup sắp xếp thông báo mới nhất trước theo DATETIME và ID."),
        ("BR-NOTI-06", "Một email chỉ sinh một thông báo logic; Message-ID được ưu tiên làm identity nguồn."),
        ("BR-NOTI-07", "Nếu email không có Message-ID, dùng ID báo cáo để tạo identity ổn định."),
        ("BR-NOTI-08", "Đánh dấu đã đọc không xóa thông báo gốc và không làm mất lịch sử của người dùng khác."),
    ])

    doc.add_heading("6.3.6. Yêu cầu phi chức năng", level=3)
    table(doc, ["Mã", "Yêu cầu"], [
        ("NFR-NOTI-01", "Chu kỳ polling mặc định 30 giây; đồng bộ Email API không thường xuyên hơn 15 giây mỗi tiến trình."),
        ("NFR-NOTI-02", "Email API timeout đọc/kết nối 5 giây; Oracle command timeout 10–15 giây tùy thao tác."),
        ("NFR-NOTI-03", "Request nền không được kích hoạt lớp loading toàn trang hoặc làm gián đoạn thao tác dashboard."),
        ("NFR-NOTI-04", "API phải đặt NoCache/NoStore để tránh trả trạng thái thông báo cũ."),
        ("NFR-NOTI-05", "Tất cả câu lệnh Oracle phải bind parameter; đồng bộ email phải dùng transaction và rollback khi lỗi."),
        ("NFR-NOTI-06", "Nội dung hiển thị phải chống XSS; lỗi máy chủ không được trả chi tiết nhạy cảm cho trình duyệt."),
        ("NFR-NOTI-07", "Chỉ hoạt động trong phiên đã xác thực; POST thay đổi trạng thái phải có header X-Requested-With."),
        ("NFR-NOTI-08", "Giao diện popup phải hỗ trợ bàn phím/Escape, aria-label và trạng thái disabled rõ ràng."),
        ("NFR-NOTI-09", "Lỗi đồng bộ nguồn ngoài phải được Trace Warning; lỗi xử lý chính phải Trace Error để giám sát."),
        ("NFR-NOTI-10", "Index phải hỗ trợ truy vấn theo thời gian, target và user để polling không làm suy giảm hiệu năng."),
    ])

    doc.add_heading("6.3.7. Tiêu chí nghiệm thu và truy vết", level=3)
    tests = [
        ("TC-NOTI-01", "Mở ChartReport bằng phiên hợp lệ", "Badge và danh sách được tải tự động"),
        ("TC-NOTI-02", "Chờ trên 30 giây sau khi tạo thông báo", "Thông báo mới xuất hiện không cần reload trang"),
        ("TC-NOTI-03", "Polling khi request trước chưa xong", "Không chạy chồng; chỉ tải lại một lượt chờ"),
        ("TC-NOTI-04", "Có trên 99 thông báo chưa đọc", "Badge hiển thị 99+ và aria-label chứa số thực"),
        ("TC-NOTI-05", "Đánh dấu một thông báo đã đọc", "Dòng biến mất, badge giảm một và DB ghi theo user"),
        ("TC-NOTI-06", "Đánh dấu tất cả đã đọc", "Badge về 0, popup rỗng; user khác không bị ảnh hưởng"),
        ("TC-NOTI-07", "Thông báo cá nhân", "Chỉ người nhận được chỉ định nhìn thấy"),
        ("TC-NOTI-08", "Nội dung chứa thẻ script/HTML", "Chỉ hiển thị dạng văn bản, không thực thi"),
        ("TC-NOTI-09", "Email API trả nhiều attachment cùng Message-ID", "Chỉ tạo một thông báo email"),
        ("TC-NOTI-10", "Đồng bộ lại cùng email", "Không tạo trùng; nội dung được MERGE khi thay đổi"),
        ("TC-NOTI-11", "Email API lỗi/timeout", "Thông báo nội bộ vẫn tải được"),
        ("TC-NOTI-12", "Phiên không hợp lệ gọi handler", "HTTP 401"),
        ("TC-NOTI-13", "POST không có X-Requested-With", "HTTP 403"),
        ("TC-NOTI-14", "Trang Xem tất cả lọc và phân trang", "Đúng trạng thái, tổng số và 100 dòng/trang"),
        ("TC-NOTI-15", "Oracle lỗi khi đồng bộ email", "Transaction rollback, không có dữ liệu dở dang"),
    ]
    table(doc, ["Mã kiểm thử", "Nội dung", "Kết quả mong đợi"], tests)
    table(doc, ["Nhóm yêu cầu", "Kiểm thử"], [
        ("FR-NOTI-01…FR-NOTI-08", "TC-NOTI-01…TC-NOTI-04, TC-NOTI-08, TC-NOTI-12"),
        ("FR-NOTI-09…FR-NOTI-13", "TC-NOTI-05…TC-NOTI-07, TC-NOTI-13…TC-NOTI-14"),
        ("FR-NOTI-14…FR-NOTI-20", "TC-NOTI-09…TC-NOTI-11, TC-NOTI-15"),
        ("NFR-NOTI-01…NFR-NOTI-10", "TC-NOTI-02…TC-NOTI-04, TC-NOTI-08, TC-NOTI-11…TC-NOTI-15"),
    ])


def add_aerosync_specification(doc):
    if not AEROSYNC_SOURCE.exists():
        raise FileNotFoundError(f"Thiếu tài liệu nguồn FR-INT-001: {AEROSYNC_SOURCE}")
    source = Document(AEROSYNC_SOURCE)

    doc.add_heading("5.3. Đặc tả chi tiết FR-INT-001 – VATM AeroSync", level=2)
    doc.add_paragraph(
        "VATM AeroSync tự động phát hiện tệp từ email hoặc thư mục Incoming, kiểm tra và "
        "chuẩn hóa dữ liệu, đồng bộ giấy phép bay phù hợp vào Oracle ATFM, lưu dấu vết "
        "xử lý và cung cấp giao diện giám sát cho người vận hành. Nội dung dưới đây được "
        "chuẩn hóa từ tài liệu SRS_VATM_AeroSync_Hien_Tai.docx."
    )

    doc.add_heading("5.3.1. Thông tin đặc tả", level=3)
    source_table(doc, source, 1)
    doc.add_heading("5.3.2. Phạm vi", level=3)
    doc.add_paragraph("Trong phạm vi:")
    for item in [
        "Quét định kỳ hộp thư IMAP và thư mục Incoming; giới hạn số lượng tệp mỗi chu kỳ.",
        "Kiểm tra trùng nội dung bằng SHA-256, tạo công việc và phân phối qua RabbitMQ.",
        "Nhận dạng giấy phép Word theo profile YAML; trích xuất, chuẩn hóa và kiểm tra dữ liệu.",
        "Ghi giấy phép vào Oracle ATFM bằng giao dịch master/detail và cơ chế chống ghi trùng.",
        "Lưu trữ tệp theo kết quả, ghi audit, cảnh báo và giám sát qua REST API/WinUI.",
        "Tra cứu báo cáo email, resend attachment và retry/replay theo các chốt an toàn.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph("Ngoài phạm vi/giới hạn hiện tại:")
    source_table(doc, source, 2)

    doc.add_heading("5.3.3. Tác nhân và trách nhiệm", level=3)
    source_table(doc, source, 3)
    doc.add_paragraph(
        "Các vai trò trên là vai trò nghiệp vụ. Phiên bản hiện tại chưa áp đặt RBAC theo "
        "từng người dùng; môi trường triển khai phải kiểm soát quyền mở UI và truy cập API."
    )

    doc.add_heading("5.3.4. Kiến trúc và luồng xử lý", level=3)
    doc.add_paragraph("Thành phần hệ thống:")
    source_table(doc, source, 4)
    doc.add_paragraph("Luồng xử lý chính:")
    source_table(doc, source, 5)
    doc.add_paragraph("Giao diện REST chính:")
    source_table(doc, source, 6)

    doc.add_heading("5.3.5. Trạng thái xử lý", level=3)
    doc.add_paragraph("Trạng thái công việc đồng bộ:")
    source_table(doc, source, 7)
    doc.add_paragraph("Trạng thái import giấy phép:")
    source_table(doc, source, 8)
    doc.add_paragraph("Trạng thái email, file, archive và acknowledgement:")
    source_table(doc, source, 9)

    doc.add_heading("5.3.6. Yêu cầu chức năng chi tiết", level=3)
    groups = [
        (10, "Tiếp nhận và tạo công việc"),
        (11, "Kiểm định, nhận dạng và chuẩn hóa"),
        (12, "Đồng bộ Oracle ATFM"),
        (13, "Lưu trữ, audit và phản hồi email"),
        (14, "Giám sát, cấu hình và báo cáo"),
        (15, "Retry, resend và test replay"),
    ]
    for index, title in groups:
        doc.add_paragraph(title).runs[0].bold = True
        source_table(doc, source, index)

    doc.add_heading("5.3.7. Yêu cầu dữ liệu", level=3)
    doc.add_paragraph("Dữ liệu theo dõi và dữ liệu mục tiêu:")
    source_table(doc, source, 16)
    doc.add_paragraph("Dữ liệu giấy phép Word:")
    source_table(doc, source, 17)
    doc.add_paragraph(
        "Danh mục hiện trạng có 127 profile YAML. Đây không phải giới hạn thiết kế; profile "
        "mới phải có regression test và mapping tham chiếu cần thiết."
    )

    doc.add_heading("5.3.8. Quy tắc nghiệp vụ", level=3)
    source_table(doc, source, 18)
    doc.add_heading("5.3.9. Yêu cầu phi chức năng", level=3)
    source_table(doc, source, 19)

    doc.add_heading("5.3.10. Tiêu chí kiểm thử và truy vết", level=3)
    table(doc, ["Nhóm yêu cầu", "Phạm vi kiểm thử", "Bằng chứng dự kiến"], [
        ("FR-AS-01…FR-AS-09", "Quét nguồn, checkpoint, SHA-256, tạo job và publish queue", "Job/File/Email metadata; log RabbitMQ"),
        ("FR-AS-10…FR-AS-19", "Validate, profile recognition, parse, normalize và quarantine", "Chi tiết job; lỗi theo dòng; file archive"),
        ("FR-AS-20…FR-AS-27", "Redis lock, dry-run, giao dịch Oracle, duplicate và revision", "PermitImport; audit; dữ liệu master/detail"),
        ("FR-AS-28…FR-AS-32", "Archive, audit, acknowledgement và retention", "Cây thư mục; audit; trạng thái email"),
        ("FR-AS-33…FR-AS-39", "REST API, cấu hình, dashboard, báo cáo và OpenAPI", "Response API; ảnh UI; OpenAPI"),
        ("FR-AS-40…FR-AS-45", "Retry, resend cleanup và test replay", "Audit; trạng thái reset; kết quả replay"),
        ("NFR-AS-01…NFR-AS-16", "Hiệu năng, an toàn, idempotency, bảo mật và quan sát", "Kết quả đo; cấu hình; log; biên bản"),
    ])


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)
    font(doc.styles["Normal"], 12)
    font(doc.styles["Title"], 20, True)
    font(doc.styles["Heading 1"], 15, True)
    font(doc.styles["Heading 2"], 13, True)
    font(doc.styles["Heading 3"], 12, True)

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(55)
    run = cover.add_run("TÀI LIỆU ĐẶC TẢ YÊU CẦU PHẦN MỀM")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(20)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("HỆ THỐNG QUẢN LÝ, ĐIỀU HÀNH BAY – ATFM_WEB")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(17)
    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    line.paragraph_format.space_before = Pt(30)
    line.add_run("KHUNG SƯỜN TÀI LIỆU – BƯỚC 1").bold = True
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_before = Pt(50)
    info.add_run(
        "Mã tài liệu: SRS-ATFM-WEB\n"
        "Phiên bản: 0.1 – Khung sườn\n"
        "Ngày lập: 12/08/2026\n"
        "Đơn vị: ........................................................\n"
        "Người lập: ...................................................."
    )
    doc.add_page_break()

    doc.add_heading("KIỂM SOÁT TÀI LIỆU", level=1)
    table(doc, ["Thuộc tính", "Nội dung"], [
        ("Tên tài liệu", "Đặc tả yêu cầu phần mềm ATFM_WEB"),
        ("Mã tài liệu", "SRS-ATFM-WEB"),
        ("Phiên bản", "0.1"),
        ("Trạng thái", "Khung sườn – chờ bổ sung đặc tả chi tiết"),
        ("Tài liệu đầu vào", "Plan_20-25.7.2026.docx; NHAT_KY_THAY_DOI.txt; mã nguồn ATFM_WEB"),
    ])
    doc.add_heading("Lịch sử thay đổi", level=2)
    table(doc, ["Phiên bản", "Ngày", "Nội dung", "Người thực hiện"], [
        ("0.1", "12/08/2026", "Tạo bìa và khung sườn SRS", "Codex"),
    ])
    doc.add_heading("Phê duyệt tài liệu", level=2)
    table(doc, ["Vai trò", "Họ tên", "Chữ ký", "Ngày"], [
        ("Người lập", "", "", ""),
        ("Người kiểm tra", "", "", ""),
        ("Người phê duyệt", "", "", ""),
    ])

    doc.add_heading("MỤC LỤC DỰ KIẾN", level=1)
    for item in [
        "1. Giới thiệu", "2. Tổng quan hệ thống", "3. Phạm vi và danh mục chức năng",
        "4. Yêu cầu chung", "5. Phân hệ tích hợp và tự động hóa dữ liệu",
        "6. Phân hệ cảnh báo và tiện ích hỗ trợ", "7. Phân hệ nâng cấp và tối ưu HTSLB",
        "8. Phân hệ phân tích và báo cáo thông minh", "9. Phân hệ ứng dụng AI",
        "10. Yêu cầu dữ liệu và tích hợp", "11. Yêu cầu phi chức năng",
        "12. Yêu cầu triển khai và vận hành", "13. Kiểm thử và nghiệm thu",
        "14. Ma trận truy vết", "15. Phụ lục",
    ]:
        doc.add_paragraph(item, style="List Number")
    doc.add_page_break()

    doc.add_heading("1. Giới thiệu", level=1)
    for title in ["1.1. Mục đích", "1.2. Phạm vi", "1.3. Đối tượng sử dụng tài liệu", "1.4. Thuật ngữ và từ viết tắt", "1.5. Tài liệu tham chiếu"]:
        doc.add_heading(title, level=2)
        placeholder(doc)

    doc.add_heading("2. Tổng quan hệ thống", level=1)
    for title in ["2.1. Bối cảnh nghiệp vụ", "2.2. Mục tiêu nâng cấp", "2.3. Các bên liên quan", "2.4. Nhóm người dùng", "2.5. Kiến trúc và sơ đồ ngữ cảnh", "2.6. Giả định và phụ thuộc"]:
        doc.add_heading(title, level=2)
        placeholder(doc)

    doc.add_heading("3. Phạm vi và danh mục chức năng", level=1)
    rows = []
    for roman, name, items in SECTIONS:
        for index, item in enumerate(items, 1):
            rows.append((f"{roman}.{index:02d}", name, item, "Chờ đặc tả"))
    table(doc, ["Mã", "Phân hệ", "Chức năng/hạng mục", "Trạng thái SRS"], rows)

    doc.add_heading("4. Yêu cầu chung", level=1)
    for title in ["4.1. Đăng nhập và quản lý phiên", "4.2. Phân quyền", "4.3. Tra cứu và phân trang", "4.4. Nhập liệu và kiểm tra dữ liệu", "4.5. Nhật ký và truy vết", "4.6. Xuất báo cáo/điện văn"]:
        doc.add_heading(title, level=2)
        placeholder(doc)

    chapter = 5
    for roman, name, items in SECTIONS:
        doc.add_heading(f"{chapter}. {name.title()}", level=1)
        doc.add_heading(f"{chapter}.1. Mục tiêu và phạm vi phân hệ", level=2)
        placeholder(doc)
        doc.add_heading(f"{chapter}.2. Danh mục yêu cầu chức năng", level=2)
        function_rows = []
        prefix = ["INT", "ALT", "OPT", "RPT", "AI"][chapter - 5]
        for index, item in enumerate(items, 1):
            if chapter == 9:
                function_rows.append((f"FR-{prefix}-{index:03d}", item, "Đã tích hợp đặc tả NL2SQL/Vanna Oracle", "Theo FR-NL/BR-NL/NFR-NL và bằng chứng kiểm thử"))
            else:
                function_rows.append((f"FR-{prefix}-{index:03d}", item, "Chờ đặc tả", "Chờ xây dựng"))
        table(doc, ["Mã yêu cầu", "Tên yêu cầu", "Nội dung", "Tiêu chí nghiệm thu"], function_rows)
        if chapter == 5:
            add_aerosync_specification(doc)
            for args in INTEGRATION_SOURCES:
                add_integration_source(doc, *args)
            next_section = 8
        elif chapter == 6:
            add_alt001_specification(doc)
            add_alt002_specification(doc)
            next_section = 5
        elif chapter == 9:
            add_ai_nl2sql_specification(doc)
            next_section = 5
        elif chapter == 7:
            add_opt001_003_specification(doc)
            add_opt006_010_specification(doc)
            next_section = 14
        else:
            next_section = 3
        doc.add_heading(f"{chapter}.{next_section}. Quy tắc nghiệp vụ chung của phân hệ", level=2)
        placeholder(doc)
        doc.add_heading(f"{chapter}.{next_section + 1}. Luồng xử lý và ngoại lệ chung", level=2)
        placeholder(doc)
        chapter += 1

    doc.add_heading("10. Yêu cầu dữ liệu và tích hợp", level=1)
    for title in ["10.1. Mô hình và từ điển dữ liệu", "10.2. Oracle và các kho dữ liệu", "10.3. ADS-B", "10.4. Email và thư mục lưu trữ", "10.5. AMHS/AFTN", "10.6. Bravo", "10.7. API chia sẻ dữ liệu", "10.8. Chuyển đổi, đối soát và lưu trữ dữ liệu"]:
        doc.add_heading(title, level=2)
        placeholder(doc)

    doc.add_heading("11. Yêu cầu phi chức năng", level=1)
    table(doc, ["Mã nhóm", "Nhóm yêu cầu", "Chỉ tiêu cần đặc tả"], [
        ("NFR-PERF", "Hiệu năng", "Thời gian phản hồi, số người dùng, thời gian export/build"),
        ("NFR-SEC", "An toàn thông tin", "Xác thực, phân quyền, bảo vệ dữ liệu, quản lý bí mật"),
        ("NFR-AVAIL", "Sẵn sàng và tin cậy", "Tỷ lệ sẵn sàng, xử lý lỗi, retry"),
        ("NFR-AUDIT", "Nhật ký và kiểm toán", "Sự kiện, nội dung, thời hạn lưu"),
        ("NFR-BACKUP", "Sao lưu và khôi phục", "RTO, RPO, chu kỳ sao lưu"),
        ("NFR-COMPAT", "Tương thích", "Trình duyệt, hệ điều hành, phiên bản Oracle"),
        ("NFR-USABLE", "Khả năng sử dụng", "Giao diện, cảnh báo, khả năng tiếp cận"),
        ("NFR-MAINT", "Khả năng bảo trì", "Cấu hình, giám sát, triển khai và rollback"),
    ])

    doc.add_heading("12. Yêu cầu triển khai và vận hành", level=1)
    for title in ["12.1. Môi trường triển khai", "12.2. Cấu hình", "12.3. Quy trình cài đặt", "12.4. Giám sát và cảnh báo", "12.5. Rollback", "12.6. Hướng dẫn vận hành"]:
        doc.add_heading(title, level=2)
        placeholder(doc)

    doc.add_heading("13. Kiểm thử và nghiệm thu", level=1)
    table(doc, ["Nhóm", "Phạm vi dự kiến", "Tiêu chí đạt"], [
        ("Chức năng", "Luồng chính, ngoại lệ và phân quyền", "Mỗi FR có tối thiểu một test case"),
        ("Tích hợp", "Oracle, ADS-B, Email, AMHS, Bravo, API", "Dữ liệu đúng và có đối soát"),
        ("Hiệu năng", "Tra cứu, export, build, đồng bộ", "Đạt chỉ tiêu NFR đã phê duyệt"),
        ("An toàn thông tin", "Xác thực, phân quyền và dữ liệu đầu vào", "Không còn lỗi nghiêm trọng"),
        ("Nghiệm thu người dùng", "Các quy trình nghiệp vụ chính", "Có biên bản và bằng chứng"),
    ])

    doc.add_heading("14. Ma trận truy vết", level=1)
    table(doc, ["Mã yêu cầu", "Thiết kế/thành phần", "Mã kiểm thử", "Bằng chứng", "Kết quả"], [
        ("[FR-...]", "[Màn hình/API/Package]", "[TC-...]", "[Ảnh/Log/Biên bản]", "Chờ kiểm thử"),
    ])

    doc.add_heading("15. Phụ lục", level=1)
    for title in ["15.1. Danh mục màn hình và URL", "15.2. Danh mục API/Package", "15.3. Danh mục bảng dữ liệu", "15.4. Ma trận phân quyền", "15.5. Danh mục báo cáo/điện văn", "15.6. Biểu mẫu bằng chứng kiểm thử"]:
        doc.add_heading(title, level=2)
        placeholder(doc)

    for item in doc.sections:
        footer = item.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Trang ")
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = "PAGE"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.extend([begin, instruction, end])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
